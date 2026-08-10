"""
Ren'Py 错误修复工具
检查并修复常见语法错误，支持深度 Lint 检查
"""

import io
import json
import os
import re
import subprocess
from collections import Counter
from typing import List, Dict, Tuple, Optional
from pathlib import Path

from base.LogManager import LogManager
from module.Renpy.renpy_tl_core import (
    TlBlockKind,
    TlStmtKind,
    match_tpl_to_target,
    pair_old_new_lines,
    parse_tl_document,
)
from utils.call_game_python import get_python_path_from_game_path, get_py_path


class ErrorRepairer:
    """错误修复器"""

    # 仅匹配单行、双引号定界的翻译或简单 say 语句。suffix 被限制为
    # 空白/注释，避免跨过 menu 条件、screen action 和同一行的其他代码。
    RE_RENPY_STRING_LINE = re.compile(
        r'^(?P<indent>[ \t]*)(?:(?P<prefix>old|new|[A-Za-z_]\w*)(?P<separator>[ \t]+))?'
        r'"(?P<body>.*?)"(?P<suffix>[ \t]*(?:#.*)?)$'
    )
    RE_CURLY_DOUBLE_STRING_LINE = re.compile(
        r'^(?P<indent>[ \t]*)(?:(?P<prefix>old|new|[A-Za-z_]\w*)(?P<separator>[ \t]+))?'
        r'“(?P<body>.*?)”(?P<suffix>[ \t]*(?:#.*)?)$'
    )
    RE_CURLY_SINGLE_STRING_LINE = re.compile(
        r'^(?P<indent>[ \t]*)(?:(?P<prefix>old|new|[A-Za-z_]\w*)(?P<separator>[ \t]+))?'
        r'‘(?P<body>.*?)’(?P<suffix>[ \t]*(?:#.*)?)$'
    )
    NON_DIALOGUE_PREFIXES = {
        "add", "bar", "button", "call", "camera", "default", "define",
        "drag", "else", "elif", "fixed", "for", "frame", "grid", "hide",
        "hbox", "if", "image", "imagebutton", "init", "jump", "key",
        "layeredimage", "menu", "on", "pass", "pause", "play", "python",
        "queue", "return", "scene", "screen", "show", "stop", "style",
        "text", "textbutton", "timer", "transform", "use", "vbar", "vbox",
        "viewport", "voice", "while", "window", "with",
    }
    SAFE_LINT_FIX_TYPES = frozenset({
        "parse_error",
        "syntax_error",
        "indentation_mismatch",
        "indentation_level",
    })
    SCREEN_BLOCK_PREFIXES = (
        "textbutton",
        "imagebutton",
        "text",
        "frame",
        "hbox",
        "vbox",
        "grid",
        "fixed",
        "window",
        "viewport",
        "use",
        "button",
        "image",
        "add",
    )

    def __init__(self):
        self.logger = LogManager.get()
        self.errors_found = []

    def _split_line_ending(self, line: str) -> tuple[str, str]:
        """拆分行内容和换行符，修复后保持原始换行风格。"""
        if line.endswith("\r\n"):
            return line[:-2], "\r\n"
        if line.endswith("\n"):
            return line[:-1], "\n"
        if line.endswith("\r"):
            return line[:-1], "\r"
        return line, ""

    def _has_unescaped_quote(self, text: str, quote_char: str) -> bool:
        return bool(self._unescaped_quote_offsets(text, quote_char))

    def _unescaped_quote_offsets(self, text: str, quote_char: str) -> list[int]:
        offsets: list[int] = []
        escaped = False
        for offset, ch in enumerate(text):
            if escaped:
                escaped = False
                continue
            if ch == "\\":
                escaped = True
                continue
            if ch == quote_char:
                offsets.append(offset)
        return offsets

    def _escape_unescaped_quotes(self, text: str, quote_char: str) -> tuple[str, bool]:
        result: list[str] = []
        escaped = False
        changed = False

        for ch in text:
            if escaped:
                result.append(ch)
                escaped = False
                continue

            if ch == "\\":
                result.append(ch)
                escaped = True
                continue

            if ch == quote_char:
                result.append("\\")
                result.append(ch)
                changed = True
                continue

            result.append(ch)

        return "".join(result), changed

    def _looks_like_renpy_string_statement(self, prefix: Optional[str]) -> bool:
        """仅对类似对话/翻译语句的行执行未转义引号修复。"""
        stripped = (prefix or "").strip()
        if stripped == "":
            return True
        return stripped in {"old", "new"} or stripped not in self.NON_DIALOGUE_PREFIXES

    def _can_repair_unescaped_quotes(self, body: str) -> bool:
        """只接受成对的正文内双引号，并拒绝明显的多个字符串参数。"""
        offsets = self._unescaped_quote_offsets(body, '"')
        if len(offsets) < 2 or len(offsets) % 2 != 0:
            return False

        for opening, closing in zip(offsets[::2], offsets[1::2]):
            between = body[opening + 1:closing].strip()
            if not between:
                return False
            if re.fullmatch(r'(?:[,;+*/%|&]|==?|!=|<=?|>=?|\+|-)+', between):
                return False
        return True

    def repair_unescaped_dialogue_quotes(self, line: str) -> tuple[str, bool]:
        """保守修复 Ren'Py 对话/old/new 行中成对的未转义双引号。"""
        content, line_ending = self._split_line_ending(line)
        if content.strip() == "" or content.lstrip().startswith("#") or '"""' in content:
            return line, False

        match = self.RE_RENPY_STRING_LINE.match(content)
        if not match:
            return line, False

        indent = match.group("indent") or ""
        prefix = match.group("prefix")
        body = match.group("body")
        suffix = match.group("suffix") or ""

        if not self._looks_like_renpy_string_statement(prefix):
            return line, False
        if not self._can_repair_unescaped_quotes(body):
            return line, False

        escaped_body, changed = self._escape_unescaped_quotes(body, '"')
        if not changed:
            return line, False

        statement_prefix = f"{prefix}{match.group('separator')}" if prefix else ""
        repaired = f'{indent}{statement_prefix}"{escaped_body}"{suffix}{line_ending}'
        return repaired, True

    def repair_curly_string_delimiters(self, line: str) -> tuple[str, bool]:
        """仅规范化明确用作字符串外层定界符的中文弯引号。"""
        content, line_ending = self._split_line_ending(line)
        if content.strip() == "" or content.lstrip().startswith("#"):
            return line, False

        for pattern, target_quote in (
            (self.RE_CURLY_DOUBLE_STRING_LINE, '"'),
            (self.RE_CURLY_SINGLE_STRING_LINE, "'"),
        ):
            match = pattern.match(content)
            if not match or not self._looks_like_renpy_string_statement(match.group("prefix")):
                continue

            body = match.group("body")
            # 变更外层定界符后，正文中的同类直引号会改变解析边界；无法
            # 无歧义判断其意图时保持原样。
            if self._has_unescaped_quote(body, target_quote):
                return line, False

            indent = match.group("indent") or ""
            prefix = match.group("prefix")
            statement_prefix = f"{prefix}{match.group('separator')}" if prefix else ""
            suffix = match.group("suffix") or ""
            repaired = (
                f"{indent}{statement_prefix}{target_quote}{body}{target_quote}"
                f"{suffix}{line_ending}"
            )
            return repaired, True

        return line, False

    def _detect_quote_issue(self, line: str) -> Optional[str]:
        """返回可高置信识别的引号问题；不对整行单双引号做奇偶计数。"""
        content, _ = self._split_line_ending(line)
        if content.strip() == "" or content.lstrip().startswith("#"):
            return None
        if '"""' in content or "'''" in content:
            return None

        for pattern in (self.RE_CURLY_DOUBLE_STRING_LINE, self.RE_CURLY_SINGLE_STRING_LINE):
            match = pattern.match(content)
            if match and self._looks_like_renpy_string_statement(match.group("prefix")):
                return "字符串使用了中文弯引号作为外层定界符"

        match = self.RE_RENPY_STRING_LINE.match(content)
        if match and self._looks_like_renpy_string_statement(match.group("prefix")):
            if self._can_repair_unescaped_quotes(match.group("body")):
                return "Ren'Py 字符串正文存在成对的未转义双引号"
            return None

        candidate = re.match(
            r'^[ \t]*(?:(?P<prefix>old|new|[A-Za-z_]\w*)[ \t]+)?"',
            content,
        )
        if candidate and self._looks_like_renpy_string_statement(candidate.group("prefix")):
            if len(self._unescaped_quote_offsets(content, '"')) % 2 != 0:
                return "双引号不匹配"
        return None

    @staticmethod
    def _decode_renpy_literal(raw_inner: str) -> str:
        """按扫描规则解码少量可靠的 Ren'Py 字符串转义。"""
        result: list[str] = []
        index = 0
        escape_map = {"n": "\n", "r": "\r", "t": "\t"}

        while index < len(raw_inner):
            char = raw_inner[index]
            if char != "\\" or index + 1 >= len(raw_inner):
                result.append(char)
                index += 1
                continue

            escaped = raw_inner[index + 1]
            if escaped in escape_map:
                result.append(escape_map[escaped])
            elif escaped in {'"', "'", "\\"}:
                result.append(escaped)
            else:
                # 未知转义保持原样，扫描器不能擅自改变其语义。
                result.extend(("\\", escaped))
            index += 2

        return "".join(result)

    @staticmethod
    def _scan_string_literals(code: str) -> list[dict]:
        """扫描一行代码中的单双引号字面量，忽略字符串外的注释。"""
        if '"""' in code or "'''" in code:
            return []

        literals: list[dict] = []
        index = 0
        while index < len(code):
            char = code[index]
            if char == "#":
                break
            if char not in {'"', "'"}:
                index += 1
                continue

            quote = char
            start = index
            index += 1
            raw: list[str] = []
            while index < len(code):
                char = code[index]
                if char == "\\" and index + 1 < len(code):
                    raw.extend((char, code[index + 1]))
                    index += 2
                    continue
                if char == quote:
                    raw_inner = "".join(raw)
                    literals.append({
                        "start": start,
                        "end": index + 1,
                        "raw": raw_inner,
                        "value": ErrorRepairer._decode_renpy_literal(raw_inner),
                    })
                    index += 1
                    break
                raw.append(char)
                index += 1
            else:
                break

        return literals

    @staticmethod
    def _extract_placeholders(text: str) -> Counter:
        """提取方括号插值与花括号文本标签，并忽略转义的双括号。"""
        placeholders: Counter = Counter()
        index = 0

        while index < len(text):
            opening = text[index]
            if opening not in {"[", "{"}:
                index += 1
                continue

            closing = "]" if opening == "[" else "}"
            doubled_open = opening * 2
            doubled_close = closing * 2
            if text.startswith(doubled_open, index):
                escaped_end = text.find(doubled_close, index + 2)
                index = escaped_end + 2 if escaped_end >= 0 else index + 2
                continue

            end = text.find(closing, index + 1)
            if end < 0:
                break

            body = text[index + 1:end]
            if body.strip() and opening not in body:
                placeholders[f"{opening}{body}{closing}"] += 1
            index = end + 1

        return placeholders

    @staticmethod
    def _classify_placeholder_diff(
        source: Counter,
        target: Counter,
    ) -> tuple[Counter, Counter, list[dict]]:
        """把 Counter 差异拆成缺失、额外和同类括号内的疑似改写。"""
        missing = source - target
        extra = target - source
        rewrite_counts: Counter = Counter()

        for opening in ("[", "{"):
            missing_tokens = sorted(
                token
                for token, count in missing.items()
                for _ in range(count)
                if token.startswith(opening)
            )
            extra_tokens = sorted(
                token
                for token, count in extra.items()
                for _ in range(count)
                if token.startswith(opening)
            )
            for source_token, target_token in zip(missing_tokens, extra_tokens):
                missing[source_token] -= 1
                extra[target_token] -= 1
                rewrite_counts[(source_token, target_token)] += 1

        missing += Counter()
        extra += Counter()
        rewrites = [
            {"source": source_token, "target": target_token, "count": count}
            for (source_token, target_token), count in sorted(rewrite_counts.items())
        ]
        return missing, extra, rewrites

    @staticmethod
    def _counter_details(counter: Counter) -> Dict[str, int]:
        return {token: counter[token] for token in sorted(counter)}

    def _iter_translation_pairs(self, lines: List[str]) -> list[dict]:
        """返回 tl 文档中可可靠配对的模板/译文语句。"""
        content_lines = [self._split_line_ending(line)[0] for line in lines]
        document = parse_tl_document(content_lines)
        pairs: list[dict] = []

        for block in document.blocks:
            if block.kind == TlBlockKind.STRINGS:
                mapping = pair_old_new_lines(block)
            elif block.kind == TlBlockKind.LABEL:
                mapping = match_tpl_to_target(block)
            else:
                continue

            statements = {statement.line_no: statement for statement in block.statements}
            for source_line, target_line in mapping.items():
                source = statements.get(source_line)
                target = statements.get(target_line)
                if source is None or target is None:
                    continue

                source_literals = self._scan_string_literals(source.code)
                target_literals = self._scan_string_literals(target.code)
                if not source_literals:
                    continue

                pairs.append({
                    "language": block.lang,
                    "is_old_new": block.kind == TlBlockKind.STRINGS,
                    "source": source,
                    "target": target,
                    "source_values": [
                        literal["value"]
                        for literal in source_literals
                    ],
                    "target_values": [
                        literal["value"]
                        for literal in target_literals
                    ],
                })

        return pairs

    def _duplicate_old_new_error(
        self,
        pair: dict,
        first_path: str,
        first_line: int,
    ) -> Dict:
        source = pair["source"]
        original = pair["source_values"]
        return {
            "line": source.line_no,
            "type": "duplicate_old_new",
            "message": "检测到重复 old/new 条目",
            "content": source.raw_line.strip(),
            "original": original[0] if len(original) == 1 else original,
            "first_file": first_path,
            "first_line": first_line,
        }

    def _scan_extra_empty_strings(self, lines: List[str]) -> List[Dict]:
        """报告简单翻译语句中与其他字面量相邻的多余空字符串。"""
        content_lines = [self._split_line_ending(line)[0] for line in lines]
        document = parse_tl_document(content_lines)
        errors: list[Dict] = []

        for block in document.blocks:
            if block.kind == TlBlockKind.PYTHON:
                continue
            for statement in block.statements:
                if statement.stmt_kind != TlStmtKind.TARGET:
                    continue

                literals = self._scan_string_literals(statement.code)
                if len(literals) < 2:
                    continue

                prefix = statement.code[:literals[0]["start"]].strip()
                if prefix and (
                    re.fullmatch(r"[A-Za-z_]\w*", prefix) is None
                    or not self._looks_like_renpy_string_statement(prefix)
                ):
                    continue

                components: list[list[dict]] = []
                component = [literals[0]]
                for previous, current in zip(literals, literals[1:]):
                    between = statement.code[previous["end"]:current["start"]]
                    if between.strip() == "":
                        component.append(current)
                    else:
                        components.append(component)
                        component = [current]
                components.append(component)

                extra_count = 0
                for component in components:
                    if len(component) < 2:
                        continue
                    empty_count = sum(literal["value"] == "" for literal in component)
                    if empty_count == 0:
                        continue
                    has_non_empty = empty_count < len(component)
                    extra_count += empty_count if has_non_empty else empty_count - 1
                if extra_count <= 0:
                    continue

                errors.append({
                    "line": statement.line_no,
                    "type": "extra_empty_string",
                    "message": f"翻译语句包含 {extra_count} 个多余空字符串",
                    "content": statement.raw_line.strip(),
                    "count": extra_count,
                })

        return errors

    def _scan_translation_issues(self, lines: List[str], file_path: str) -> List[Dict]:
        """纯读取扫描占位符、换行、空字符串和同文件重复条目。"""
        errors: list[Dict] = []
        seen_old: dict[tuple, tuple[str, int]] = {}

        for pair in self._iter_translation_pairs(lines):
            source = pair["source"]
            target = pair["target"]
            source_values = pair["source_values"]
            target_values = pair["target_values"]

            if pair["is_old_new"]:
                duplicate_key = (pair["language"], tuple(source_values))
                first = seen_old.get(duplicate_key)
                if first is None:
                    seen_old[duplicate_key] = (file_path, source.line_no)
                else:
                    errors.append(self._duplicate_old_new_error(pair, first[0], first[1]))

            source_placeholders: Counter = Counter()
            target_placeholders: Counter = Counter()
            source_linebreaks = 0
            target_linebreaks = 0
            compared = False
            for source_value, target_value in zip(source_values, target_values):
                # 空 new 是合法的未翻译占位，不属于翻译损坏。
                if target_value == "":
                    continue
                compared = True
                source_placeholders.update(self._extract_placeholders(source_value))
                target_placeholders.update(self._extract_placeholders(target_value))
                source_linebreaks += source_value.count("\n")
                target_linebreaks += target_value.count("\n")

            if not compared:
                continue

            missing, extra, rewrites = self._classify_placeholder_diff(
                source_placeholders,
                target_placeholders,
            )
            common = {
                "line": target.line_no,
                "content": target.raw_line.strip(),
                "source_line": source.line_no,
                "source_content": source.raw_line.strip(),
            }
            if rewrites:
                errors.append({
                    **common,
                    "type": "placeholder_rewritten",
                    "message": "译文中的占位符疑似被改写",
                    "rewrites": rewrites,
                })
            if missing:
                errors.append({
                    **common,
                    "type": "placeholder_missing",
                    "message": "译文缺少原文占位符",
                    "placeholders": self._counter_details(missing),
                })
            if extra:
                errors.append({
                    **common,
                    "type": "placeholder_extra",
                    "message": "译文包含原文没有的占位符",
                    "placeholders": self._counter_details(extra),
                })
            if source_linebreaks != target_linebreaks:
                errors.append({
                    **common,
                    "type": "linebreak_mismatch",
                    "message": "译文与原文的换行数量不一致",
                    "source_count": source_linebreaks,
                    "target_count": target_linebreaks,
                })

        errors.extend(self._scan_extra_empty_strings(lines))
        return errors

    def _get_indent_width(self, line: str) -> int:
        expanded = line.expandtabs(4)
        return len(expanded) - len(expanded.lstrip(" "))

    def _is_blank_or_comment(self, line: str) -> bool:
        stripped = line.strip()
        return stripped == "" or stripped.startswith("#")

    def _find_previous_non_empty_index(self, lines: List[str], idx: int) -> Optional[int]:
        for current in range(idx - 1, -1, -1):
            if not self._is_blank_or_comment(lines[current]):
                return current
        return None

    def _find_previous_block_indent_less_than(self, lines: List[str], idx: int, indent_limit: int) -> Optional[int]:
        for current in range(idx - 1, -1, -1):
            line = lines[current]
            if self._is_blank_or_comment(line):
                continue
            content, _ = self._split_line_ending(line)
            stripped = content.lstrip(" \t")
            indent = self._get_indent_width(content)
            if indent < indent_limit and stripped.endswith(":"):
                return indent
        return None

    def _looks_like_screen_block_line(self, stripped: str) -> bool:
        token = stripped.split(maxsplit = 1)[0] if stripped else ""
        return token in self.SCREEN_BLOCK_PREFIXES or stripped.startswith(("if ", "elif ", "else:", "for ", "while "))

    def repair_indentation_level(self, lines: List[str], idx: int) -> tuple[str, bool]:
        """尝试修复明显错位的块级缩进。"""
        if idx < 0 or idx >= len(lines):
            return lines[idx], False

        original_line = lines[idx]
        content, line_ending = self._split_line_ending(original_line)
        if self._is_blank_or_comment(content):
            return original_line, False

        stripped = content.lstrip(" \t")
        current_indent = self._get_indent_width(content)
        prev_idx = self._find_previous_non_empty_index(lines, idx)
        if prev_idx is None:
            return original_line, False

        prev_content, _ = self._split_line_ending(lines[prev_idx])
        prev_stripped = prev_content.lstrip(" \t")
        prev_indent = self._get_indent_width(prev_content)

        # 情况 1：上一行是块起始，当前行应作为其子块，统一为 +4 空格。
        if prev_stripped.endswith(":") and not re.match(r"^(elif|else|except|finally)\b", stripped):
            expected_indent = prev_indent + 4
            if current_indent != expected_indent:
                return (" " * expected_indent) + stripped + line_ending, True

        # 情况 2：当前行自身是块起始，但前一段刚从深层子块退出。
        # 这时按最近父块缩进回退一层，适合修复 screen 里 displayable / textbutton 错位。
        if stripped.endswith(":") and self._looks_like_screen_block_line(stripped) and prev_indent >= current_indent + 8:
            parent_indent = self._find_previous_block_indent_less_than(lines, idx, current_indent)
            if parent_indent is not None and parent_indent != current_indent:
                return (" " * parent_indent) + stripped + line_ending, True

        return original_line, False

    def check_file(
        self,
        file_path: str,
        check_syntax: bool = True,
        check_indent: bool = True,
        check_indent_level: bool = True,
        check_quotes: bool = True,
        check_dialogue_quotes: bool = True,
        encoding: str = "utf-8",
        check_translation_issues: bool = True,
    ) -> List[Dict[str, any]]:
        """
        检查单个文件

        Args:
            file_path: 文件路径
            check_syntax: 是否检查语法
            check_indent: 是否检查缩进
            check_quotes: 是否检查引号匹配
            encoding: 文件编码
            check_translation_issues: 是否扫描译文占位符、换行、空字符串和重复项

        Returns:
            错误列表 [{"line": 行号, "type": 错误类型, "message": 错误信息}, ...]
        """
        errors = []

        try:
            with open(file_path, "r", encoding=encoding, errors="ignore") as f:
                lines = f.readlines()

            for line_num, line in enumerate(lines, 1):
                # 语法检查
                if check_syntax:
                    # 检查 label 后是否有冒号
                    if line.strip().startswith("label ") and not line.strip().endswith(":"):
                        errors.append({
                            "line": line_num,
                            "type": "syntax",
                            "message": "label 语句缺少冒号",
                            "content": line.strip()
                        })

                    # 检查 if/elif/else/menu 后是否有冒号
                    if re.match(r'^\s*(if|elif|else|menu|while|for)\s', line) and \
                       not line.strip().endswith(":"):
                        errors.append({
                            "line": line_num,
                            "type": "syntax",
                            "message": "控制流语句缺少冒号",
                            "content": line.strip()
                        })

                # 缩进检查
                if check_indent:
                    # 检查是否使用 Tab
                    if "\t" in line:
                        errors.append({
                            "line": line_num,
                            "type": "indent",
                            "message": "使用了 Tab 而非空格缩进",
                            "content": line.strip()
                        })

                if check_indent_level:
                    _, changed = self.repair_indentation_level(lines, line_num - 1)
                    if changed:
                        errors.append({
                            "line": line_num,
                            "type": "indentation_level",
                            "message": "疑似块级缩进层级错误",
                            "content": line.strip()
                        })

                quote_issue_found = False
                if check_quotes:
                    quote_message = self._detect_quote_issue(line)
                    if quote_message:
                        errors.append({
                            "line": line_num,
                            "type": "quotes",
                            "message": quote_message,
                            "content": line.strip()
                        })
                        quote_issue_found = True

                if check_dialogue_quotes and not quote_issue_found:
                    _, changed = self.repair_unescaped_dialogue_quotes(line)
                    if changed:
                        errors.append({
                            "line": line_num,
                            "type": "unescaped_dialogue_quotes",
                            "message": "Ren'Py 对话行存在未转义引号",
                            "content": line.strip()
                        })

            if check_translation_issues:
                errors.extend(self._scan_translation_issues(lines, file_path))

        except Exception as e:
            self.logger.error(f"检查文件失败 {file_path}: {e}")

        return errors

    def check_folder(
        self,
        folder_path: str,
        check_syntax: bool = True,
        check_indent: bool = True,
        check_indent_level: bool = True,
        check_quotes: bool = True,
        check_dialogue_quotes: bool = True,
        encoding: str = "utf-8",
        check_translation_issues: bool = True,
    ) -> Dict[str, List[Dict]]:
        """
        批量检查文件夹

        Args:
            folder_path: 文件夹路径
            check_syntax: 是否检查语法
            check_indent: 是否检查缩进
            check_quotes: 是否检查引号匹配
            encoding: 文件编码
            check_translation_issues: 是否扫描译文占位符、换行、空字符串和重复项

        Returns:
            {文件路径: 错误列表}
        """
        all_errors = {}
        rpy_files = sorted(Path(folder_path).rglob("*.rpy"))

        self.logger.info(f"检查 {len(rpy_files)} 个 .rpy 文件")

        for file_path in rpy_files:
            errors = self.check_file(
                str(file_path),
                check_syntax=check_syntax,
                check_indent=check_indent,
                check_indent_level=check_indent_level,
                check_quotes=check_quotes,
                check_dialogue_quotes=check_dialogue_quotes,
                encoding=encoding,
                check_translation_issues=check_translation_issues,
            )

            if errors:
                all_errors[str(file_path)] = errors

        if check_translation_issues:
            seen_old: dict[tuple, tuple[str, int]] = {}
            for file_path in rpy_files:
                try:
                    with open(file_path, "r", encoding=encoding, errors="ignore") as source_file:
                        pairs = self._iter_translation_pairs(source_file.readlines())
                except Exception as e:
                    self.logger.error(f"检查跨文件重复失败 {file_path}: {e}")
                    continue

                for pair in pairs:
                    if not pair["is_old_new"]:
                        continue
                    duplicate_key = (pair["language"], tuple(pair["source_values"]))
                    first = seen_old.get(duplicate_key)
                    if first is None:
                        seen_old[duplicate_key] = (str(file_path), pair["source"].line_no)
                        continue

                    file_errors = all_errors.setdefault(str(file_path), [])
                    line_number = pair["source"].line_no
                    already_reported = any(
                        error.get("type") == "duplicate_old_new"
                        and error.get("line") == line_number
                        for error in file_errors
                    )
                    if not already_reported:
                        file_errors.append(
                            self._duplicate_old_new_error(pair, first[0], first[1])
                        )

        total_errors = sum(len(errs) for errs in all_errors.values())
        self.logger.info(f"检查完成: 发现 {total_errors} 个错误")

        return all_errors

    def auto_fix_file(
        self,
        file_path: str,
        fix_indent: bool = True,
        fix_indent_level: bool = False,
        fix_quotes: bool = False,
        fix_dialogue_quotes: bool = False,
        encoding: str = "utf-8"
    ) -> Tuple[bool, int]:
        """
        自动修复文件

        Args:
            file_path: 文件路径
            fix_indent: 是否修复缩进
            fix_quotes: 是否尝试修复引号 (危险操作)
            encoding: 文件编码

        Returns:
            (是否成功, 修复数量)
        """
        try:
            with open(file_path, "r", encoding=encoding, errors="ignore") as f:
                lines = f.readlines()

            new_lines = []
            fix_count = 0

            for idx, line in enumerate(lines):
                new_line = line

                # 修复缩进: Tab 转空格
                if fix_indent and "\t" in line:
                    new_line = new_line.replace("\t", "    ")
                    fix_count += 1

                if fix_dialogue_quotes:
                    repaired_line, changed = self.repair_unescaped_dialogue_quotes(new_line)
                    if changed:
                        new_line = repaired_line
                        fix_count += 1

                # 中文弯引号只有在明确充当整条字符串的外层定界符时才转换。
                # 字符串正文中的正常中文引号必须原样保留。
                if fix_quotes:
                    repaired_line, changed = self.repair_curly_string_delimiters(new_line)
                    if changed:
                        new_line = repaired_line
                        fix_count += 1

                if fix_indent_level:
                    temp_lines = list(lines)
                    temp_lines[idx] = new_line
                    repaired_line, changed = self.repair_indentation_level(temp_lines, idx)
                    if changed:
                        new_line = repaired_line
                        fix_count += 1

                new_lines.append(new_line)

            # 写回文件
            if fix_count > 0:
                with open(file_path, "w", encoding=encoding) as f:
                    f.writelines(new_lines)
                self.logger.info(f"修复完成: {file_path} ({fix_count} 处)")

            return True, fix_count

        except Exception as e:
            self.logger.error(f"修复失败 {file_path}: {e}")
            return False, 0

    def export_error_report(self, errors: Dict[str, List[Dict]], output_path: str):
        """
        导出错误报告为 Excel

        Args:
            errors: 错误字典 {文件路径: 错误列表}
            output_path: 输出路径
        """
        try:
            import openpyxl
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Error Report"

            # 写入表头
            headers = ["文件", "行号", "错误类型", "错误信息", "内容", "详情"]
            ws.append(headers)

            # 写入数据
            for file_path, file_errors in errors.items():
                for error in file_errors:
                    details = {
                        key: value
                        for key, value in error.items()
                        if key not in {"line", "type", "message", "content"}
                    }
                    ws.append([
                        file_path,
                        error.get("line", 0),
                        error.get("type", ""),
                        error.get("message", ""),
                        error.get("content", ""),
                        json.dumps(details, ensure_ascii=False, sort_keys=True),
                    ])

            wb.save(output_path)
            self.logger.info(f"错误报告已导出: {output_path}")

        except Exception as e:
            self.logger.error(f"导出错误报告失败: {e}")

    # ========== 深度 Lint 检查 (调用 Ren'Py 内置 lint) ==========
    
    def exec_renpy_lint(self, game_path: str) -> Optional[str]:
        """
        执行 Ren'Py 内置 lint 命令
        
        Args:
            game_path: 游戏可执行文件路径 (.exe)
            
        Returns:
            None 表示执行失败，空字符串表示执行成功且无输出，非空字符串
            表示执行成功并获得 Lint 输出。
        """
        try:
            python_path = get_python_path_from_game_path(game_path)
            py_path = get_py_path(game_path)
            game_dir = os.path.dirname(game_path)
            
            if not python_path or not os.path.isfile(python_path):
                self.logger.error(f"找不到游戏 Python: {python_path}")
                return None
                
            if not os.path.isfile(py_path):
                self.logger.error(f"找不到游戏 .py 文件: {py_path}")
                return None
            
            # 错误输出文件
            error_output = os.path.join(game_dir, "lint_errors.txt")
            
            # 使用参数列表执行，避免 shell 对游戏路径做二次解释。
            command = [python_path, "-O", py_path, game_dir, "lint"]
            
            self.logger.info(f"执行 Lint 命令: {command}")
            
            # 执行命令
            result = subprocess.run(
                command,
                shell=False,
                capture_output=True,
                text=True,
                encoding='utf-8',
                errors='ignore',
                creationflags=0x08000000 if os.name == 'nt' else 0,  # CREATE_NO_WINDOW
                cwd=game_dir
            )
            
            # 合并 stdout 和 stderr
            output = (result.stdout or "") + (result.stderr or "")

            if result.returncode != 0:
                self.logger.error(f"Lint 执行失败，退出码: {result.returncode}")
                if output.strip():
                    self.logger.error(output.strip())
                return None
            
            if output.strip():
                # 保存到文件
                with open(error_output, "w", encoding="utf-8") as f:
                    f.write(output)
                self.logger.info(f"Lint 结果已保存到: {error_output}")
                return output
            else:
                self.logger.info("Lint 检查完成，未发现错误")
                return ""
                
        except Exception as e:
            self.logger.error(f"执行 Lint 失败: {e}")
            return None
    
    def parse_lint_errors(self, lint_output: str) -> List[Dict]:
        """
        解析 Lint 输出内容
        
        Args:
            lint_output: lint 命令的输出
            
        Returns:
            错误列表
        """
        errors = []
        
        if not lint_output:
            return errors
            
        for line in lint_output.split('\n'):
            line = line.strip()
            if not line:
                continue
                
            error_info = {"raw": line}
            
            # 解析常见错误格式，保留路径中的空格。
            location_match = re.match(
                r'^(?:File\s+)?["\']?(?P<file>.+?\.rpy)["\']?,\s*'
                r'line\s+(?P<line>\d+):\s*(?P<message>.*)$',
                line,
                re.IGNORECASE,
            )
            if location_match:
                error_info["file"] = location_match.group("file")
                error_info["line"] = int(location_match.group("line"))
                error_info["message"] = location_match.group("message").strip()

            # 解析翻译重复错误
            elif line.startswith('Exception: A translation for '):
                error_info["type"] = "duplicate_translation"
                if 'already exists at ' in line:
                    idx = line.rindex('already exists at ')
                    location = line[idx + len('already exists at '):].rstrip('.')
                    if ':' in location:
                        file_part, line_part = location.rsplit(':', 1)
                        error_info["file"] = file_part.strip()
                        try:
                            error_info["line"] = int(line_part.strip())
                        except ValueError:
                            pass
                error_info["message"] = line
            
            # 标记错误类型
            if 'is not terminated with a newline' in line:
                error_info["type"] = "unterminated_string"
            elif 'end of line expected' in line:
                error_info["type"] = "syntax_error"
            elif 'expects a non-empty block' in line:
                error_info["type"] = "empty_block"
            elif 'unknown statement' in line:
                error_info["type"] = "unknown_statement"
            elif 'expected statement' in line:
                error_info["type"] = "expected_statement"
            elif 'Could not parse string' in line:
                error_info["type"] = "parse_error"
            elif 'Indentation mismatch' in line:
                error_info["type"] = "indentation_mismatch"
                 
            errors.append(error_info)
            
        return errors
    
    def fix_by_lint(self, game_path: str, max_iterations: int = 16) -> Tuple[bool, int]:
        """
        通过 Lint 检查自动修复错误（递归修复）
        
        Args:
            game_path: 游戏可执行文件路径
            max_iterations: 最大迭代次数
            
        Returns:
            (是否已由 Lint 验证无错误, 安全修复的数量)
        """
        if max_iterations <= 0:
            return False, 0

        total_fixed = 0
        game_dir = os.path.dirname(game_path)
        game_root = Path(game_dir).resolve()

        for iteration in range(max_iterations):
            self.logger.info(f"开始第 {iteration + 1}/{max_iterations} 轮 Lint 检查...")

            # 执行 lint
            lint_output = self.exec_renpy_lint(game_path)

            if lint_output is None:
                self.logger.error("Lint 执行失败，停止自动修复")
                return False, total_fixed
            if lint_output == "":
                self.logger.info("没有更多错误，修复完成!")
                return True, total_fixed

            # 解析错误
            errors = self.parse_lint_errors(lint_output)

            if not errors:
                self.logger.warning("Lint 有输出，但没有可解析的错误；未修改文件")
                return False, total_fixed

            # 每轮最多改一处，然后重新执行 lint，避免继续使用已经失效的行号。
            fixed_in_round = False
            processed_locations: set[tuple[str, int]] = set()
            for error in errors:
                file_path = error.get("file")
                line_num = error.get("line")
                error_type = error.get("type", "")

                if error_type not in self.SAFE_LINT_FIX_TYPES:
                    continue
                if not file_path or not isinstance(line_num, int) or line_num < 1:
                    continue

                # 构建完整路径
                candidate = Path(file_path)
                if not candidate.is_absolute():
                    candidate = game_root / candidate
                try:
                    candidate = candidate.resolve()
                    candidate.relative_to(game_root)
                except (OSError, ValueError):
                    self.logger.warning(f"忽略游戏目录外的 Lint 路径: {file_path}")
                    continue

                if not candidate.is_file():
                    self.logger.warning(f"文件不存在: {candidate}")
                    continue

                location = (str(candidate), line_num)
                if location in processed_locations:
                    continue
                processed_locations.add(location)

                if self._fix_single_lint_error(str(candidate), line_num, error_type):
                    fixed_in_round = True
                    total_fixed += 1
                    break

            if not fixed_in_round:
                self.logger.warning("没有可安全自动修复的 Lint 错误；未继续修改文件")
                return False, total_fixed

            self.logger.info(f"第 {iteration + 1} 轮安全修复了 1 个错误")

        self.logger.warning("达到 Lint 自动修复轮次上限，结果尚未验证为无错误")
        return False, total_fixed
    
    def _fix_single_lint_error(self, file_path: str, line_num: int, error_type: str) -> bool:
        """
        修复单个 Lint 错误
        
        Args:
            file_path: 文件路径
            line_num: 错误行号 (1-based)
            error_type: 错误类型
            
        Returns:
            是否修复成功
        """
        if error_type not in self.SAFE_LINT_FIX_TYPES:
            return False

        try:
            with io.open(file_path, 'r', encoding='utf-8') as f:
                lines = f.readlines()
            
            if line_num < 1 or line_num > len(lines):
                return False
                
            idx = line_num - 1  # 转为 0-based
            self.logger.info(f"修复 {file_path} 第 {line_num} 行: {lines[idx].strip()[:50]}...")

            changed = False
            if error_type in {"parse_error", "syntax_error"}:
                candidate_indices = [idx]
                if error_type == "syntax_error" and idx > 0:
                    candidate_indices.append(idx - 1)
                for candidate_idx in candidate_indices:
                    repaired_line, repaired = self.repair_unescaped_dialogue_quotes(
                        lines[candidate_idx]
                    )
                    if repaired:
                        lines[candidate_idx] = repaired_line
                        changed = True
                        break
            elif error_type in {"indentation_mismatch", "indentation_level"}:
                candidate_indices = [idx] + ([idx - 1] if idx > 0 else [])
                for candidate_idx in candidate_indices:
                    repaired_line, repaired = self.repair_indentation_level(
                        lines,
                        candidate_idx,
                    )
                    if repaired:
                        lines[candidate_idx] = repaired_line
                        changed = True
                        break

            if not changed:
                return False

            # 保持行数不变，避免本轮错误位置整体漂移。
            with io.open(file_path, 'w', encoding='utf-8') as f:
                f.writelines(lines)

            return True

        except Exception as e:
            self.logger.error(f"修复错误失败 {file_path}:{line_num} - {e}")
            return False
    
    def _remove_consecutive_empty_lines(self, lines: List[str]) -> List[str]:
        """移除连续的空行，最多保留一个"""
        result = []
        prev_empty = False
        
        for line in lines:
            is_empty = not line.strip()
            if is_empty and prev_empty:
                continue
            result.append(line)
            prev_empty = is_empty
            
        return result
