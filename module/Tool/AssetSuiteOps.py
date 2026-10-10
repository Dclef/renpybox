"""资源套件共享业务：批量修正与姓名提取（供 API 与旧 Qt 共用）。

不含 Qt / FastAPI 依赖。路径写入前校验根目录边界；Excel 单元格按纯文本写出。
"""

from __future__ import annotations

import ast
import json
import os
import re
import shutil
import stat
import tempfile
import time
from pathlib import Path, PureWindowsPath
from typing import Any, Callable, Iterable, Mapping, Sequence

import openpyxl
from openpyxl.cell.cell import Cell

from module.File.AtomicWrite import atomic_write_text
from module.Renpy.renpy_tl_core import escape_tl_string

FILE_NAME_WHITELIST = re.compile(
    r"^(结果检查_|result_check_)([^\\/]+)\.json$",
    flags=re.IGNORECASE,
)
FILE_NAME_BLACKLIST = ("result_check_untranslated.json", "结果检查_未翻译.json")
BATCH_WORKBOOK_NAME = "批量修正.xlsx"
RE_RENPY_CHARACTER = re.compile(
    r'define\s+(\w+)\s*=\s*Character\s*\(\s*["\']([^"\']+)["\']',
    re.MULTILINE,
)
RE_JSON_NAME_FIELD = re.compile(r'"name"\s*:\s*"((?:\\.|[^"\\])*)"')
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_REPARSE_POINT = 0x400
HEADERS_ZH = (
    "文件名",
    "错误类型",
    "原文（勿修改此列）",
    "译文（勿修改此列）",
    "修正（请修改此列）",
)
HEADERS_EN = (
    "File",
    "Error Type",
    "Source (Do Not Edit)",
    "Translation (Do Not Edit)",
    "Correction (Edit This Column)",
)


class AssetSuiteError(Exception):
    """共享逻辑内部异常；界面文案由调用方负责。"""

    def __init__(self, message: str, *, code: str = "", result: dict[str, Any] | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.result = result or {}


class AssetSuiteCancelled(AssetSuiteError):
    def __init__(self, message: str = "任务已取消", *, result: dict[str, Any] | None = None) -> None:
        super().__init__(message, code="cancelled")
        self.result = result or {}


def raise_if_cancelled(cancel_check: Callable[[], bool] | None, *, result: dict[str, Any] | None = None) -> None:
    if cancel_check is not None and cancel_check():
        raise AssetSuiteCancelled(result=result)


def resolve_existing_dir(raw: str | Path, *, label: str = "目录") -> Path:
    path = Path(raw).expanduser()
    if not str(raw).strip():
        raise AssetSuiteError(f"{label}不能为空", code="bad_path")
    if not path.exists():
        raise AssetSuiteError(f"{label}不存在：{path}", code="missing_path")
    if _is_link(path) and not path.is_dir():
        raise AssetSuiteError(f"{label}符号链接无效：{path}", code="escape")
    resolved = path.resolve(strict=False)
    if not resolved.is_dir():
        raise AssetSuiteError(f"{label}不是目录：{path}", code="bad_path")
    return resolved


def resolve_existing_file(raw: str | Path, *, label: str = "文件") -> Path:
    path = Path(raw).expanduser()
    if not str(raw).strip():
        raise AssetSuiteError(f"{label}不能为空", code="bad_path")
    if not path.exists():
        raise AssetSuiteError(f"{label}不存在：{path}", code="missing_path")
    if _is_link(path) and not path.is_file():
        raise AssetSuiteError(f"{label}符号链接无效：{path}", code="escape")
    resolved = path.resolve(strict=False)
    if not resolved.is_file():
        raise AssetSuiteError(f"{label}不是文件：{path}", code="bad_path")
    return resolved


def resolve_under_root(root: Path, relative: str | Path) -> Path:
    """把工作簿中的相对路径解析到 root 内；绝对路径、..、盘符/UNC、越界链接一律拒绝。"""
    text = str(relative or "").strip().replace("\\", "/")
    if not text:
        raise AssetSuiteError("目标相对路径为空", code="bad_rel")
    pure = PureWindowsPath(text)
    if pure.is_absolute() or text.startswith("/") or text.startswith("\\\\") or (len(text) >= 2 and text[1] == ":"):
        raise AssetSuiteError(f"拒绝绝对路径：{relative}", code="escape")
    parts = [part for part in Path(text).parts if part not in ("", ".")]
    if any(part == ".." for part in parts):
        raise AssetSuiteError(f"拒绝越界相对路径：{relative}", code="escape")
    candidate = (root / Path(*parts)).resolve(strict=False)
    if not _within(root, candidate):
        raise AssetSuiteError(f"路径越出译文根目录：{relative}", code="escape")
    if _is_link(candidate) and not _within(root, candidate):
        raise AssetSuiteError(f"符号链接越出译文根目录：{relative}", code="escape")
    return candidate


def sanitize_excel_text(value: Any) -> str:
    text = "" if value is None else str(value)
    return text


def write_text_cell(cell: Cell, value: Any) -> None:
    cell.value = sanitize_excel_text(value)
    cell.data_type = "s"
    cell.number_format = "@"


def batch_headers(*, english: bool = False) -> tuple[str, ...]:
    return HEADERS_EN if english else HEADERS_ZH


def collect_batch_correction_rows(
    input_dir: str | Path,
    *,
    cancel_check: Callable[[], bool] | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    folder = resolve_existing_dir(input_dir, label="输入目录")
    data_dict: dict[tuple[str, str], dict[str, Any]] = {}
    warnings: list[str] = []

    for entry in os.scandir(folder):
        raise_if_cancelled(cancel_check)
        if not entry.is_file():
            continue
        if FILE_NAME_WHITELIST.search(entry.name) is None:
            continue
        if entry.name.casefold() in {name.casefold() for name in FILE_NAME_BLACKLIST}:
            continue
        try:
            with open(entry.path, "r", encoding="utf-8-sig") as reader:
                json_data = json.load(reader)
        except Exception as exc:
            warnings.append(f"读取失败 {entry.name}: {exc}")
            continue
        if not isinstance(json_data, dict):
            warnings.append(f"跳过非字典 JSON：{entry.name}")
            continue
        for file_path, items_by_path in json_data.items():
            if not isinstance(items_by_path, dict):
                continue
            chunks = str(file_path).split("|")
            if len(chunks) == 1:
                group = FILE_NAME_WHITELIST.sub(r"\2", entry.name)
                real_path = chunks[0].strip()
            else:
                group = FILE_NAME_WHITELIST.sub(r"\2", entry.name) + " " + chunks[1].strip()
                real_path = chunks[0].strip()
            for src, dst in items_by_path.items():
                key = (real_path, str(src))
                slot = data_dict.setdefault(
                    key,
                    {"src": str(src), "dst": str(dst), "group": [], "file_path": real_path},
                )
                slot["dst"] = str(dst)
                slot["group"] = list(slot.get("group") or []) + [group]

    items = sorted(data_dict.values(), key=lambda item: (item.get("file_path", ""), str(item.get("group", ""))))
    return items, warnings


def export_batch_correction_workbook(
    input_dir: str | Path,
    output_dir: str | Path,
    *,
    english: bool = False,
    confirm_overwrite: bool = False,
    cancel_check: Callable[[], bool] | None = None,
) -> dict[str, Any]:
    items, warnings = collect_batch_correction_rows(input_dir, cancel_check=cancel_check)
    if not items:
        raise AssetSuiteError("未找到需要修正的数据", code="empty")

    out_root = resolve_existing_dir(output_dir, label="输出目录") if Path(output_dir).exists() else Path(output_dir).expanduser().resolve(strict=False)
    out_root.mkdir(parents=True, exist_ok=True)
    output_path = out_root / BATCH_WORKBOOK_NAME
    backup_path = ""
    if output_path.exists():
        if not confirm_overwrite:
            raise AssetSuiteError("输出文件已存在，需要 confirm_overwrite=true", code="need_confirm")
        ensure_write_path(output_path, out_root)
        raise_if_cancelled(cancel_check)
        backup_path = str(_backup_file(output_path))

    raise_if_cancelled(cancel_check)
    book = openpyxl.Workbook()
    sheet = book.active
    headers = batch_headers(english=english)
    for col, header in enumerate(headers, 1):
        write_text_cell(sheet.cell(row=1, column=col), header)
    sheet.auto_filter.ref = "A1:E1"
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 64
    sheet.column_dimensions["D"].width = 64
    sheet.column_dimensions["E"].width = 64
    for index, item in enumerate(items, 2):
        raise_if_cancelled(cancel_check)
        write_text_cell(sheet.cell(row=index, column=1), item.get("file_path"))
        write_text_cell(sheet.cell(row=index, column=2), "\n".join(item.get("group") or []))
        write_text_cell(sheet.cell(row=index, column=3), item.get("src"))
        write_text_cell(sheet.cell(row=index, column=4), item.get("dst"))
        write_text_cell(sheet.cell(row=index, column=5), item.get("dst"))

    ensure_write_path(output_path, out_root)
    _atomic_save_workbook(book, output_path, cancel_check=cancel_check, allowed_root=out_root)
    return {
        "success": True,
        "path": str(output_path),
        "count": len(items),
        "backup_path": backup_path,
        "warnings": warnings,
        "message": f"已生成修正数据文件（共 {len(items)} 条）",
    }


def read_batch_corrections(workbook: str | Path, *, cancel_check: Callable[[], bool] | None = None) -> dict[str, list[dict[str, str]]]:
    excel_path = resolve_existing_file(workbook, label="工作簿")
    book = openpyxl.load_workbook(excel_path, data_only=False)
    try:
        sheet = book.active
        if sheet.max_row in (0, None) or sheet.max_column in (0, None):
            raise AssetSuiteError("Excel 文件为空", code="empty")
        data_dict: dict[str, list[dict[str, str]]] = {}
        for row in range(2, (sheet.max_row or 1) + 1):
            raise_if_cancelled(cancel_check)
            file_path = sheet.cell(row=row, column=1).value
            src = sheet.cell(row=row, column=3).value
            dst = sheet.cell(row=row, column=4).value
            fix = sheet.cell(row=row, column=5).value
            if file_path is None or not str(file_path).strip():
                continue
            src_text = "" if src is None else str(src)
            dst_text = "" if dst is None else str(dst)
            fix_text = "" if fix is None else str(fix)
            if fix_text == "" or fix_text == dst_text:
                continue
            data_dict.setdefault(str(file_path), []).append(
                {"src": src_text, "dst": dst_text, "fix": fix_text}
            )
        return data_dict
    finally:
        book.close()


def apply_batch_corrections(
    workbook: str | Path,
    translation_root: str | Path,
    *,
    confirm: bool = False,
    cancel_check: Callable[[], bool] | None = None,
) -> dict[str, Any]:
    if not confirm:
        raise AssetSuiteError("注入会原地修改译文文件，需要 confirm=true", code="need_confirm")
    root = resolve_existing_dir(translation_root, label="译文根目录")
    raise_if_cancelled(cancel_check)
    data_dict = read_batch_corrections(workbook, cancel_check=cancel_check)
    if not data_dict:
        raise AssetSuiteError("没有需要修正的数据（修正列与译文列相同）", code="empty")

    total_files = len(data_dict)
    total_corrections = sum(len(items) for items in data_dict.values())
    applied_files = 0
    applied_changes = 0
    skipped_unchanged = 0
    unmatched: list[dict[str, Any]] = []
    written: list[str] = []
    backups: list[str] = []
    warnings: list[str] = []

    def cancelled_result() -> dict[str, Any]:
        return {
            "success": False, "level": "warning", "cancelled": True, "partial": bool(written),
            "total_files": total_files, "total_corrections": total_corrections,
            "applied_files": applied_files, "applied_changes": applied_changes,
            "skipped_unchanged": skipped_unchanged, "unmatched": unmatched,
            "written": list(written), "backups": list(backups), "warnings": list(warnings),
            "unwritten": [path for path in data_dict if str(root / path) not in written],
            "message": f"任务已取消，已应用 {applied_changes} 处修正，已写入 {len(written)} 个文件",
        }

    for rel_path, corrections in data_dict.items():
        raise_if_cancelled(cancel_check, result=cancelled_result())
        rel = str(rel_path).strip()
        if not rel:
            continue
        try:
            target = resolve_under_root(root, rel)
        except AssetSuiteError as exc:
            unmatched.append({"path": rel, "items": corrections, "reason": str(exc)})
            warnings.append(str(exc))
            continue
        if not target.exists():
            unmatched.append({"path": rel, "items": corrections, "reason": "文件不存在"})
            warnings.append(f"未找到目标文件: {target}")
            continue
        if _is_link(target) and not _within(root, target):
            unmatched.append({"path": rel, "items": corrections, "reason": "符号链接越界"})
            continue
        try:
            raise_if_cancelled(cancel_check)
            changed, applied, remaining, backup = apply_corrections_to_file(target, corrections, cancel_check=cancel_check, allowed_root=root)
        except AssetSuiteCancelled as exc:
            exc.result = cancelled_result()
            raise
        except Exception as exc:
            unmatched.append({"path": rel, "items": corrections, "reason": str(exc)})
            warnings.append(f"修正失败 {target}: {exc}")
            continue
        if backup:
            backups.append(backup)
        if applied > 0:
            applied_changes += applied
            applied_files += 1
            if changed:
                written.append(str(target))
            else:
                skipped_unchanged += applied
        if remaining:
            unmatched.append({"path": rel, "items": remaining, "reason": "未匹配"})

    level = "success"
    if applied_changes == 0:
        level = "error" if unmatched else "warning"
        message = "未在目标文件中找到可匹配的修正项"
        success = False
    elif unmatched:
        level = "warning"
        message = f"部分完成：已更新 {applied_files} 个文件，应用 {applied_changes} 处；{len(unmatched)} 个文件仍有未匹配项"
        success = True
    else:
        message = f"已更新 {applied_files} 个文件，共应用 {applied_changes} 处修正"
        success = True

    preview = [
        {"path": item["path"], "count": len(item["items"]), "reason": item.get("reason", "")}
        for item in unmatched[:5]
    ]
    return {
        "success": success,
        "level": level,
        "message": message,
        "total_files": total_files,
        "total_corrections": total_corrections,
        "applied_files": applied_files,
        "applied_changes": applied_changes,
        "skipped_unchanged": skipped_unchanged,
        "unmatched": unmatched,
        "unmatched_preview": preview,
        "written": written,
        "backups": backups,
        "warnings": warnings,
        "partial": bool(unmatched) and applied_changes > 0,
    }


def apply_corrections_to_file(
    file_path: Path,
    corrections: Sequence[Mapping[str, str]],
    *, cancel_check: Callable[[], bool] | None = None, allowed_root: Path | None = None,
) -> tuple[bool, int, list[dict[str, str]], str]:
    backup_before = file_path.with_suffix(file_path.suffix + ".bak")
    suffix = file_path.suffix.lower()
    if suffix == ".rpy":
        changed, applied, remaining = apply_rpy_corrections(file_path, corrections, cancel_check=cancel_check, allowed_root=allowed_root)
    else:
        changed, applied, remaining = apply_text_corrections(file_path, corrections, cancel_check=cancel_check, allowed_root=allowed_root)
    backup = ""
    if changed and backup_before.exists():
        backup = str(backup_before)
    return changed, applied, remaining, backup


def apply_rpy_corrections(
    file_path: Path,
    corrections: Sequence[Mapping[str, str]],
    *, cancel_check: Callable[[], bool] | None = None, allowed_root: Path | None = None,
) -> tuple[bool, int, list[dict[str, str]]]:
    raw, text, newline, has_bom = _read_text_bytes(file_path)
    lines = text.splitlines(keepends=True)
    state = [
        {
            "src": str(item.get("src", "") or ""),
            "dst": str(item.get("dst", "") or ""),
            "fix": str(item.get("fix", "") or ""),
            "applied": False,
        }
        for item in corrections
    ]
    changed = False
    applied = 0
    index = 0
    while index < len(lines):
        raise_if_cancelled(cancel_check)
        stripped = lines[index].strip()
        if stripped.startswith("old "):
            original_value = decode_renpy_literal(stripped[4:].strip())
            cursor = index + 1
            while cursor < len(lines):
                stripped_new = lines[cursor].strip()
                if stripped_new == "":
                    cursor += 1
                    continue
                if not stripped_new.startswith("new "):
                    break
                translation_value = decode_renpy_literal(stripped_new[4:].strip())
                for item in state:
                    if item["applied"]:
                        continue
                    if (
                        normalize_text(original_value) == normalize_text(item["src"])
                        and normalize_text(translation_value) == normalize_text(item["dst"])
                    ):
                        fixed_text = auto_convert_line_break(item["src"], item["fix"])
                        encoded = escape_tl_string(fixed_text)
                        indent = lines[cursor][: len(lines[cursor]) - len(lines[cursor].lstrip())]
                        ending = "\r\n" if lines[cursor].endswith("\r\n") else "\n" if lines[cursor].endswith("\n") else "\r" if lines[cursor].endswith("\r") else ""
                        lines[cursor] = f'{indent}new "{encoded}"{ending}'
                        item["applied"] = True
                        changed = True
                        applied += 1
                        break
                break
        index += 1

    remaining = [
        {"src": item["src"], "dst": item["dst"], "fix": item["fix"]}
        for item in state
        if not item["applied"]
    ]
    if changed:
        raise_if_cancelled(cancel_check)
        ensure_write_path(file_path, allowed_root or file_path.parent)
        _ensure_bak(file_path, allowed_root=allowed_root)
        _atomic_write_bytes(file_path, "".join(lines), newline=newline, has_bom=has_bom,
                            cancel_check=cancel_check, allowed_root=allowed_root)
    return changed, applied, remaining


def apply_text_corrections(
    file_path: Path,
    corrections: Sequence[Mapping[str, str]],
    *, cancel_check: Callable[[], bool] | None = None, allowed_root: Path | None = None,
) -> tuple[bool, int, list[dict[str, str]]]:
    raw, text, newline, has_bom = _read_text_bytes(file_path)
    content = text
    changed = False
    applied = 0
    remaining: list[dict[str, str]] = []
    for item in corrections:
        raise_if_cancelled(cancel_check)
        dst = str(item.get("dst", "") or "")
        fix = str(item.get("fix", "") or "")
        src = str(item.get("src", "") or "")
        fix_value = auto_convert_line_break(src, fix)
        if dst and dst in content:
            if fix_value != dst:
                content = content.replace(dst, fix_value, 1)
                changed = True
            applied += 1
        else:
            remaining.append({"src": src, "dst": dst, "fix": fix})
    if changed:
        raise_if_cancelled(cancel_check)
        ensure_write_path(file_path, allowed_root or file_path.parent)
        _ensure_bak(file_path, allowed_root=allowed_root)
        _atomic_write_bytes(file_path, content, newline=newline, has_bom=has_bom, preserve_text=True,
                            cancel_check=cancel_check, allowed_root=allowed_root)
    return changed, applied, remaining


def auto_convert_line_break(src: str, fix: str) -> str:
    if "_x000D_" not in src and "\r" not in src:
        return fix
    return fix.replace("_x000D_", "").replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\r\n")


def decode_renpy_literal(literal: str) -> str:
    text = (literal or "").strip()
    if text.startswith("u\"") or text.startswith("u'"):
        text = text[1:]
    try:
        return ast.literal_eval(text)
    except Exception:
        return text.strip('"').strip("'")


def normalize_text(text: str | None) -> str:
    if text is None:
        return ""
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("_x000D_", "")


def extract_character_names(
    input_dir: str | Path,
    *,
    cancel_check: Callable[[], bool] | None = None,
) -> dict[str, Any]:
    folder = resolve_existing_dir(input_dir, label="输入目录")
    names: dict[str, str] = {}
    warnings: list[str] = []
    order: list[str] = []

    for rpy_file in folder.rglob("*.rpy"):
        raise_if_cancelled(cancel_check)
        if _is_link(rpy_file):
            warnings.append(f"跳过链接文件：{rpy_file}")
            continue
        try:
            content = rpy_file.read_text(encoding="utf-8")
        except Exception as exc:
            warnings.append(f"读取失败 {rpy_file.name}: {exc}")
            continue
        lines = content.splitlines()
        for match in RE_RENPY_CHARACTER.finditer(content):
            raise_if_cancelled(cancel_check)
            display_name = match.group(2)
            index = content.count("\n", 0, match.start())
            context = "\n".join(lines[index:index + 4])
            previous = names.get(display_name)
            if previous is None:
                names[display_name] = context
                order.append(display_name)
            elif len(context) > len(previous):
                names[display_name] = context

    for json_file in folder.rglob("*.json"):
        raise_if_cancelled(cancel_check)
        if _is_link(json_file):
            warnings.append(f"跳过链接文件：{json_file}")
            continue
        try:
            raw_text = json_file.read_text(encoding="utf-8")
        except Exception as exc:
            warnings.append(f"读取失败 {json_file.name}: {exc}")
            continue
        found: list[str] = []
        try:
            payload = json.loads(raw_text)
            found = _collect_json_names(payload)
        except Exception:
            warnings.append(f"JSON 解析失败，回退正则：{json_file.name}")
            for match in RE_JSON_NAME_FIELD.findall(raw_text):
                try:
                    found.append(json.loads(f'"{match}"'))
                except Exception:
                    found.append(match.encode("utf-8").decode("unicode_escape"))
        for name in found:
            if name and name not in names:
                names[name] = f"[从 {json_file.name} 提取]"
                order.append(name)

    entries = [{"src": name, "context": names[name]} for name in order if name in names]
    # 若 order 与 dict 因去重不一致，补齐剩余
    for name, context in names.items():
        if name not in {item["src"] for item in entries}:
            entries.append({"src": name, "context": context})

    return {
        "success": True,
        "count": len(entries),
        "entries": entries,
        "warnings": warnings,
        "message": f"找到 {len(entries)} 个角色姓名" if entries else "未找到任何角色姓名定义",
        "empty": len(entries) == 0,
    }


def export_name_glossary(
    entries: Sequence[Mapping[str, Any]],
    *,
    format: str = "txt",
    output_file: str | Path | None = None,
    confirm_overwrite: bool = False,
    dst_overrides: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    if not entries:
        raise AssetSuiteError("没有可导出的姓名条目", code="empty")
    fmt = (format or "txt").strip().lower()
    if fmt not in {"txt", "json"}:
        raise AssetSuiteError("导出格式仅支持 txt 或 json", code="bad_format")

    rows: list[dict[str, str]] = []
    for item in entries:
        src = str(item.get("src", "") or "").strip()
        if not src:
            continue
        dst = str((dst_overrides or {}).get(src) or item.get("dst") or src)
        comment = str(item.get("comment") or item.get("info") or "角色姓名")
        rows.append({"src": src, "dst": dst, "info": comment, "comment": comment, "type": str(item.get("type") or "角色")})

    if fmt == "txt":
        content = "\n".join(f"{row['src']} -> {row['dst']} #{row['comment']}" for row in rows)
        media_type = "text/plain; charset=utf-8"
        filename = "glossary_names.txt"
    else:
        content = json.dumps(rows, ensure_ascii=False, indent=2)
        media_type = "application/json; charset=utf-8"
        filename = "glossary_names.json"

    path_text = ""
    backup_path = ""
    if output_file is not None:
        target = Path(output_file).expanduser()
        scope = target.parent.resolve(strict=False)
        ensure_write_path(target, scope)
        if target.exists():
            if not confirm_overwrite:
                raise AssetSuiteError("输出文件已存在，需要 confirm_overwrite=true", code="need_confirm")
            backup_path = str(_backup_file(target))
        target.parent.mkdir(parents=True, exist_ok=True)
        atomic_write_text(target, content if content.endswith("\n") else content + "\n", encoding="utf-8", newline="\n", allowed_roots=[scope])
        path_text = str(target.resolve(strict=False))
        filename = target.name

    return {
        "success": True,
        "format": fmt,
        "filename": filename,
        "media_type": media_type,
        "content": content,
        "path": path_text,
        "backup_path": backup_path,
        "count": len(rows),
        "message": f"已生成术语表（共 {len(rows)} 个条目）",
    }


def backup_directory_safe(src_folder: Path, *, project_root: Path | None = None, cancel_check: Callable[[], bool] | None = None) -> Path:
    """备份目录；若源位于 game/ 内，备份放到游戏根外侧，避免生成可执行 .rpy 副本。"""
    src = src_folder.resolve(strict=False)
    parent = src.parent
    game = next((ancestor for ancestor in (src, *src.parents) if ancestor.name.lower() == "game"), None)
    if game is not None:
        parent = game.parent / "_renpybox_backups"
        parent.mkdir(parents=True, exist_ok=True)
        ensure_write_path(parent, game.parent)
    timestamp = int(time.time())
    backup_path = parent / f"{src.name}_backup_{timestamp}"
    counter = 1
    while backup_path.exists():
        backup_path = parent / f"{src.name}_backup_{timestamp}_{counter}"
        counter += 1
    _copytree_checked(src, backup_path, cancel_check=cancel_check)
    return backup_path


def backup_translate_output_if_exists(base_out: Path, *, cancel_check: Callable[[], bool] | None = None) -> str:
    if not base_out.exists():
        return ""
    stamp = time.strftime("%Y%m%d_%H%M%S")
    dest = base_out.parent / f"translate_output_backup_{stamp}"
    counter = 1
    while dest.exists():
        dest = base_out.parent / f"translate_output_backup_{stamp}_{counter}"
        counter += 1
    _copytree_checked(base_out, dest, cancel_check=cancel_check)
    return str(dest)


def _collect_json_names(payload: Any) -> list[str]:
    found: list[str] = []

    def walk(node: Any) -> None:
        if isinstance(node, dict):
            for key, value in node.items():
                if key == "name" and isinstance(value, str) and value:
                    found.append(value)
                else:
                    walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(payload)
    return found


def _read_text_bytes(path: Path) -> tuple[bytes, str, str, bool]:
    raw = path.read_bytes()
    has_bom = raw.startswith(b"\xef\xbb\xbf")
    body = raw[3:] if has_bom else raw
    newline = "\r\n" if b"\r\n" in body else "\n"
    text = body.decode("utf-8")
    return raw, text, newline, has_bom


def _atomic_write_bytes(
    path: Path,
    text: str,
    *,
    newline: str,
    has_bom: bool,
    preserve_text: bool = False,
    cancel_check: Callable[[], bool] | None = None, allowed_root: Path | None = None,
) -> None:
    # 按原字符串字节写回，保留混合换行且避免 CRCRLF。
    data = text.encode("utf-8")
    if has_bom:
        data = b"\xef\xbb\xbf" + data
    ensure_write_path(path, allowed_root or path.parent)
    raise_if_cancelled(cancel_check)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent))
    temp_path = Path(temp_name)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        raise_if_cancelled(cancel_check)
        ensure_write_path(path, allowed_root or path.parent)
        os.replace(str(temp_path), str(path))
        temp_path = None
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _ensure_bak(path: Path, *, allowed_root: Path | None = None) -> str:
    backup = path.with_suffix(path.suffix + ".bak")
    ensure_write_path(backup, allowed_root or path.parent)
    if backup.exists() and not backup.is_file():
        raise AssetSuiteError(f"已有备份不是文件：{backup}", code="bad_backup")
    if not backup.exists():
        shutil.copy2(path, backup)
    return str(backup)


def _backup_file(path: Path) -> Path:
    stamp = int(time.time())
    backup = path.with_name(f"{path.name}.bak_{stamp}")
    counter = 1
    while backup.exists() or backup.is_symlink():
        backup = path.with_name(f"{path.name}.bak_{stamp}_{counter}")
        counter += 1
    ensure_write_path(backup, path.parent)
    shutil.copy2(path, backup)
    return backup


def _atomic_save_workbook(book: openpyxl.Workbook, output_path: Path, *, cancel_check: Callable[[], bool] | None = None, allowed_root: Path | None = None) -> None:
    ensure_write_path(output_path, allowed_root or output_path.parent)
    raise_if_cancelled(cancel_check)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{output_path.name}.", suffix=".tmp", dir=str(output_path.parent))
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        book.save(temp_path)
        raise_if_cancelled(cancel_check)
        ensure_write_path(output_path, allowed_root or output_path.parent)
        os.replace(str(temp_path), str(output_path))
        temp_path = None
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _within(root: Path, path: Path) -> bool:
    try:
        root_real = Path(os.path.normcase(os.path.abspath(root.resolve(strict=False))))
        real = Path(os.path.normcase(os.path.abspath(path.resolve(strict=False))))
    except OSError:
        return False
    return real == root_real or real.is_relative_to(root_real)


def _is_link(path: Path) -> bool:
    try:
        info = path.lstat()
    except OSError:
        return False
    if stat.S_ISLNK(info.st_mode):
        return True
    attributes = getattr(info, "st_file_attributes", 0) or 0
    return bool(attributes & _REPARSE_POINT)


__all__ = [
    "AssetSuiteCancelled",
    "AssetSuiteError",
    "BATCH_WORKBOOK_NAME",
    "FILE_NAME_BLACKLIST",
    "FILE_NAME_WHITELIST",
    "RE_RENPY_CHARACTER",
    "apply_batch_corrections",
    "apply_corrections_to_file",
    "apply_rpy_corrections",
    "apply_text_corrections",
    "auto_convert_line_break",
    "backup_directory_safe",
    "backup_translate_output_if_exists",
    "batch_headers",
    "collect_batch_correction_rows",
    "decode_renpy_literal",
    "export_batch_correction_workbook",
    "export_name_glossary",
    "extract_character_names",
    "normalize_text",
    "raise_if_cancelled",
    "read_batch_corrections",
    "resolve_existing_dir",
    "resolve_existing_file",
    "resolve_under_root",
    "sanitize_excel_text",
    "write_text_cell",
]


def ensure_write_path(path: Path, root: Path) -> Path:
    """写入前检查真实目录边界，阻止预存符号链接和 junction 逃逸。"""
    if not _within(root, path):
        raise AssetSuiteError(f"写入路径越出已选目录：{path}", code="escape")
    return path


def _copytree_checked(source: Path, destination: Path, *, cancel_check=None) -> None:
    for current, dirs, files in os.walk(source, followlinks=False):
        raise_if_cancelled(cancel_check)
        for name in dirs + files:
            ensure_write_path(Path(current) / name, source)
    def copy_file(src, dst):
        raise_if_cancelled(cancel_check)
        ensure_write_path(Path(src), source)
        ensure_write_path(Path(dst), destination)
        return shutil.copy2(src, dst)
    raise_if_cancelled(cancel_check)
    try:
        shutil.copytree(source, destination, copy_function=copy_file)
    except (shutil.Error, AssetSuiteCancelled) as exc:
        if isinstance(exc, AssetSuiteCancelled) or (cancel_check and cancel_check()):
            raise AssetSuiteCancelled(result={"backup_path": str(destination), "backup_complete": False,
                                               "written": [], "partial": False}) from exc
        raise
