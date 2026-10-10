"""词库 / 禁翻纯业务：与 LocalGlossaryPage、TextPreservePage 对齐，无 Qt 依赖。"""

from __future__ import annotations

import base64
import dataclasses
import io
import re
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell

from base.PathHelper import get_resource_path
from module.Config import Config
from module.RuleStatistics import (
    count_glossary_hit_counts,
    count_text_preserve_hit_counts,
    get_statistics_cache_dir,
    load_counted_source_texts,
)
from module.Text.SkipRules import should_skip_text

# Excel 导入/导出上限（字节）
EXCEL_MAX_BYTES = 5 * 1024 * 1024

# 过滤器关键字（参考 AiNiee NER 过滤规则），命中则跳过
FILTER_KEYWORDS = (
    "-",
    "…",
    "一",
    "―",
    "？",
    "©",
    "章　",
    "ー",
    "http",
    "！",
    "=",
    '"',
    "＋",
    "：",
    "『",
    "ぃ",
    "～",
    "♦",
    "〇",
    "└",
    "'",
    "/",
    "｢",
    "）",
    "（",
    "♥",
    "●",
    "!",
    "】",
    "【",
    "<",
    ">",
    "*",
    "〜",
    "EV",
    "♪",
    "^",
    "★",
    "※",
    ".",
    "|",
    "ｰ",
    "%",
    "if",
    "Lv",
    "(",
    "\\",
    "]",
    "[",
    "◆",
    ":",
    "_",
    "ｗｗｗ",
    "、",
    "ぁぁ",
    "んえ",
    "んんん",
)

RE_VARIABLE_IN_TEXT = re.compile(r"\[([\w.]+)\]")
RE_CHARACTER_CALL = re.compile(
    r"Character\s*\(\s*(?:_\(\s*)?(['\"])((?:\\\1|.)*?)\1",
    re.MULTILINE,
)
_FORMULA_PREFIXES = ("=", "+", "-", "@")


class LexiconOpsError(ValueError):
    """词库业务可预期错误。"""


# ---------------------------------------------------------------------------
# 术语库：规范化 / 合并 / Excel
# ---------------------------------------------------------------------------


def normalize_glossary_src(text: str) -> str:
    if not text:
        return ""
    normalized = re.sub(r"\s+", " ", text)
    normalized = normalized.strip().strip("\"'“”‘’")
    return normalized.lower()


def merge_glossary_entries(base: Mapping[str, Any], incoming: Mapping[str, Any]) -> dict[str, Any]:
    def _clean(value: Any) -> str:
        return value.strip() if isinstance(value, str) else ""

    merged = {
        "record_id": str(base.get("record_id", "") or ""),
        "src": _clean(base.get("src")),
        "dst": _clean(base.get("dst")),
        "type": _clean(base.get("type")),
        "comment": _clean(base.get("comment")),
        "case_sensitive": bool(base.get("case_sensitive", False)),
        "candidate": bool(base.get("candidate", False)),
        "candidate_confirmed": bool(base.get("candidate_confirmed", False)),
        "enabled": bool(base.get("enabled", True)),
        "regex": bool(base.get("regex", False)),
    }
    incoming_cleaned = {
        "record_id": str(incoming.get("record_id", "") or ""),
        "src": _clean(incoming.get("src")),
        "dst": _clean(incoming.get("dst")),
        "type": _clean(incoming.get("type")),
        "comment": _clean(incoming.get("comment")),
        "case_sensitive": bool(incoming.get("case_sensitive", False)),
        "candidate": bool(incoming.get("candidate", False)),
        "candidate_confirmed": bool(incoming.get("candidate_confirmed", False)),
        "enabled": bool(incoming.get("enabled", True)),
        "regex": bool(incoming.get("regex", False)),
    }

    if not merged["record_id"] and incoming_cleaned["record_id"]:
        merged["record_id"] = incoming_cleaned["record_id"]
        merged["candidate"] = incoming_cleaned["candidate"]
        merged["candidate_confirmed"] = incoming_cleaned["candidate_confirmed"]

    if incoming_cleaned["dst"]:
        if not merged["dst"] or (merged["src"] and merged["dst"].lower() == merged["src"].lower()):
            merged["dst"] = incoming_cleaned["dst"]

    if incoming_cleaned["type"] and not merged["type"]:
        merged["type"] = incoming_cleaned["type"]

    if incoming_cleaned["comment"]:
        if not merged["comment"]:
            merged["comment"] = incoming_cleaned["comment"]
        elif (
            incoming_cleaned["comment"] not in merged["comment"]
            and len(incoming_cleaned["comment"]) > len(merged["comment"])
        ):
            merged["comment"] = incoming_cleaned["comment"]

    if incoming_cleaned["src"] and not merged["src"]:
        merged["src"] = incoming_cleaned["src"]

    if incoming_cleaned["case_sensitive"]:
        merged["case_sensitive"] = True
    if incoming_cleaned["regex"]:
        merged["regex"] = True
    if incoming_cleaned["candidate_confirmed"]:
        merged["candidate_confirmed"] = True

    return merged


def build_glossary_header_map(headers: Sequence[str]) -> dict[str, int]:
    alias = {
        "src": {"原文", "原始文本", "source", "src"},
        "dst": {"译文", "翻译", "target", "translation", "dst"},
        "type": {"类别", "分类", "type", "category"},
        "comment": {"备注", "说明", "comment", "note", "备注信息"},
    }
    mapping: dict[str, int] = {}
    for index, name in enumerate(headers):
        lower_name = str(name or "").lower()
        for key, options in alias.items():
            if lower_name in {opt.lower() for opt in options} and key not in mapping:
                mapping[key] = index
    return mapping


def safe_excel_cell(row: Sequence[Any], index: int | None) -> str:
    if index is None:
        return ""
    if index >= len(row):
        return ""
    value = row[index]
    return "" if value is None else str(value).strip()


def map_language_to_fasttranslator_code(
    lang: str,
    *,
    is_target: bool,
    traditional_chinese_enable: bool,
) -> str:
    key = str(lang or "").strip() or "auto"
    upper = key.upper()

    if upper in {"AUTO", "自动", "NONE"}:
        return "auto"

    if upper in {"ZH", "CHINESE", "CN", "ZH-CN", "ZH_CN", "ZH-HANS", "ZH_HANS"}:
        return "zh-TW" if is_target and traditional_chinese_enable else "zh-CN"
    if upper in {"EN", "ENGLISH"}:
        return "en"
    if upper in {"JA", "JP", "JAPANESE"}:
        return "ja"
    if upper in {"KO", "KR", "KOREAN"}:
        return "ko"
    if upper in {"RU", "RUSSIAN"}:
        return "ru"
    if upper in {"AR", "ARABIC"}:
        return "ar"
    if upper in {"DE", "GERMAN"}:
        return "de"
    if upper in {"FR", "FRENCH"}:
        return "fr"
    if upper in {"PL", "POLISH"}:
        return "pl"
    if upper in {"ES", "SPANISH"}:
        return "es"
    if upper in {"IT", "ITALIAN"}:
        return "it"
    if upper in {"PT", "PORTUGUESE"}:
        return "pt"
    if upper in {"HU", "HUNGARIAN"}:
        return "hu"
    if upper in {"TR", "TURKISH"}:
        return "tr"
    if upper in {"TH", "THAI"}:
        return "th"
    if upper in {"ID", "INDONESIAN"}:
        return "id"
    if upper in {"VI", "VIETNAMESE"}:
        return "vi"

    return key


def collect_glossary_translate_tasks(rows: Sequence[Mapping[str, Any]]) -> list[tuple[int, str]]:
    """从行列表收集待译任务：译文为空或等于原文。"""
    tasks: list[tuple[int, str]] = []
    for index, row in enumerate(rows):
        if not isinstance(row, Mapping):
            continue
        src = str(row.get("src", "") or "").strip()
        if not src:
            continue
        dst = str(row.get("dst", "") or "").strip()
        if dst and dst != src:
            continue
        tasks.append((index, src))
    return tasks


def prepare_candidate_entry(entry: Mapping[str, Any]) -> dict[str, Any]:
    type_raw = str(entry.get("type", "") or "").strip()
    return {
        "record_id": str(entry.get("record_id", "") or ""),
        "src": str(entry.get("source", entry.get("src", "")) or "").strip(),
        "dst": str(entry.get("target", entry.get("dst", "")) or "").strip(),
        "type": f"候选 / {type_raw}" if type_raw else "候选",
        "comment": str(
            entry.get("note", entry.get("comment", "")) or entry.get("info", "") or "术语候选 (自动提取)"
        ).strip(),
        "case_sensitive": bool(entry.get("case_sensitive", False)),
        "candidate": True,
        "candidate_confirmed": False,
        "enabled": bool(entry.get("enabled", True)),
        "regex": bool(entry.get("regex", False)),
    }


def merge_candidate_entries(
    current_entries: Sequence[Mapping[str, Any]],
    entries: Sequence[Mapping[str, Any]],
) -> tuple[list[dict[str, Any]], int, int, dict[str, int]]:
    """合并候选到现有行；返回 (merged, added, updated, count_map)。"""
    merged_entries: list[dict[str, Any]] = []
    key_index: dict[tuple[str, bool], int] = {}

    for item in current_entries:
        if not isinstance(item, Mapping):
            continue
        copied = {
            "record_id": str(item.get("record_id", "") or ""),
            "src": str(item.get("src", "") or "").strip(),
            "dst": str(item.get("dst", "") or "").strip(),
            "type": str(item.get("type", "") or "").strip(),
            "comment": str(item.get("comment", "") or "").strip(),
            "case_sensitive": bool(item.get("case_sensitive", False)),
            "candidate": bool(item.get("candidate", False)),
            "candidate_confirmed": bool(item.get("candidate_confirmed", False)),
            "enabled": bool(item.get("enabled", True)),
            "regex": bool(item.get("regex", False)),
        }
        source_key = normalize_glossary_src(copied.get("src", ""))
        key = (source_key, copied["regex"])
        if not source_key:
            continue
        if key not in key_index:
            merged_entries.append(copied)
            key_index[key] = len(merged_entries) - 1
            continue
        merged_entries[key_index[key]] = merge_glossary_entries(merged_entries[key_index[key]], copied)

    added_count = 0
    updated_count = 0
    count_map: dict[str, int] = {}

    for entry in entries:
        if not isinstance(entry, Mapping):
            continue
        prepared = prepare_candidate_entry(entry)
        source_key = normalize_glossary_src(prepared.get("src", ""))
        key = (source_key, prepared["regex"])
        if not source_key:
            continue
        count_map[source_key] = max(count_map.get(source_key, 0), int(entry.get("count", 0) or 0))
        if key not in key_index:
            merged_entries.append(prepared)
            key_index[key] = len(merged_entries) - 1
            added_count += 1
            continue
        current = merged_entries[key_index[key]]
        merged = merge_glossary_entries(current, prepared)
        if merged != current:
            updated_count += 1
        merged_entries[key_index[key]] = merged

    return merged_entries, added_count, updated_count, count_map


def dedupe_glossary_entries(entries: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    key_index: dict[str, int] = {}
    deduped: list[dict[str, Any]] = []
    for item in entries:
        if not isinstance(item, Mapping):
            continue
        key = normalize_glossary_src(str(item.get("src", "") or ""))
        if not key:
            continue
        if key not in key_index:
            deduped.append(dict(item))
            key_index[key] = len(deduped) - 1
            continue
        deduped[key_index[key]] = merge_glossary_entries(deduped[key_index[key]], item)
    return deduped


def clean_text_for_classify(text: str) -> str:
    if not text:
        return ""
    cleaned = re.sub(r"\{/?[^}]+\}", "", text)
    cleaned = cleaned.replace("\u3000", " ").strip()
    return cleaned


def categorize_term(text: str, default: str = "") -> str:
    if not text:
        return default
    t = text.strip()
    lower = t.lower()
    place_keywords = [
        "city",
        "village",
        "town",
        "forest",
        "mountain",
        "hill",
        "park",
        "garden",
        "school",
        "academy",
        "college",
        "campus",
        "church",
        "temple",
        "shrine",
        "castle",
        "tower",
        "dungeon",
        "cave",
        "ruins",
        "harbor",
        "port",
        "station",
        "beach",
        "island",
        "lake",
        "river",
        "bridge",
        "street",
        "road",
        "avenue",
        "hotel",
        "inn",
        "bar",
        "cafe",
        "shop",
        "market",
        "library",
    ]
    item_keywords = [
        "sword",
        "blade",
        "dagger",
        "bow",
        "gun",
        "rifle",
        "pistol",
        "armor",
        "shield",
        "ring",
        "necklace",
        "amulet",
        "bracelet",
        "crown",
        "helmet",
        "boots",
        "gloves",
        "potion",
        "elixir",
        "herb",
        "scroll",
        "book",
        "map",
        "key",
        "card",
        "ticket",
        "coin",
        "gem",
        "crystal",
        "stone",
        "orb",
        "staff",
        "wand",
        "medal",
    ]
    if any(k in lower for k in place_keywords):
        return "地名"
    if any(k in lower for k in item_keywords):
        return "物品"
    words = t.split()
    if words and all(w[:1].isupper() for w in words if w):
        return default or ""
    return default


def is_probable_name(text: str) -> bool:
    if not text:
        return False
    if len(text) > 40 or "\n" in text:
        return False
    t = clean_text_for_classify(text)
    if not t:
        return False
    if any(p in t for p in (".", "?", "!", ":", "；", "。", "！", "？")):
        return False
    if any(ch.isdigit() for ch in t):
        return False
    words = t.split()
    if not words or len(words) > 4:
        return False
    allow_connectors = {"of", "the", "and"}
    cjk = any("\u3040" <= ch <= "\u30ff" or "\u4e00" <= ch <= "\u9fff" for ch in t)
    if cjk:
        return True
    for w in words:
        if w.lower() in allow_connectors:
            continue
        if not (w[:1].isupper() or w.isupper()):
            return False
    return True


def auto_categorize_entries(rows: Sequence[Mapping[str, Any]]) -> tuple[int, list[dict[str, Any]]]:
    """关键词填充空白类别；返回 (changed, new_rows)。"""
    new_rows: list[dict[str, Any]] = []
    changed = 0
    for row in rows:
        item = dict(row) if isinstance(row, Mapping) else {}
        type_text = str(item.get("type", "") or "").strip()
        if type_text:
            new_rows.append(item)
            continue
        cleaned = clean_text_for_classify(str(item.get("src", "") or ""))
        if not cleaned:
            new_rows.append(item)
            continue
        guess = categorize_term(cleaned)
        if guess:
            item["type"] = guess
            changed += 1
        new_rows.append(item)
    return changed, new_rows


def guess_ner_preference_from_texts(texts: Sequence[str]) -> str:
    cjk = 0
    latin = 0
    for text in list(texts)[:200]:
        for ch in str(text or ""):
            if "\u4e00" <= ch <= "\u9fff" or "\u3040" <= ch <= "\u30ff":
                cjk += 1
            elif ch.isalpha():
                latin += 1
    if cjk > latin:
        return "ja"
    return "en"


def find_ner_model_path(texts: Sequence[str] | None = None) -> Path | None:
    """查找本地 NER 模型路径（Resource/Models/ner），按语言偏好选择。"""
    candidates: list[Path] = []
    for model_root in [Path(get_resource_path("resource", "Models", "ner")), Path(get_resource_path("resource", "models", "ner"))]:
        if model_root.exists():
            for path in model_root.iterdir():
                if path.is_dir() and (path / "meta.json").exists():
                    candidates.append(path)
    if not candidates:
        return None

    preferred = guess_ner_preference_from_texts(texts or [])

    def _score(path: Path) -> int:
        name = path.name.lower()
        if preferred == "ja":
            if name.startswith("ja_core"):
                return 0
        if preferred == "en":
            if name.startswith("en_core_web_md"):
                return 0
            if name.startswith("en_core_web_"):
                return 1
        return 5

    candidates.sort(key=_score)
    return candidates[0]


def ner_categorize_entries(rows: Sequence[Mapping[str, Any]]) -> tuple[int, list[dict[str, Any]]]:
    """使用本地 spaCy NER 填充空白类别；失败静默返回 0。"""
    new_rows = [dict(row) if isinstance(row, Mapping) else {} for row in rows]
    try:
        import spacy
    except Exception:
        return 0, new_rows

    src_texts = [str(item.get("src", "") or "") for item in new_rows]
    model_path = find_ner_model_path(src_texts)
    if not model_path:
        return 0, new_rows

    try:
        nlp = spacy.load(
            str(model_path),
            exclude=["parser", "tagger", "lemmatizer", "attribute_ruler", "tok2vec"],
        )
    except Exception:
        return 0, new_rows

    label_map = {
        "PER": "角色",
        "PERSON": "角色",
        "PER_NO": "角色",
        "LOC": "地名",
        "GPE": "地名",
        "ORG": "组织",
        "FAC": "地名",
        "PRODUCT": "物品",
        "ITEM": "物品",
    }

    changed = 0
    for item in new_rows:
        if str(item.get("type", "") or "").strip():
            continue
        text = clean_text_for_classify(str(item.get("src", "") or ""))
        if not text:
            continue
        if any(k in text for k in FILTER_KEYWORDS):
            continue
        try:
            doc = nlp(text)
        except Exception:
            continue
        guessed = ""
        for ent in doc.ents:
            if any(k in ent.text for k in FILTER_KEYWORDS):
                continue
            if ent.text.strip().lower() == text.strip().lower():
                guessed = label_map.get(ent.label_, ent.label_)
                break
        if not guessed and doc.ents:
            for ent in doc.ents:
                if any(k in ent.text for k in FILTER_KEYWORDS):
                    continue
                guessed = label_map.get(ent.label_, ent.label_)
                if guessed:
                    break
        if guessed:
            item["type"] = guessed
            changed += 1
    return changed, new_rows


def auto_classify_entries(rows: Sequence[Mapping[str, Any]]) -> tuple[int, int, list[dict[str, Any]]]:
    """先 NER 再关键词；返回 (ner_count, kw_count, rows)。"""
    ner_count, after_ner = ner_categorize_entries(rows)
    kw_count, after_kw = auto_categorize_entries(after_ner)
    return ner_count, kw_count, after_kw


# ---------------------------------------------------------------------------
# Ren'Py 路径与角色扫描
# ---------------------------------------------------------------------------


def _normalize_existing_path(path_text: str) -> Path | None:
    raw = str(path_text or "").strip()
    if raw == "":
        return None
    try:
        path = Path(raw).expanduser().resolve()
    except Exception:
        path = Path(raw)
    return path if path.exists() else None


def list_renpy_scan_candidates(config: Any) -> list[Path]:
    """根据配置推断 Ren'Py 扫描候选目录。"""
    raws = [
        getattr(config, "input_folder", ""),
        getattr(config, "output_folder", ""),
        getattr(config, "renpy_tl_folder", ""),
        getattr(config, "renpy_game_folder", ""),
        getattr(config, "renpy_project_path", ""),
    ]
    candidates: list[Path] = []
    for raw in raws:
        path = _normalize_existing_path(raw)
        if path is None:
            continue

        if path.is_file():
            candidates.append(path.parent)
            continue

        if path.parent.name.lower() == "tl":
            if path.parent.parent.name.lower() == "game":
                candidates.append(path.parent.parent)
                candidates.append(path.parent.parent.parent)
            else:
                candidates.append(path.parent.parent)
            continue

        if path.name.lower() == "tl":
            if path.parent.name.lower() == "game":
                candidates.append(path.parent)
                candidates.append(path.parent.parent)
            else:
                candidates.append(path.parent)
            continue

        game_child = path / "game"
        if game_child.is_dir():
            candidates.append(game_child)
            candidates.append(path)
            continue

        candidates.append(path)

    deduped: list[Path] = []
    seen: set[str] = set()
    for candidate in candidates:
        try:
            key = str(candidate.resolve()).lower()
        except Exception:
            key = str(candidate).lower()
        if key in seen:
            continue
        seen.add(key)
        deduped.append(candidate)
    return deduped


def count_source_rpy_files(root: Path) -> int:
    count = 0
    try:
        for rpy_file in root.rglob("*.rpy"):
            if "tl" in [part.lower() for part in rpy_file.parts]:
                continue
            count += 1
    except OSError as exc:
        raise LexiconOpsError(f"源码目录读取失败：{root}") from exc
    return count


def is_tl_dir_for_source_root(source_root: Path, tl_root: Path) -> bool:
    try:
        source = source_root.resolve()
        tl_dir = tl_root.resolve()
    except Exception:
        source = source_root
        tl_dir = tl_root

    if source.name.lower() == "game":
        allowed_prefixes = (
            source / "tl",
            source.parent / "tl",
        )
    else:
        allowed_prefixes = (
            source / "game" / "tl",
            source / "tl",
        )

    for prefix in allowed_prefixes:
        try:
            prefix_resolved = prefix.resolve()
        except Exception:
            prefix_resolved = prefix
        if tl_dir.parent == prefix_resolved:
            return True
    return False


def resolve_renpy_scan_paths(config: Any) -> tuple[Path | None, Path | None]:
    """解析源码扫描目录与 tl 语言目录。"""
    candidates = list_renpy_scan_candidates(config)
    if not candidates:
        return None, None

    source_root: Path | None = None
    fallback_root: Path | None = None
    for candidate in candidates:
        if fallback_root is None:
            fallback_root = candidate
        count = count_source_rpy_files(candidate)
        if count > 0:
            source_root = candidate
            break

    if source_root is None:
        source_root = fallback_root
    if source_root is None:
        return None, None

    tl_root: Path | None = None
    configured_tl = _normalize_existing_path(getattr(config, "renpy_tl_folder", ""))
    if (
        configured_tl is not None
        and configured_tl.is_dir()
        and is_tl_dir_for_source_root(source_root, configured_tl)
    ):
        tl_root = configured_tl
    else:
        possible_roots = [source_root]
        if source_root.name.lower() == "game" and source_root.parent.exists():
            possible_roots.append(source_root.parent)
        for base in possible_roots:
            for tl_dir in (base / "game" / "tl", base / "tl"):
                if not tl_dir.exists() or not tl_dir.is_dir():
                    continue
                for preferred in ("chinese", "schinese", "tchinese", "english", "japanese", "korean"):
                    candidate = tl_dir / preferred
                    if candidate.is_dir():
                        tl_root = candidate
                        break
                if tl_root is None:
                    children = sorted(
                        [child for child in tl_dir.iterdir() if child.is_dir()],
                        key=lambda item: item.name.lower(),
                    )
                    if children:
                        tl_root = children[0]
                if tl_root is not None:
                    break
            if tl_root is not None:
                break

    return source_root, tl_root


def effective_tl_name(config: Any, tl_root: Path | None = None) -> str:
    if tl_root is not None and tl_root.name:
        return tl_root.name
    raw_value = str(getattr(config, "renpy_tl_folder", "") or "").strip()
    return Path(raw_value).name if raw_value else "chinese"


def extract_names_from_source(game_path: Path, cancel_callback=None) -> set[str]:
    names: set[str] = set()
    try:
        for rpy_file in game_path.rglob("*.rpy"):
            if cancel_callback:
                cancel_callback()
            rel_parts = rpy_file.relative_to(game_path).parts
            if any(p.casefold() in ("tl", "cache", "__pycache__") for p in rel_parts):
                continue
            try:
                content = rpy_file.read_text(encoding="utf-8", errors="ignore")
                for match in RE_CHARACTER_CALL.finditer(content):
                    raw_name = match.group(2)
                    name = raw_name.replace('\\"', '"').replace("\\'", "'").replace("\\n", " ").strip()
                    if not name:
                        continue
                    if name.startswith("[") and name.endswith("]"):
                        continue
                    if len(name) > 50:
                        continue
                    if name.isdigit() or all(c in "!@#$%^&*()_+-=[]{}|;:'\",.<>?/\\" for c in name):
                        continue
                    cleaned = clean_text_for_classify(name)
                    if cleaned and len(cleaned) >= 2:
                        if is_probable_name(cleaned):
                            names.add(cleaned)
                        elif len(cleaned) <= 20:
                            names.add(cleaned)
            except OSError as exc:
                raise LexiconOpsError(f"角色源码读取失败：{rpy_file}") from exc
    except OSError as exc:
        raise LexiconOpsError("角色源码扫描失败") from exc
    return names


def extract_names_from_miss_files(game_path: Path, tl_name: str, cancel_callback=None) -> set[str]:
    names: set[str] = set()
    tl_root = game_path / "tl" / tl_name
    if not tl_root.exists():
        alt_tl_root = game_path / "game" / "tl" / tl_name
        if alt_tl_root.exists():
            tl_root = alt_tl_root
    if not tl_root.exists():
        return names
    candidates: list[Path] = []
    for base in (tl_root, tl_root / "miss"):
        candidates.extend(base.glob("miss_ready_replace*.rpy"))
        candidates.extend(base.glob("miss_ready_replace*.txt"))
    if not candidates:
        return names
    for miss_file in candidates:
        if cancel_callback:
            cancel_callback()
        try:
            for line in miss_file.read_text(encoding="utf-8", errors="ignore").splitlines():
                line = line.strip()
                if line.startswith("old "):
                    match = re.search(r'old\s+"(.*)"', line)
                    if not match:
                        continue
                    raw = match.group(1)
                    text = raw.replace('\\"', '"').replace("\\n", "\n").replace("\\\\", "\\")
                    clean = clean_text_for_classify(text)
                    if not clean:
                        continue
                    if is_probable_name(clean):
                        names.add(clean)
        except OSError as exc:
            raise LexiconOpsError(f"角色补漏文件读取失败：{miss_file}") from exc
    return names


def _is_auto_character_entry(item: Mapping[str, Any]) -> bool:
    info = str(item.get("info", "") or item.get("comment", "") or item.get("note", "") or "")
    return "自动提取" in info and ("character" in info.lower() or "角色" in info)


def build_character_scan_entries(
    manual_entries: Sequence[Mapping[str, Any]],
    found_names: Iterable[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    """过滤自动角色候选后合并扫描结果；返回 (kept_manual, new_entries, removed_auto)。"""
    kept: list[dict[str, Any]] = []
    removed_auto = 0
    for item in manual_entries:
        if not isinstance(item, Mapping):
            continue
        if _is_auto_character_entry(item):
            removed_auto += 1
            continue
        kept.append(dict(item))

    existing_src = {str(item.get("src", "") or "") for item in kept}
    new_entries: list[dict[str, Any]] = []
    for name in found_names:
        if name in existing_src:
            continue
        cleaned = clean_text_for_classify(name)
        if not cleaned or should_skip_text(cleaned):
            continue
        type_guess = categorize_term(cleaned, default="")
        new_entries.append(
            {
                "src": cleaned,
                "dst": "",
                "info": "角色名 (自动提取)",
                "comment": "角色名 (自动提取)",
                "type": f"候选 / {type_guess}" if type_guess else "候选",
                "candidate": True,
                "candidate_confirmed": False,
            }
        )
    return kept, new_entries, removed_auto


def apply_character_scan_to_rows(
    rows: Sequence[Mapping[str, Any]],
    found_names: Iterable[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    """与 Qt 一致：过滤自动提取角色行后拼上新候选。"""
    kept, new_entries, removed_auto = build_character_scan_entries(rows, found_names)
    return kept + new_entries, new_entries, removed_auto


# ---------------------------------------------------------------------------
# Excel 导入导出（防公式注入、尺寸限制）
# ---------------------------------------------------------------------------


def _ensure_excel_size(data: bytes) -> None:
    if len(data) > EXCEL_MAX_BYTES:
        raise LexiconOpsError(f"Excel 文件超过 {EXCEL_MAX_BYTES // (1024 * 1024)}MB 限制")


def _sanitize_excel_text(value: Any) -> str:
    text = "" if value is None else str(value)
    if text and text[0] in _FORMULA_PREFIXES:
        return "\t" + text
    return text


def _write_text_cell(cell: Cell, value: Any) -> None:
    text = _sanitize_excel_text(value)
    cell.value = text
    cell.data_type = "s"
    cell.number_format = "@"


def import_glossary_excel_bytes(data: bytes) -> list[dict[str, str]]:
    _ensure_excel_size(data)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        sheet = workbook.active
        if sheet.max_row > 50000 or sheet.max_column > 64:
            raise LexiconOpsError("Excel 超过50000行或64列限制")
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if first_row is None:
            raise LexiconOpsError("Excel 为空")
        headers = [str(cell).strip() if cell is not None else "" for cell in first_row]
        header_map = build_glossary_header_map(headers)
        if "src" not in header_map or "dst" not in header_map:
            raise LexiconOpsError("未找到原文/译文列，请检查表头")
        items: list[dict[str, str]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            src = safe_excel_cell(row, header_map.get("src"))
            dst = safe_excel_cell(row, header_map.get("dst"))
            type_ = safe_excel_cell(row, header_map.get("type"))
            comment = safe_excel_cell(row, header_map.get("comment"))
            if not src:
                continue
            items.append({"src": src, "dst": dst, "type": type_, "comment": comment})
        return items
    finally:
        workbook.close()


def export_glossary_excel_bytes(entries: Sequence[Mapping[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Glossary"
    headers = ("原文", "译文", "类别", "备注")
    for col, header in enumerate(headers, start=1):
        _write_text_cell(sheet.cell(row=1, column=col), header)
    for row_index, item in enumerate(entries, start=2):
        if not isinstance(item, Mapping):
            continue
        values = (
            item.get("src", ""),
            item.get("dst", ""),
            item.get("type", ""),
            item.get("comment", item.get("info", "")),
        )
        for col, value in enumerate(values, start=1):
            _write_text_cell(sheet.cell(row=row_index, column=col), value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    _ensure_excel_size(data)
    return data


def decode_llm_jsonline(response_text: str, expected: int) -> list[str]:
    try:
        import json_repair as repair
    except Exception:
        repair = None

    mapping: dict[str, str] = {}
    for raw in (response_text or "").splitlines():
        line = raw.strip()
        if not line or line.startswith("```"):
            continue
        try:
            data = repair.loads(line) if repair else None  # type: ignore[attr-defined]
        except Exception:
            data = None
        if isinstance(data, dict) and len(data) == 1:
            key, value = next(iter(data.items()))
            if isinstance(key, str) and isinstance(value, str):
                mapping[key] = value

    if not mapping:
        try:
            data = repair.loads(response_text) if repair else None  # type: ignore[attr-defined]
        except Exception:
            data = None
        if isinstance(data, dict):
            for key, value in data.items():
                if isinstance(key, str) and isinstance(value, str):
                    mapping[key] = value

    return [mapping.get(str(i), "") for i in range(expected)]


def apply_translate_results(
    rows: Sequence[Mapping[str, Any]],
    results: Sequence[tuple[int, str] | list[Any]],
) -> tuple[list[dict[str, Any]], int]:
    """将翻译结果写入行；不覆盖已有且不同于原文的译文。"""
    new_rows = [dict(row) if isinstance(row, Mapping) else {} for row in rows]
    applied = 0
    for item in results:
        if not item or len(item) < 2:
            continue
        row_index = int(item[0])
        dst = str(item[1] or "").strip()
        if row_index < 0 or row_index >= len(new_rows) or not dst:
            continue
        src_text = str(new_rows[row_index].get("src", "") or "").strip()
        current_dst = str(new_rows[row_index].get("dst", "") or "").strip()
        if current_dst and current_dst != src_text:
            continue
        new_rows[row_index]["dst"] = dst
        applied += 1
    return new_rows, applied


def glossary_statistics_key(entry: Mapping[str, Any]) -> str:
    src = re.sub(r"\s+", " ", str(entry.get("src", "") or "")).strip()
    case_sensitive = bool(entry.get("case_sensitive", False))
    return f"{src}|{int(case_sensitive)}"


def run_glossary_hit_statistics(config: Config, entries: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    texts = load_counted_source_texts(config)
    if len(texts) == 0:
        return {
            "success": False,
            "counts": [],
            "counted_item_total": 0,
            "cache_dir": get_statistics_cache_dir(config),
            "message": "未找到可统计的缓存条目",
        }
    row_entries = [dict(entry) for entry in entries if isinstance(entry, Mapping)]
    counts = count_glossary_hit_counts(row_entries, texts)
    return {
        "success": True,
        "counts": counts,
        "counted_item_total": len(texts),
        "snapshot_keys": [glossary_statistics_key(entry) for entry in row_entries],
        "message": "统计完成",
    }


# ---------------------------------------------------------------------------
# 禁翻
# ---------------------------------------------------------------------------


def normalize_preserve_src(text: str) -> str:
    if not text:
        return ""
    return text.strip().strip("\"'“”‘’").lower()


def merge_preserve_entries(base: Mapping[str, Any], incoming: Mapping[str, Any]) -> dict[str, Any]:
    def _clean(value: Any) -> str:
        return value.strip() if isinstance(value, str) else ""

    merged = {"src": _clean(base.get("src")), "comment": _clean(base.get("comment") or base.get("info"))}
    incoming_clean = {
        "src": _clean(incoming.get("src")),
        "comment": _clean(incoming.get("comment") or incoming.get("info")),
    }
    if incoming_clean["comment"]:
        if not merged["comment"] or len(incoming_clean["comment"]) > len(merged["comment"]):
            merged["comment"] = incoming_clean["comment"]
    if incoming_clean["src"] and not merged["src"]:
        merged["src"] = incoming_clean["src"]
    return merged


def build_preserve_header_map(headers: Sequence[str]) -> dict[str, int]:
    alias = {
        "src": {"原文", "原始文本", "source", "src", "text"},
        "comment": {"备注", "说明", "comment", "note", "备注信息"},
    }
    mapping: dict[str, int] = {}
    for index, name in enumerate(headers):
        lower_name = str(name or "").lower()
        for key, options in alias.items():
            if lower_name in {opt.lower() for opt in options} and key not in mapping:
                mapping[key] = index
    return mapping


def dedupe_preserve_entries(entries: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    key_index: dict[str, int] = {}
    deduped: list[dict[str, Any]] = []
    for item in entries:
        if not isinstance(item, Mapping):
            continue
        key = normalize_preserve_src(str(item.get("src", "") or ""))
        if not key:
            continue
        normalized = {
            "src": str(item.get("src", "") or "").strip(),
            "comment": str(item.get("comment", item.get("info", "")) or "").strip(),
        }
        if key not in key_index:
            deduped.append(normalized)
            key_index[key] = len(deduped) - 1
        else:
            deduped[key_index[key]] = merge_preserve_entries(deduped[key_index[key]], item)
    return deduped


def normalize_preserve_rows(data: Sequence[Any]) -> list[dict[str, str]]:
    converted: list[dict[str, str]] = []
    for item in data or []:
        if isinstance(item, dict):
            converted.append(
                {
                    "src": str(item.get("src", "") or ""),
                    "comment": str(item.get("comment", item.get("info", "")) or ""),
                }
            )
        elif isinstance(item, str):
            converted.append({"src": item, "comment": ""})
    return converted


def import_preserve_excel_bytes(data: bytes) -> list[dict[str, str]]:
    _ensure_excel_size(data)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        sheet = workbook.active
        if sheet.max_row > 50000 or sheet.max_column > 64:
            raise LexiconOpsError("Excel 超过50000行或64列限制")
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if first_row is None:
            raise LexiconOpsError("Excel 为空")
        headers = [str(cell).strip() if cell is not None else "" for cell in first_row]
        header_map = build_preserve_header_map(headers)
        if "src" not in header_map:
            raise LexiconOpsError("未找到原文列，请检查表头")
        items: list[dict[str, str]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            src = safe_excel_cell(row, header_map.get("src"))
            comment = safe_excel_cell(row, header_map.get("comment"))
            if not src:
                continue
            items.append({"src": src, "comment": comment})
        return items
    finally:
        workbook.close()


def export_preserve_excel_bytes(entries: Sequence[Mapping[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TextPreserve"
    headers = ("原文", "备注")
    for col, header in enumerate(headers, start=1):
        _write_text_cell(sheet.cell(row=1, column=col), header)
    for row_index, item in enumerate(entries, start=2):
        if not isinstance(item, Mapping):
            continue
        values = (item.get("src", ""), item.get("comment", item.get("info", "")))
        for col, value in enumerate(values, start=1):
            _write_text_cell(sheet.cell(row=row_index, column=col), value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    _ensure_excel_size(data)
    return data


def list_preserve_scan_candidates(config: Any) -> list[Path]:
    raws = [
        getattr(config, "input_folder", ""),
        getattr(config, "output_folder", ""),
        getattr(config, "renpy_game_folder", ""),
    ]
    candidates: list[Path] = []
    for raw in raws:
        if not raw:
            continue
        path = Path(raw)
        if not path.exists():
            continue
        if path.is_file():
            if path.suffix.lower() == ".rpy":
                candidates.append(path.parent)
            continue
        if path.name.lower() == "tl" and path.parent.exists():
            candidates.append(path.parent)
        game_child = path / "game"
        if game_child.exists() and game_child.is_dir():
            candidates.append(game_child)
        candidates.append(path)

    deduped: list[Path] = []
    seen: set[str] = set()
    for path in candidates:
        try:
            key = str(path.resolve()).lower()
        except Exception:
            key = str(path).lower()
        if key in seen:
            continue
        seen.add(key)
        deduped.append(path)
    return deduped


def count_rpy_files_without_tl(root: Path) -> int:
    count = 0
    for rpy_file in root.rglob("*.rpy"):
        if "tl" in [part.lower() for part in rpy_file.parts]:
            continue
        count += 1
    return count


def scan_preserve_variables(config: Any, cancel_callback=None) -> dict[str, Any]:
    """扫描 [variable]；返回 {entries, game_path, cleared}。"""
    candidates = list_preserve_scan_candidates(config)
    if not candidates:
        raise LexiconOpsError("没有可用于扫描的文件夹，请先设置输入/输出或游戏目录")

    game_path: Path | None = None
    fallback_path: Path | None = None
    for candidate in candidates:
        if cancel_callback:
            cancel_callback()
        if fallback_path is None:
            fallback_path = candidate
        try:
            count = count_rpy_files_without_tl(candidate)
        except OSError as exc:
            raise LexiconOpsError(f"源码目录读取失败：{candidate}") from exc
        if count > 0:
            game_path = candidate
            break
    if game_path is None:
        game_path = fallback_path
    if game_path is None or not game_path.exists():
        raise LexiconOpsError("无法确定要扫描的文件夹")

    found_preserves: set[str] = set()
    try:
        for rpy_file in game_path.rglob("*.rpy"):
            if cancel_callback:
                cancel_callback()
            if "tl" in [part.lower() for part in rpy_file.parts]:
                continue
            try:
                content = rpy_file.read_text(encoding="utf-8", errors="ignore")
                for var_name in RE_VARIABLE_IN_TEXT.findall(content):
                    found_preserves.add(f"[{var_name}]")
            except OSError as exc:
                raise LexiconOpsError(f"变量源码读取失败：{rpy_file}") from exc
    except Exception as exc:
        raise LexiconOpsError(f"扫描失败: {exc}") from exc

    if not found_preserves:
        return {
            "entries": [],
            "game_path": str(game_path),
            "cleared": True,
            "enabled": False,
            "message": f"未找到变量引用，已清空禁翻表（扫描：{game_path}）",
        }

    entries = [{"src": text} for text in sorted(found_preserves)]
    return {
        "entries": entries,
        "game_path": str(game_path),
        "cleared": False,
        "enabled": True,
        "message": f"找到 {len(entries)} 条变量引用（扫描：{game_path}）",
    }


def run_preserve_hit_statistics(config: Config, entries: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    texts = load_counted_source_texts(config)
    if len(texts) == 0:
        return {
            "success": False,
            "counts": [],
            "counted_item_total": 0,
            "cache_dir": get_statistics_cache_dir(config),
            "message": "未找到可统计的缓存条目",
        }
    row_entries = [dict(entry) for entry in entries if isinstance(entry, Mapping)]
    counts = count_text_preserve_hit_counts(row_entries, texts)
    return {
        "success": True,
        "counts": counts,
        "counted_item_total": len(texts),
        "snapshot_keys": [normalize_preserve_src(str(entry.get("src", "") or "")) for entry in row_entries],
        "message": "统计完成",
    }


# ---------------------------------------------------------------------------
# 翻译
# ---------------------------------------------------------------------------


def translate_glossary_fast(
    tasks: Sequence[tuple[int, str]],
    source_lang: str,
    target_lang: str,
    engine: str = "bing",
) -> list[tuple[int, str]]:
    from module.Engine.FastTranslator import FastTranslator

    if not tasks:
        return []
    translator = FastTranslator(engine=engine)
    srcs = [src for _, src in tasks]
    translated = translator.translate_batch(srcs, target_lang=target_lang, source_lang=source_lang)
    results: list[tuple[int, str]] = []
    for index, (row, _) in enumerate(tasks):
        dst = translated[index] if index < len(translated) else ""
        results.append((row, dst))
    return results


def _convert_chinese_form(config: Config, text: str) -> str:
    try:
        from base.BaseLanguage import BaseLanguage
    except Exception:
        return text
    if str(getattr(config, "target_language", "")).upper() != str(BaseLanguage.Enum.ZH):
        return text
    # 词库批量翻译阶段不做运行时简繁转换
    return text


def translate_glossary_llm(
    tasks: Sequence[tuple[int, str]],
    config: Config,
    platform: Mapping[str, Any],
    batch_size: int = 30,
    cancel_callback: Callable[[], bool] | None = None,
) -> list[tuple[int, str]]:
    """LLM 批量翻译术语；失败/跳过时保留原文，不伪造成功译文。"""
    from base.Base import Base
    from module.Engine.TaskRequester import TaskRequester
    from module.PromptBuilder import PromptBuilder

    if not tasks:
        return []
    if not platform:
        raise LexiconOpsError("未选择翻译引擎，请先配置并启用平台")

    config_for_prompt = dataclasses.replace(
        config,
        glossary_enable=False,
        auto_glossary_enable=False,
    )
    prompt_builder = PromptBuilder(config_for_prompt)
    batch_size = max(1, int(batch_size))
    all_results: list[tuple[int, str]] = []
    total = len(tasks)
    total_batches = (total + batch_size - 1) // batch_size

    for batch_index in range(total_batches):
        if cancel_callback and cancel_callback():
            raise InterruptedError("术语库翻译已取消")

        start = batch_index * batch_size
        batch = list(tasks[start : start + batch_size])
        srcs = [src for _, src in batch]

        if platform.get("api_format") != Base.APIFormat.SAKURALLM:
            messages, _ = prompt_builder.generate_prompt(srcs, [], [], False)
        else:
            messages, _ = prompt_builder.generate_prompt_sakura(srcs)

        requester = TaskRequester(config_for_prompt, dict(platform), batch_index)
        response_shape = (
            "json_object"
            if config_for_prompt.structured_output_enable
            and platform.get("api_format") != Base.APIFormat.SAKURALLM
            and platform.get("api_format") not in Base.MACHINE_API_FORMATS
            else "none"
        )
        skip, _, response_text, _, _ = requester.request(messages, response_shape=response_shape)

        if skip or not response_text:
            # 失败/跳过：保留原文，不当作已翻译成功
            translated = ["" for _ in srcs]
        else:
            translated = decode_llm_jsonline(response_text, len(srcs))
            translated = [
                t if isinstance(t, str) and t.strip() else "" for t, src in zip(translated, srcs)
            ]

        translated = [_convert_chinese_form(config_for_prompt, t) for t in translated]
        for (row, _), dst in zip(batch, translated):
            all_results.append((row, dst))

    return all_results


def decode_content_base64(content_base64: str, *, max_bytes: int = EXCEL_MAX_BYTES) -> bytes:
    raw = str(content_base64 or "").strip()
    if not raw:
        raise LexiconOpsError("content_base64 不能为空")
    # 粗估 base64 膨胀，避免解码前占用过大内存
    if len(raw) > int(max_bytes * 1.4) + 64:
        raise LexiconOpsError("Base64 内容过大")
    try:
        data = base64.b64decode(raw, validate=True)
    except Exception as exc:
        raise LexiconOpsError(f"Base64 解码失败: {exc}") from exc
    if len(data) > max_bytes:
        raise LexiconOpsError(f"解码后数据超过 {max_bytes // (1024 * 1024)}MB 限制")
    return data


def encode_content_base64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")
