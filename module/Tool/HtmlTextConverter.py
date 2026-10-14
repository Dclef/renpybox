# -*- coding: utf-8 -*-
"""HTML / TXT / Excel 互转纯业务，供 Qt HtmlImportPage 与 API 共用。

对齐 frontend/RenpyToolbox/HtmlImportPage.py 高级区真实行为：
- HTML→TXT：普通模式读可见文本；对照模式需 data payload
- TXT→HTML：每行 h6；可选隐藏 div#data JSON payload
- Excel→TXT：跳过元数据表，双语列别名，空内容不写
payload 解析只用 json.loads / ast.literal_eval，禁止 eval。
"""

from __future__ import annotations

import ast
import json
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from bs4 import BeautifulSoup

# Excel 列业务键（与 Qt ComboBox userData 一致，语言切换不改）
EXCEL_COLUMN_ORIGINAL = "原文"
EXCEL_COLUMN_TRANSLATION = "译文"

EXCEL_COLUMN_ALIASES: dict[str, set[str]] = {
    EXCEL_COLUMN_ORIGINAL: {
        "原文",
        "original",
        "原文（勿修改此列）",
        "Source (Do Not Edit)",
    },
    EXCEL_COLUMN_TRANSLATION: {
        "译文",
        "translation",
        "译文（勿修改此列）",
        "Translation (Do Not Edit)",
    },
}

_METADATA_SHEETS = frozenset({"元数据", "Metadata"})


def read_html_payload_texts(soup: BeautifulSoup) -> List[str] | None:
    """从 div#data 解析 payload，返回按 line 排序的文本列表。

    普通模式文本优先非空 current，否则 original（与 Qt 实际代码一致，无 target 回退）。
    """
    data_div = soup.find("div", id="data")
    if data_div is None:
        return None
    raw = (data_div.string or data_div.get_text() or "").strip()
    if not raw:
        return None
    payload = _parse_payload(raw)
    if payload is None:
        return None

    entries: list[tuple[int, str]] = []
    for item in payload:
        if not isinstance(item, dict):
            continue
        line_no = item.get("line")
        current = item.get("current")
        original = item.get("original")
        if current and str(current).strip():
            text = str(current)
        elif original and str(original).strip():
            text = str(original)
        else:
            text = ""
        entries.append((line_no if isinstance(line_no, int) else 0, text))
    entries.sort(key=lambda pair: pair[0])
    return [text for _line, text in entries]


def read_html_strings(path: str | Path) -> List[str]:
    """读取 HTML 可见文本行：优先 payload，否则 h6，再 p/div；过滤空串。"""
    with open(path, "r", encoding="utf-8") as reader:
        soup = BeautifulSoup(reader, "html.parser")

    payload = read_html_payload_texts(soup)
    if payload is not None:
        return payload

    strings = [tag.get_text() for tag in soup.find_all("h6")]
    if not strings:
        strings = [tag.get_text() for tag in soup.find_all(["p", "div"])]
    return [s.replace("\r", "").strip() for s in strings if s and s.strip()]


def read_html_translation_pairs(path: str | Path) -> List[Tuple[str, str]]:
    """读取 (original, current) 对照；译文取 current or target or original；无 payload 返回 []。"""
    with open(path, "r", encoding="utf-8") as reader:
        soup = BeautifulSoup(reader, "html.parser")
    data_div = soup.find("div", id="data")
    if data_div is None:
        return []
    raw = (data_div.string or data_div.get_text() or "").strip()
    if not raw:
        return []
    payload = _parse_payload(raw)
    if payload is None:
        return []

    pairs: list[tuple[int, str, str]] = []
    for item in payload:
        if not isinstance(item, dict):
            continue
        line_no = item.get("line")
        original = item.get("original") or ""
        current = item.get("current") or item.get("target") or original
        pairs.append(
            (
                line_no if isinstance(line_no, int) else 0,
                str(original),
                str(current),
            )
        )
    pairs.sort(key=lambda triple: triple[0])
    return [(orig, cur) for _line, orig, cur in pairs]


def build_html_content(lines: Sequence[str], keep_data: bool = False) -> str:
    """按行生成 h6；keep_data 时写入隐藏 JSON payload。"""
    soup = BeautifulSoup(
        "<html><head><meta charset='utf-8'></head><body></body></html>",
        "html.parser",
    )
    body = soup.body
    data_payload: list[dict] = []
    for idx, text in enumerate(lines):
        value = text or ""
        h6 = soup.new_tag("h6")
        h6.string = value
        body.append(h6)
        if keep_data:
            data_payload.append(
                {
                    "line": idx,
                    "original": value,
                    "target": value,
                    "current": value,
                }
            )
    if keep_data and data_payload:
        data_div = soup.new_tag("div", id="data", style="display: none;")
        data_div.string = json.dumps(data_payload, ensure_ascii=False)
        body.append(data_div)
    return str(soup)


def html_to_txt(
    html_path: str | Path,
    txt_output: str | Path,
    *,
    pairs: bool = False,
) -> tuple[int, Path]:
    """HTML→TXT。pairs=True 需 payload；无内容失败不写。"""
    html_path = Path(html_path)
    txt_output = Path(txt_output)
    if not html_path.is_file():
        raise FileNotFoundError(f"HTML 文件不存在: {html_path}")
    if html_path.resolve() == txt_output.resolve():
        raise ValueError("HTML 输入与 TXT 输出不能是同一文件")

    if pairs:
        pair_list = read_html_translation_pairs(html_path)
        if not pair_list:
            raise ValueError("对照导出需要 HTML 含 data payload，且至少有一条可用文本")
        lines = [f"{orig}\t{cur}" for orig, cur in pair_list]
    else:
        lines = read_html_strings(html_path)
        if not lines:
            raise ValueError("HTML 中没有可导出的文本")

    txt_output.parent.mkdir(parents=True, exist_ok=True)
    txt_output.write_text("\n".join(lines), encoding="utf-8")
    return len(lines), txt_output


def txt_to_html(
    txt_path: str | Path,
    html_output: str | Path,
    *,
    keep_data: bool = False,
) -> tuple[int, Path]:
    """TXT→HTML。保留空行；空文件失败不写。"""
    txt_path = Path(txt_path)
    html_output = Path(html_output)
    if not txt_path.is_file():
        raise FileNotFoundError(f"TXT 文件不存在: {txt_path}")
    if txt_path.resolve() == html_output.resolve():
        raise ValueError("TXT 输入与 HTML 输出不能是同一文件")

    with open(txt_path, "r", encoding="utf-8") as reader:
        lines = [line.rstrip("\n") for line in reader]
    if not lines:
        raise ValueError("TXT 文件为空")

    html = build_html_content(lines, keep_data=keep_data)
    html_output.parent.mkdir(parents=True, exist_ok=True)
    html_output.write_text(html, encoding="utf-8")
    return len(lines), html_output


def excel_to_txt(
    excel_path: str | Path,
    txt_output: str | Path,
    *,
    column: str = EXCEL_COLUMN_TRANSLATION,
) -> tuple[int, Path]:
    """Excel→TXT。跳过元数据表；无内容警告式失败且不写目标。"""
    from openpyxl import load_workbook

    excel_path = Path(excel_path)
    txt_output = Path(txt_output)
    if not excel_path.is_file():
        raise FileNotFoundError(f"Excel 文件不存在: {excel_path}")
    if excel_path.resolve() == txt_output.resolve():
        raise ValueError("Excel 输入与 TXT 输出不能是同一文件")

    target_col = str(column or EXCEL_COLUMN_TRANSLATION).strip() or EXCEL_COLUMN_TRANSLATION
    aliases = {name.lower() for name in EXCEL_COLUMN_ALIASES.get(target_col, {target_col})}

    book = load_workbook(excel_path)
    lines: list[str] = []
    for sheet_name in book.sheetnames:
        if sheet_name in _METADATA_SHEETS:
            continue
        sheet = book[sheet_name]
        if sheet.max_row == 0 or sheet.max_column == 0:
            continue
        headers = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in sheet[1]
        ]
        col_index = None
        for idx, name in enumerate(headers):
            if name.lower() in aliases:
                col_index = idx
                break
        if col_index is None:
            continue
        for row in sheet.iter_rows(min_row=2, values_only=True):
            val = row[col_index] if col_index < len(row) else None
            if val is None:
                continue
            text = str(val).strip()
            if text:
                lines.append(text)

    if not lines:
        raise ValueError("未在 Excel 中找到可导出的内容")

    txt_output.parent.mkdir(parents=True, exist_ok=True)
    txt_output.write_text("\n".join(lines), encoding="utf-8")
    return len(lines), txt_output


def default_sibling_output(input_path: str | Path, suffix: str) -> Path:
    """空输出时的同名旁路路径。"""
    path = Path(input_path)
    return path.with_suffix(suffix)


def _parse_payload(raw: str) -> list | None:
    """解析 JSON 或旧版 literal_eval(list[dict])；绝不 eval。"""
    try:
        payload = json.loads(raw)
    except (ValueError, TypeError):
        try:
            payload = ast.literal_eval(raw)
        except (ValueError, SyntaxError):
            return None
    if not isinstance(payload, list):
        return None
    return payload


__all__ = [
    "EXCEL_COLUMN_ALIASES",
    "EXCEL_COLUMN_ORIGINAL",
    "EXCEL_COLUMN_TRANSLATION",
    "build_html_content",
    "default_sibling_output",
    "excel_to_txt",
    "html_to_txt",
    "read_html_payload_texts",
    "read_html_strings",
    "read_html_translation_pairs",
    "txt_to_html",
]
