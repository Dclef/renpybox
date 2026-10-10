# -*- coding: utf-8 -*-
"""HakimiSuiteRunner

特性：
- 扫描 game 目录（排除 tl）提取角色名/文本/变量/replace
- 外部数据挖掘(.json/.yaml/.yml) + 疯狗模式深度扫描
- 三档模式（标准 / 外部文件 / 外部+疯狗）
- 对比 tl/<lang> 已有 old 翻译，排除重复
- 输出到 translate_output/{1_Excels,2_RPY_Files}，附 AI_Prompt
- 可选生成 Emoji/Tag 保护表（译前/译后）
- 可选先执行官方抽取以刷新 tl
"""

from __future__ import annotations

import os
import re
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, List, Optional, Sequence, Set, Tuple

from openpyxl import Workbook, load_workbook

from base.LogManager import LogManager
from module.Extract.EmojiReplacer import generate_emoji_replacement_sheets
from module.Extract.RenpyExtractor import RenpyExtractor
from module.Renpy.renpy_tl_core import escape_tl_string
from module.Tool.AssetSuiteOps import (
    AssetSuiteCancelled,
    AssetSuiteError,
    _atomic_write_bytes,
    backup_translate_output_if_exists,
    ensure_write_path,
    _atomic_save_workbook,
    raise_if_cancelled,
    sanitize_excel_text,
)
from module.Tool.LanguageTools import require_language
from module.Tool.LanguageTools import LanguageToolsError

try:
    import yaml  # type: ignore
except Exception:  # pragma: no cover - 可选依赖
    yaml = None


# -------------------- 数据结构 -------------------- #
@dataclass
class HakimiResult:
    names_count: int = 0
    others_count: int = 0
    replace_count: int = 0
    deleted_count: int = 0
    emoji_replacements: int = 0
    base_dir: Optional[Path] = None
    excel_dir: Optional[Path] = None
    rpy_dir: Optional[Path] = None
    emoji_dir: Optional[Path] = None
    backup_path: str = ""
    warnings: List[str] = field(default_factory=list)
    written: List[str] = field(default_factory=list)
    partial: bool = False


# -------------------- 核心实现 -------------------- #
class HakimiSuiteRunner:
    def __init__(self, logger: Optional[LogManager] = None, renpy_extractor: Optional[RenpyExtractor] = None) -> None:
        self.logger = logger or LogManager.get()
        self.renpy_extractor = renpy_extractor or RenpyExtractor()

    def run(
        self,
        target_path: str | Path,
        tl_name: str,
        *,
        use_official: bool = False,
        exe_path: str | Path | None = None,
        gen_emoji: bool = False,
        mode: str | int = "1",
        confirm_overwrite: bool = False,
        cancel_check: Callable[[], bool] | None = None,
        progress_callback: Callable[[str, int | None, int | None], None] | None = None,
    ) -> HakimiResult:
        try:
            language = require_language(str(tl_name or "").strip() or "chinese")
        except LanguageToolsError as exc:
            raise ValueError(str(exc)) from exc

        project_root, game_dir, auto_exe = self._resolve_paths(target_path)
        tl_dir = game_dir / "tl" / language
        ensure_write_path(tl_dir, project_root)
        self.warnings: List[str] = []
        mode_str = str(mode).strip() or "1"
        if mode_str not in {"1", "2", "3"}:
            mode_str = "1"
        include_external = mode_str in {"2", "3"}
        mad_dog = mode_str == "3"

        def report(message: str, done: int | None = None, total: int | None = None) -> None:
            if progress_callback is not None:
                progress_callback(message, done, total)

        raise_if_cancelled(cancel_check)

        # 可选官方抽取
        if use_official:
            exe = self._pick_exe(exe_path, auto_exe, project_root)
            if exe is None:
                raise FileNotFoundError("开启官方抽取但未找到可执行文件，请手动选择 exe")
            self.logger.info(f"执行官方抽取: {exe}")
            report("官方抽取中…")
            self.renpy_extractor.official_extract(
                str(exe),
                language,
                generate_empty=False,
                force=True,
                should_stop=cancel_check,
                progress_callback=(lambda message: report(message)),
            )
            raise_if_cancelled(cancel_check)

        # 扫描源码：实际仅跳过 tl（与 _scan_files 一致）
        extensions = (".rpy", ".json", ".yaml", ".yml") if include_external else (".rpy",)
        file_list = self._scan_files(game_dir, extensions, cancel_check=cancel_check)
        rpy_files = [p for p in file_list if p.suffix.lower() == ".rpy"]
        external_files = [p for p in file_list if p.suffix.lower() in {".json", ".yaml", ".yml"}]
        names: List[str] = []
        texts: List[str] = []
        variables: List[str] = []
        replaces: List[str] = []
        sandbox_strings: List[str] = []

        self.logger.info(">>> 开始提取...")
        report("提取文本中…", 0, len(rpy_files) or 1)
        for index, rpy in enumerate(rpy_files, 1):
            raise_if_cancelled(cancel_check)
            n, t, v, r = self._extract_strings_from_rpy(rpy)
            names.extend(n)
            texts.extend(t)
            variables.extend(v)
            replaces.extend(r)
            if mad_dog:
                sandbox_strings.extend(self._extract_deep_python_strings(rpy))
            report(f"提取 {rpy.name}", index, len(rpy_files))

        if include_external and external_files:
            raise_if_cancelled(cancel_check)
            sandbox_strings.extend(self._extract_from_external_files(external_files, cancel_check=cancel_check))

        # 过滤
        self.logger.info(">>> 手术级清洗垃圾中...")
        raise_if_cancelled(cancel_check)
        f_names, d_names = self._filter_strings(names)
        f_texts, d_texts = self._filter_strings(texts)
        f_vars, d_vars = self._filter_strings(variables)
        f_replaces, d_replaces = self._filter_strings(replaces)
        f_sandbox, d_sandbox = self._filter_strings(sandbox_strings, strict_mode=True)
        deleted_total = len(d_names) + len(d_texts) + len(d_vars) + len(d_replaces) + len(d_sandbox)
        if mad_dog and d_sandbox:
            self.logger.info("  >>> 疯狗模式：已拦截 %s 条垃圾变量/字符串", len(d_sandbox))

        # 去除已有翻译
        existing_set = self._extract_existing_translations(tl_dir, cancel_check=cancel_check)

        self.logger.info(">>> 对比 SDK 翻译 (%s)...", language)
        final_names = [s for s in f_names if s not in existing_set]
        others_pool = set(f_texts + f_vars)
        final_others = sorted([s for s in others_pool if s not in existing_set and s not in set(final_names)])

        rep_pool = set(f_replaces)
        if f_sandbox:
            used = set(final_names) | set(final_others)
            for s in f_sandbox:
                if s not in existing_set and s not in used:
                    rep_pool.add(s)
        final_replace = sorted(list(rep_pool))

        # 输出目录
        base_out = project_root / "translate_output"
        ensure_write_path(base_out, project_root)
        backup_path = ""
        written: List[str] = []
        for parent in (base_out, base_out / "1_Excels", base_out / "2_RPY_Files", base_out / "3_Emoji_Tools"):
            ensure_write_path(parent, project_root)
        if base_out.exists():
            for existing in base_out.rglob("*"):
                ensure_write_path(existing, project_root)
        if base_out.exists() and any(base_out.iterdir()):
            if not confirm_overwrite:
                raise FileExistsError("translate_output 已存在内容，需要 confirm_overwrite=true")
            raise_if_cancelled(cancel_check)
            backup_path = backup_translate_output_if_exists(base_out, cancel_check=cancel_check)

        # Emoji 预检：避免结构已写完才发现失败
        emoji_preflight_error = ""
        if gen_emoji:
            try:
                if not tl_dir.exists():
                    raise FileNotFoundError(f"语言目录不存在: {tl_dir}")
                if not list(tl_dir.rglob("*.rpy")):
                    raise ValueError("未找到 .rpy 文件")
            except Exception as exc:
                emoji_preflight_error = str(exc)

        excel_dir = base_out / "1_Excels"
        rpy_dir = base_out / "2_RPY_Files"
        try:
            excel_dir.mkdir(parents=True, exist_ok=True)
            rpy_dir.mkdir(parents=True, exist_ok=True)

            raise_if_cancelled(cancel_check)
            report("写入 Excel…")
            self._save_to_excel(final_names, excel_dir / "names.xlsx", ["Original", "Translation"], cancel_check=cancel_check, allowed_root=base_out)
            written.append(str(excel_dir / "names.xlsx"))
            raise_if_cancelled(cancel_check)
            self._save_to_excel(final_others, excel_dir / "others.xlsx", ["Original", "Translation"], cancel_check=cancel_check, allowed_root=base_out)
            written.append(str(excel_dir / "others.xlsx"))
            raise_if_cancelled(cancel_check)
            self._save_to_excel(final_replace, excel_dir / "replace_text.xlsx", ["Text", "Replacement"], cancel_check=cancel_check, allowed_root=base_out)
            written.append(str(excel_dir / "replace_text.xlsx"))

            # AI Prompt
            ai_prompt = base_out / "AI_Prompt_Names.txt"
            prompt_lines = ["请翻译以下游戏角色名："]
            prompt_lines.extend(final_names)
            raise_if_cancelled(cancel_check)
            _atomic_write_bytes(ai_prompt, "\n".join(prompt_lines) + "\n", newline="\n", has_bom=False, cancel_check=cancel_check, allowed_root=base_out)
            written.append(str(ai_prompt))

            # RPY：空池时删除旧残留，避免吞掉人工编辑后仍留下过期文件
            raise_if_cancelled(cancel_check)
            if self._write_or_clear_rpy(final_names, rpy_dir / "translate_names.rpy", language, cancel_check=cancel_check, allowed_root=base_out):
                written.append(str(rpy_dir / "translate_names.rpy"))
            raise_if_cancelled(cancel_check)
            if self._write_or_clear_rpy(final_others, rpy_dir / "translate_others.rpy", language, cancel_check=cancel_check, allowed_root=base_out):
                written.append(str(rpy_dir / "translate_others.rpy"))
            raise_if_cancelled(cancel_check)
            replace_path = rpy_dir / "replace.rpy"
            if final_replace:
                self._generate_replace_rpy(final_replace, replace_path, language, cancel_check=cancel_check, allowed_root=base_out)
                written.append(str(replace_path))
            elif replace_path.exists():
                ensure_write_path(replace_path, base_out)
                raise_if_cancelled(cancel_check)
                replace_path.unlink()
                written.append(str(replace_path))

            emoji_dir: Optional[Path] = None
            emoji_count = 0
            if gen_emoji:
                emoji_dir = base_out / "3_Emoji_Tools"
                emoji_dir.mkdir(parents=True, exist_ok=True)
                if emoji_preflight_error:
                    self.logger.warning("Emoji 对照表跳过：%s", emoji_preflight_error)
                else:
                    try:
                        raise_if_cancelled(cancel_check)
                        report("生成 Emoji 对照表…")
                        emoji_count, pre_path, post_path = generate_emoji_replacement_sheets(
                            tl_dir,
                            emoji_dir,
                            cancel_check=cancel_check,
                        )
                        written.extend([str(pre_path), str(post_path)])
                        for source, alias in ((pre_path, "Tag_Protection_Pre(译前).xlsx"), (post_path, "Tag_Protection_Post(译后).xlsx")):
                            book = load_workbook(source)
                            try:
                                _atomic_save_workbook(book, emoji_dir / alias, cancel_check=cancel_check, allowed_root=base_out)
                            finally:
                                book.close()
                            written.append(str(emoji_dir / alias))
                    except AssetSuiteCancelled:
                        raise
                    except Exception as exc:
                        self.logger.warning("Emoji 对照表生成失败（结构输出已写入）：%s", exc)
                        emoji_preflight_error = str(exc)

            result = HakimiResult(
                names_count=len(final_names),
                others_count=len(final_others),
                replace_count=len(final_replace),
                deleted_count=deleted_total,
                emoji_replacements=emoji_count,
                base_dir=base_out,
                excel_dir=excel_dir,
                rpy_dir=rpy_dir,
                emoji_dir=emoji_dir,
                backup_path=backup_path,
                warnings=self.warnings + ([emoji_preflight_error] if emoji_preflight_error else []),
                written=written,
                partial=bool(self.warnings or emoji_preflight_error),
            )
            if emoji_preflight_error:
                setattr(result, "emoji_warning", emoji_preflight_error)
            return result
        except Exception as exc:
            payload = {"success": False, "level": "warning" if isinstance(exc, AssetSuiteCancelled) else "error",
                       "output_dir": str(base_out), "backup_path": backup_path,
                       "written": written, "partial": bool(written), "cancelled": isinstance(exc, AssetSuiteCancelled),
                       "message": str(exc)}
            if isinstance(exc, AssetSuiteCancelled):
                exc.result.update(payload)
                raise
            raise AssetSuiteError(str(exc), result=payload) from exc


    # -------------------- 辅助函数 -------------------- #
    def _resolve_paths(self, target: str | Path) -> Tuple[Path, Path, Optional[Path]]:
        path = Path(target).expanduser().resolve()
        exe_path: Optional[Path] = None
        base = path
        if path.is_file():
            exe_path = path
            base = path.parent
        if base.name.lower() == "game":
            base = base.parent
        game_dir = base / "game"
        ensure_write_path(game_dir, base)
        if not game_dir.exists():
            raise FileNotFoundError(f"未找到 game 目录: {game_dir}")
        return base, game_dir, exe_path

    def _pick_exe(self, explicit: str | Path | None, auto: Path | None, project_root: Path) -> Optional[Path]:
        if explicit:
            candidate = Path(explicit).expanduser().resolve()
            if candidate.exists():
                return candidate
        if auto and auto.exists():
            return auto
        return self._auto_find_exe(project_root)

    def _auto_find_exe(self, root: Path) -> Optional[Path]:
        for pattern in ("*.exe", "*.py"):
            candidates = [p for p in root.glob(pattern) if p.is_file()]
            if candidates:
                candidates.sort(key=lambda p: p.stat().st_size if p.exists() else 0, reverse=True)
                return candidates[0]
        return None

    def _scan_files(
        self,
        root: Path,
        extensions: Tuple[str, ...],
        *,
        cancel_check: Callable[[], bool] | None = None,
    ) -> List[Path]:
        # v7.5：仅跳过 tl 目录
        ignored_dirs = {"tl"}
        files: List[Path] = []
        self.logger.info(">>> 开始建立扫描索引 (已屏蔽 tl)...")
        for dirpath, dirnames, filenames in os.walk(root):
            raise_if_cancelled(cancel_check)
            dirnames[:] = [d for d in dirnames if d.lower() not in ignored_dirs]
            for filename in filenames:
                if filename.lower().endswith(extensions):
                    files.append(Path(dirpath) / filename)
        self.logger.info("✔ 索引建立完成。共找到 %s 个有效文件。", len(files))
        return files

    def _filter_strings(self, strings: Sequence[str], *, strict_mode: bool = False) -> Tuple[List[str], List[str]]:
        """过滤逻辑（含 strict_mode：沙盒/疯狗模式）。"""
        filtered_list: List[str] = []
        deleted_list: List[str] = []

        file_extensions = (
            ".mp3",
            ".png",
            ".jpg",
            ".jpeg",
            ".ogg",
            ".wav",
            ".webp",
            ".gif",
            ".avi",
            ".mp4",
            ".mov",
            ".webm",
            ".flv",
            ".wmv",
            ".rpy",
            ".py",
            ".json",
            ".yaml",
            ".yml",
            ".ttf",
            ".otf",
            ".xml",
            ".csv",
        )
        code_keywords = {
            "true",
            "false",
            "none",
            "null",
            "return",
            "jump",
            "call",
            "label",
            "screen",
            "style",
            "transform",
            "image",
            "define",
            "default",
            "init",
            "python",
            "if",
            "else",
            "elif",
            "for",
            "while",
            "in",
            "and",
            "or",
            "not",
            "pass",
            "break",
            "continue",
            "set",
            "get",
            "music",
            "sound",
            "play",
            "stop",
            "scene",
            "show",
            "hide",
            "with",
            "at",
            "persistent",
        }

        for raw in strings:
            s = raw if isinstance(raw, str) else str(raw)
            original = s
            if not s or s.strip() == "":
                deleted_list.append(original)
                continue

            # 含汉字（中日共用区段）默认不抽取（避免把已汉化内容当成待翻译）
            if self._has_chinese(s):
                deleted_list.append(original)
                continue

            lower = s.lower()
            if any(ext in lower for ext in file_extensions) and " " not in s:
                deleted_list.append(original)
                continue
            if s.isdigit():
                deleted_list.append(original)
                continue
            if lower in code_keywords:
                deleted_list.append(original)
                continue

            temp = re.sub(r"\{.*?\}", "", s)
            temp = re.sub(r"\[.*?\]", "", temp)
            if not re.search(r"[\u4e00-\u9fa5a-zA-Z]", temp):
                deleted_list.append(original)
                continue

            if (s.startswith("[") and s.endswith("]")) or (s.startswith("{") and s.endswith("}")):
                if not ("{/" in s or " " in s):
                    deleted_list.append(original)
                    continue

            if s.startswith("#"):
                deleted_list.append(original)
                continue

            if strict_mode:
                if " " not in s and len(s) < 4:
                    deleted_list.append(original)
                    continue
                if s.islower() and "_" in s and " " not in s:
                    deleted_list.append(original)
                    continue
                if s.isupper() and " " not in s:
                    deleted_list.append(original)
                    continue
                if "/" in s and " " not in s:
                    deleted_list.append(original)
                    continue

            filtered_list.append(s)

        return sorted(list(set(filtered_list))), deleted_list

    def _unescape(self, s: str) -> str:
        return s.replace('\\"', '"').replace("\\'", "'").replace("\\\\", "\\")

    def _has_chinese(self, s: str) -> bool:
        return bool(re.search(r"[\u4e00-\u9fa5]", s))

    def _is_camel_case(self, s: str) -> bool:
        return bool(s) and s[0].islower() and any(x.isupper() for x in s) and " " not in s

    def _extract_strings_from_rpy(self, file_path: Path) -> Tuple[List[str], List[str], List[str], List[str]]:
        name_strings: List[str] = []
        text_strings: List[str] = []
        variable_strings: List[str] = []
        replace_strings: List[str] = []

        try:
            content = file_path.read_text(encoding="utf-8", errors="replace")
        except Exception as exc:
            self.warnings.append(f"读取失败 {file_path}: {exc}")
            self.logger.warning(self.warnings[-1])
            return name_strings, text_strings, variable_strings, replace_strings

        # 角色名
        char_patterns = [
            r'Character\s*\(\s*(["\'])((?:\\\1|.)*?)\1',
            r'define\s+\w+\s*=\s*Character\s*\(\s*(["\'])((?:\\\1|.)*?)\1',
        ]
        for pattern in char_patterns:
            for match in re.finditer(pattern, content, re.IGNORECASE):
                string = self._unescape(match.group(2))
                if not re.search(r'_\s*\(\s*["\']' + re.escape(string), content):
                    name_strings.append(string)

        # 文本
        text_patterns = [
            r'\btext\s+(["\'])((?:\\\1|.)*?)\1\s*:',
            r'\b(text|textbutton|show\s+text)\s+(["\'])((?:\\\2|.)*?)\2',
            r'renpy\.input\s*\(\s*(["\'])((?:\\\1|.)*?)\1',
        ]
        for pattern in text_patterns:
            for match in re.finditer(pattern, content, re.IGNORECASE | re.MULTILINE):
                idx = 3 if pattern == text_patterns[1] else 2
                string = self._unescape(match.group(idx))
                start_pos = match.start()
                line_start = content.rfind('\n', 0, start_pos)
                preceding = content[line_start:start_pos]
                if not re.search(r'_\s*\(\s*$', preceding.strip()):
                    text_strings.append(string)

        # 字典字符串（如 "safe": "..."、"text": "..."、"lines": ["..."] 等）
        dict_patterns = [
            r'"safe"\s*:\s*(["\'])((?:\\\1|.)*?)\1',
            r'"text"\s*:\s*(["\'])((?:\\\1|.)*?)\1',
            r'"lines"\s*:\s*\[\s*(["\'])((?:\\\1|.)*?)\1',
        ]
        for pattern in dict_patterns:
            for match in re.finditer(pattern, content, re.IGNORECASE):
                string = self._unescape(match.group(2))
                # 检查是否已经被翻译标记（_()）排除
                start_pos = match.start()
                line_start = content.rfind('\n', 0, start_pos)
                preceding = content[line_start:start_pos]
                if not re.search(r'_\s*\(\s*$', preceding.strip()):
                    text_strings.append(string)

        # 变量 / 特殊调用
        variable_keywords = [r'default\s+\w+\s*=\s*', r'define\s+\w+\s*=\s*', r'\$\s*\w+\s*=\s*']
        for line in content.split('\n'):
            for keyword in variable_keywords:
                if re.search(keyword, line) and "Character" not in line and not re.search(r'_\s*\(', line):
                    for match in re.finditer(r'(["\'])((?:\\\1|.)*?)\1', line):
                        variable_strings.append(self._unescape(match.group(2)))

            if ('f"' in line or "f'" in line) and not re.search(r'_\s*\(\s*f', line):
                for match in re.finditer(r'f(["\'])((?:\\\1|.)*?)\1', line):
                    replace_strings.append(self._unescape(match.group(2)))

            if re.search(r'^\s*\$\s*(renpy\.notify|csay)\s*\(', line):
                for match in re.finditer(r'(["\'])((?:\\\1|.)*?)\1', line):
                    if not re.search(r'_\s*\(\s*$', line[:match.start()].rstrip()):
                        text_strings.append(self._unescape(match.group(2)))

        tooltip_pattern = r'\btooltip\s*\(\s*(["\'])((?:\\\1|.)*?)\1'
        for match in re.finditer(tooltip_pattern, content, re.IGNORECASE):
            replace_strings.append(self._unescape(match.group(2)))

        return name_strings, text_strings, variable_strings, replace_strings

    def _extract_from_external_files(
        self,
        file_list: Sequence[Path],
        *,
        cancel_check: Callable[[], bool] | None = None,
    ) -> List[str]:
        self.logger.info("  >>> 启动 [外部挖掘机] ...")
        found: List[str] = []
        for path in file_list:
            raise_if_cancelled(cancel_check)
            ext = path.suffix.lower()
            if ext not in {".json", ".yaml", ".yml"}:
                continue

            try:
                text = path.read_text(encoding="utf-8", errors="replace")
            except Exception as exc:
                self.warnings.append(f"读取失败 {path}: {exc}")
                continue

            try:
                data: Any
                if ext == ".json":
                    data = json.loads(text)
                else:
                    if yaml is None:
                        self.logger.warning("未安装 PyYAML，已跳过: %s", path.name)
                        continue
                    data = yaml.safe_load(text)
                if data is not None:
                    self._recursive_find_strings(data, found)
            except Exception as exc:
                self.warnings.append(f"外部文件解析失败 {path}: {exc}")
                continue

        self.logger.info("  >>> 外部提取: %s 条", len(found))
        return found

    def _recursive_find_strings(self, data: Any, found_list: List[str]) -> None:
        if isinstance(data, str):
            found_list.append(data)
        elif isinstance(data, list):
            for item in data:
                self._recursive_find_strings(item, found_list)
        elif isinstance(data, dict):
            for value in data.values():
                self._recursive_find_strings(value, found_list)

    def _extract_deep_python_strings(self, path: Path) -> List[str]:
        """疯狗模式：抽取 rpy 里所有引号字符串（后续用 strict_mode 强力过滤）。"""
        found: List[str] = []
        try:
            content = path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            return found
        matches = re.findall(r'(["\'])((?:\\\1|.)*?)\1', content)
        for _, s in matches:
            found.append(self._unescape(s))
        return found

    def _extract_existing_translations(
        self,
        tl_dir: Path,
        *,
        cancel_check: Callable[[], bool] | None = None,
    ) -> Set[str]:
        existing: Set[str] = set()
        if not tl_dir.exists():
            return existing
        for rpy in tl_dir.rglob("*.rpy"):
            raise_if_cancelled(cancel_check)
            try:
                content = rpy.read_text(encoding="utf-8", errors="replace")
            except Exception:
                continue
            olds = re.findall(r'^\s*old\s+(["\'])((?:\\\1|.)*?)\1', content, re.MULTILINE)
            for _, s in olds:
                existing.add(self._unescape(s))
        return existing

    def _save_to_excel(self, strings: Sequence[str], path: Path, headers: Sequence[str], *, cancel_check=None, allowed_root=None) -> None:
        book = Workbook()
        sheet = book.active
        sheet.append([sanitize_excel_text(headers[0]), sanitize_excel_text(headers[1])])
        for value in strings:
            raise_if_cancelled(cancel_check)
            sheet.append([sanitize_excel_text(value), ""])
        sheet.column_dimensions["A"].width = 50
        sheet.column_dimensions["B"].width = 50
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=2):
            for cell in row:
                cell.data_type = "s"
                cell.number_format = "@"
        _atomic_save_workbook(book, path, cancel_check=cancel_check, allowed_root=allowed_root or path.parent)
        self.logger.info("已保存 %s 条到 %s", len(strings), path)

    def _write_or_clear_rpy(self, strings: Sequence[str], output_path: Path, lang_folder: str, *, cancel_check=None, allowed_root=None) -> bool:
        raise_if_cancelled(cancel_check)
        ensure_write_path(output_path, allowed_root or output_path.parent)
        if not strings:
            if output_path.exists():
                output_path.unlink()
                return True
            return False
        self._generate_rpy_file(strings, output_path, lang_folder, cancel_check=cancel_check, allowed_root=allowed_root)
        return True

    def _generate_rpy_file(self, strings: Sequence[str], output_path: Path, lang_folder: str, *, cancel_check=None, allowed_root=None) -> None:
        if not strings:
            return
        lines = [f"translate {lang_folder} strings:", ""]
        for s in strings:
            raise_if_cancelled(cancel_check)
            escaped = escape_tl_string(s)
            lines.append(f'    old "{escaped}"')
            lines.append('    new ""')
            lines.append("")
        _atomic_write_bytes(output_path, "\n".join(lines), newline="\n", has_bom=False, cancel_check=cancel_check, allowed_root=allowed_root)

    def _generate_replace_rpy(self, strings: Sequence[str], output_path: Path, lang_folder: str, *, cancel_check=None, allowed_root=None) -> None:
        if not strings:
            return
        sorted_strings = sorted(list(set(strings)), key=len, reverse=True)
        safe_lang = escape_tl_string(lang_folder)
        lines = [
            "init python:",
            "    # Generated by Sandbox Special Edition",
            f'    if preferences.language == "{safe_lang}":',
            "        def replace_text(s):",
            "            if not isinstance(s, str): return s",
        ]
        for s in sorted_strings:
            raise_if_cancelled(cancel_check)
            escaped = escape_tl_string(s)
            lines.append(f'            s = s.replace("{escaped}", "{escaped}") # 待翻译: {escaped}')
        lines.extend([
            "            return s",
            "        config.replace_text = replace_text",
        ])
        _atomic_write_bytes(output_path, "\n".join(lines), newline="\n", has_bom=False, cancel_check=cancel_check, allowed_root=allowed_root)


__all__ = ["HakimiSuiteRunner", "HakimiResult"]
