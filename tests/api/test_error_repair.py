"""错误修复 API：真实临时目录扫描/修复/导出/lint 边界。"""
from __future__ import annotations

import json
import os
import subprocess
import threading
import time
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from api.jobs import JobManager
from api.routes.error_repair import router
from api.routes.jobs import router as jobs_router
from module.Config import Config
from module.Engine.Engine import Engine
from module.Renpy.ProjectPaths import RenpyProjectPaths
from module.Tool.ErrorRepairer import ErrorRepairer


@pytest.fixture
def repair_client(tmp_path, monkeypatch):
    config = Config()
    engine = Engine()
    monkeypatch.setattr(Engine, "get", classmethod(lambda cls: engine))
    monkeypatch.setattr(Config, "save", lambda self, *args, **kwargs: self)
    app = FastAPI()
    app.state.config = config
    app.state.jobs = JobManager()
    app.include_router(router)
    app.include_router(jobs_router)
    with TestClient(app) as client:
        yield client, engine, tmp_path, app.state.jobs, config
        app.state.jobs.shutdown()


def _finish(client, job_id, timeout=8.0):
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in {"done", "failed", "cancelled"}:
            return job
        time.sleep(0.05)
    raise AssertionError(f"job {job_id} did not finish")


def _tl(tmp_path: Path) -> Path:
    tl = tmp_path / "game" / "tl" / "chinese"
    tl.mkdir(parents=True)
    (tl / "script.rpy").write_text(
        'translate chinese strings:\n'
        '    old "Hello"\n'
        '    new "你好"\n'
        '\n'
        'translate chinese empty_block:\n'
        '    old "x"\n'
        '    new ""\n',
        encoding="utf-8",
    )
    # 含行首 Tab，供修复
    (tl / "tabs.rpy").write_text('\told "A"\n\tnew "A"\n', encoding="utf-8")
    return tl


def _bind(config: Config, project: Path, language: str = "chinese") -> str:
    game = project / "game"
    tl = game / "tl" / language
    tl.mkdir(parents=True, exist_ok=True)
    config.renpy_project_path = str(project)
    config.renpy_game_folder = str(game)
    config.renpy_tl_folder = str(tl)
    config.input_folder = str(tl)
    paths = RenpyProjectPaths.from_config(config)
    assert paths is not None
    return paths.project_key


def _make_junction(link: Path, target: Path) -> None:
    completed = subprocess.run(
        ["cmd", "/c", "mklink", "/J", str(link), str(target)],
        capture_output=True,
        text=True,
    )
    if completed.returncode != 0 or not link.exists():
        pytest.skip(f"无法创建目录联接: {completed.stderr or completed.stdout}")


def test_scan_quotes_always_on_and_empty_report_exportable(repair_client):
    client, engine, tmp_path, _, _ = repair_client
    tl = _tl(tmp_path)
    response = client.post("/api/error-repair/scan", json={
        "folder": str(tl),
        "check_indent": True,
        "check_indent_level": False,
    })
    assert response.status_code == 200, response.text
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "done", job
    assert job["result"]["action"] == "scan"
    assert "report" not in job["result"]
    assert job["result"]["has_scan_report"] is True
    assert job["result"]["affected_files"] == 1
    # 空报告也应可导出
    out = tmp_path / "report.xlsx"
    exported = client.post("/api/error-repair/export", json={
        "folder": str(tl),
        "output_path": str(out),
        "report": {},
    })
    assert exported.status_code == 200, exported.text
    assert out.is_file()
    assert engine.get_status() == Engine.Status.IDLE


def test_fix_partial_counts_and_confirm_required(repair_client):
    client, engine, tmp_path, _, _ = repair_client
    tl = _tl(tmp_path)
    assert client.post("/api/error-repair/fix", json={
        "folder": str(tl),
        "confirm": False,
    }).status_code == 400
    response = client.post("/api/error-repair/fix", json={
        "folder": str(tl),
        "fix_indent": True,
        "fix_indent_level": False,
        "fix_quotes": False,
        "fix_dialogue_quotes": False,
        "confirm": True,
    })
    assert response.status_code == 200, response.text
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "done", job
    result = job["result"]
    assert "fixed_files" in result and "fixed_items" in result and "failed_files" in result
    # Only explicitly selected operations are applied by the desktop API.
    assert result["action"] == "fix"
    assert engine.get_status() == Engine.Status.IDLE


def test_fix_can_run_lint_recheck_in_the_same_job(repair_client, monkeypatch):
    """Lint 复检并入「扫描并修复」：同一任务里出结果，误报提示单独计数。"""
    client, _, tmp_path, _, config = repair_client
    project = tmp_path / "89G"
    owner_key = _bind(config, project)
    (project / "LustfulDesires.py").write_text("", encoding="utf-8")
    (project / "LustfulDesires.exe").write_bytes(b"MZ\0\0")
    (project / "unins000.exe").write_bytes(b"MZ\0\0")
    tl = project / "game" / "tl" / "chinese"
    (tl / "script.rpy").write_text('\told "A"\n', encoding="utf-8")

    seen: dict = {}
    report = (
        "game/tl/chinese/script.rpy:1 this file contains orphan translations at line 1 (id x).\n"
        "game/tl/chinese/script.rpy:2 Could not evaluate 'suspect1' in the who part of a say statement.\n"
        "game/script.rpy:2 Could not evaluate 'suspect1' in the who part of a say statement.\n"
    )

    def fake_lint(self, path, **kwargs):
        seen["exe"] = path
        return report

    monkeypatch.setattr("api.routes.error_repair.ErrorRepairer.exec_renpy_lint", fake_lint)
    response = client.post("/api/error-repair/fix", json={
        "folder": str(tl), "fix_indent": True, "confirm": True, "lint_after": True,
        "owner_project_key": owner_key,
    })
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "done", job
    lint = job["result"]["lint"]
    # 游戏主程序自动识别（不含 unins000.exe），译文只留真问题，who 提示单独计数。
    assert seen["exe"].endswith("LustfulDesires.exe")
    assert lint["issue_count"] == 1
    assert lint["issues"][0]["file"] == "game/tl/chinese/script.rpy"
    assert lint["ignored_hints"] == 1
    assert "LustfulDesires.exe" in job["result"]["message"] or "Lint 复检" in job["result"]["message"]


def test_export_failure_not_fake_success(repair_client, monkeypatch):
    client, _, tmp_path, _, _ = repair_client
    tl = _tl(tmp_path)
    monkeypatch.setattr(
        "api.routes.error_repair.ErrorRepairer.export_error_report",
        lambda self, errors, output_path: False,
    )
    response = client.post("/api/error-repair/export", json={
        "folder": str(tl),
        "output_path": str(tmp_path / "nope.xlsx"),
        "report": {},
    })
    assert response.status_code == 500


def test_lint_none_fails_empty_succeeds(repair_client, monkeypatch):
    client, engine, tmp_path, _, _ = repair_client
    tl = _tl(tmp_path)
    exe = tmp_path / "game.exe"
    exe.write_bytes(b"MZ\0\0")

    monkeypatch.setattr(
        "api.routes.error_repair.ErrorRepairer.exec_renpy_lint",
        lambda self, path, **kwargs: None,
    )
    response = client.post("/api/error-repair/lint", json={
        "folder": str(tl),
        "game_exe": str(exe),
    })
    assert response.status_code == 200
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "failed"

    monkeypatch.setattr(
        "api.routes.error_repair.ErrorRepairer.exec_renpy_lint",
        lambda self, path, **kwargs: "",
    )
    monkeypatch.setattr(
        "api.routes.error_repair.ErrorRepairer.parse_lint_errors",
        lambda self, output: [
            {"file": "game/tl/chinese/script.rpy", "line": 2, "message": "译文缺图"},
            {"file": "game/script.rpy", "line": 1, "message": "本体缺图"},
            {"raw": "统计行"},
        ],
    )
    response = client.post("/api/error-repair/lint", json={
        "folder": str(tl),
        "game_exe": str(exe),
    })
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "done"
    # 译文条目走 error_count（旧字段含义不变），源码条目单独给一组。
    assert job["result"]["error_count"] == 1
    assert job["result"]["source_error_count"] == 1
    assert job["result"]["tl_errors"][0]["file"] == "game/tl/chinese/script.rpy"
    assert job["result"]["source_errors"][0]["file"] == "game/script.rpy"
    assert engine.get_status() == Engine.Status.IDLE


def test_status_suggests_game_exe_without_manual_browse(repair_client):
    """深度 Lint 的「游戏主程序」按项目自动识别：同名 .py + .exe 才算启动器。"""
    client, _, tmp_path, _, config = repair_client
    project = tmp_path / "89G"
    _bind(config, project)
    (project / "LustfulDesires.py").write_text("", encoding="utf-8")
    exe = project / "LustfulDesires.exe"
    exe.write_bytes(b"MZ\0\0")
    # 卸载器这类不该被当成启动器。
    (project / "unins000.exe").write_bytes(b"MZ\0\0")

    payload = client.get("/api/error-repair").json()
    assert payload["game_exe"] == str(exe)
    assert payload["folder"].endswith("chinese")


def test_lint_scope_checks_translation_and_source_separately(repair_client, monkeypatch):
    """译文 / 源码分开检查：同一份报告，按 scope 只回填对应那一组的明细。"""
    client, _, tmp_path, _, _ = repair_client
    tl = _tl(tmp_path)
    exe = tmp_path / "game.exe"
    exe.write_bytes(b"MZ\0\0")
    monkeypatch.setattr(
        "api.routes.error_repair.ErrorRepairer.exec_renpy_lint",
        lambda self, path, **kwargs: "game/tl/chinese/script.rpy:1 译文问题\ngame/script.rpy:2 源码问题",
    )

    def run(scope):
        response = client.post("/api/error-repair/lint", json={
            "folder": str(tl), "game_exe": str(exe), "scope": scope,
        })
        return _finish(client, response.json()["job"]["id"])["result"]

    only_tl = run("tl")
    assert only_tl["scope"] == "tl"
    assert [item["file"] for item in only_tl["tl_errors"]] == ["game/tl/chinese/script.rpy"]
    assert only_tl["source_errors"] == []
    assert only_tl["error_count"] == 1 and only_tl["source_error_count"] == 1

    only_source = run("source")
    assert [item["file"] for item in only_source["source_errors"]] == ["game/script.rpy"]
    assert only_source["tl_errors"] == []


def test_lint_reports_timeout_reason_and_forwards_wait(repair_client, monkeypatch):
    """超时不是「跑不起来」：要把上限与原因如实带给界面，且默认等待远大于旧的 120 秒。"""
    client, _, tmp_path, _, _ = repair_client
    tl = _tl(tmp_path)
    exe = tmp_path / "game.exe"
    exe.write_bytes(b"MZ\0\0")
    seen: dict = {}
    limit = int(ErrorRepairer.LINT_TIMEOUT)

    def fake(self, path, **kwargs):
        seen.update(kwargs)
        self.last_lint_error = f"超过 {limit} 秒上限仍未结束，已中止（大型项目首次 lint 常需数分钟）"
        return None

    monkeypatch.setattr("api.routes.error_repair.ErrorRepairer.exec_renpy_lint", fake)
    response = client.post("/api/error-repair/lint", json={"folder": str(tl), "game_exe": str(exe)})
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "failed"
    assert str(limit) in job["error"]
    assert seen["timeout"] == ErrorRepairer.LINT_TIMEOUT
    assert limit > 120
    assert callable(seen["on_tick"])


def test_busy_engine_blocks_export_without_side_effects(repair_client):
    """探针 1：busy 下 export 必须 409 且不写盘；IDLE 可导出。"""
    client, engine, tmp_path, manager, _ = repair_client
    tl = _tl(tmp_path)
    out_busy = tmp_path / "busy.xlsx"
    out_idle = tmp_path / "idle.xlsx"
    assert engine.try_set_status(Engine.Status.IDLE, Engine.Status.TRANSLATING)
    busy = client.post("/api/error-repair/export", json={
        "folder": str(tl),
        "output_path": str(out_busy),
        "report": {},
    })
    assert busy.status_code == 409
    assert not out_busy.exists()
    assert not manager.list()
    assert engine.try_set_status(Engine.Status.TRANSLATING, Engine.Status.IDLE)
    idle = client.post("/api/error-repair/export", json={
        "folder": str(tl),
        "output_path": str(out_idle),
        "report": {},
    })
    assert idle.status_code == 200, idle.text
    assert out_idle.is_file()
    assert engine.get_status() == Engine.Status.IDLE


def test_busy_engine_no_side_effects(repair_client):
    client, engine, tmp_path, manager, _ = repair_client
    tl = _tl(tmp_path)
    assert engine.try_set_status(Engine.Status.IDLE, Engine.Status.TRANSLATING)
    assert client.post("/api/error-repair/scan", json={"folder": str(tl)}).status_code == 409
    assert not manager.list()


def test_status_exposes_owner_snapshot(repair_client):
    client, _, tmp_path, _, config = repair_client
    owner = tmp_path / "Owner"
    owner_key = _bind(config, owner)
    tl = owner / "game" / "tl" / "chinese"
    (tl / "a.rpy").write_text('old "x"\n', encoding="utf-8")
    status = client.get("/api/error-repair", params={"folder": str(tl)}).json()
    assert status["owner_project_key"] == owner_key
    assert status["project_key"]


def test_manual_target_with_owner_then_replay_rejected(repair_client):
    """探针 3：owner A 时手工 target B 合法；owner→C 后旧 body 重放 409。"""
    client, engine, tmp_path, _, config = repair_client
    owner = tmp_path / "Owner"
    other = tmp_path / "Other"
    manual = tmp_path / "ManualB-repair"
    owner_key = _bind(config, owner)
    _bind_other = other / "game" / "tl" / "chinese"
    _bind_other.mkdir(parents=True)
    manual.mkdir(parents=True)
    (manual / "scene.rpy").write_text('\told "A"\n\tnew "A"\n', encoding="utf-8")

    status = client.get("/api/error-repair", params={"folder": str(manual)}).json()
    assert status["owner_project_key"] == owner_key
    target_key = status["project_key"]
    body = {
        "folder": str(manual),
        "project_key": target_key,
        "owner_project_key": owner_key,
        "fix_indent": True,
        "confirm": True,
    }
    legal = client.post("/api/error-repair/fix", json=body)
    assert legal.status_code == 200, legal.text
    job = _finish(client, legal.json()["job"]["id"])
    assert job["status"] == "done"
    assert "\t" not in (manual / "scene.rpy").read_text(encoding="utf-8")

    # 切到 owner C，旧 body 必须拒绝且不再写入
    (manual / "scene.rpy").write_text('\told "B"\n\tnew "B"\n', encoding="utf-8")
    other_key = _bind(config, other)
    assert other_key != owner_key
    delayed = client.post("/api/error-repair/fix", json=body)
    assert delayed.status_code == 409
    assert "\t" in (manual / "scene.rpy").read_text(encoding="utf-8")
    assert engine.get_status() == Engine.Status.IDLE


def test_unbound_manual_target_allowed(repair_client):
    """未绑定 owner 时手工目录可合法执行。"""
    client, engine, tmp_path, _, config = repair_client
    config.renpy_project_path = config.renpy_game_folder = config.renpy_tl_folder = ""
    config.input_folder = ""
    manual = tmp_path / "UnboundB"
    manual.mkdir()
    (manual / "scene.rpy").write_text('\told "A"\n', encoding="utf-8")
    status = client.get("/api/error-repair", params={"folder": str(manual)}).json()
    assert status["owner_project_key"] == ""
    response = client.post("/api/error-repair/fix", json={
        "folder": str(manual),
        "project_key": status["project_key"],
        "owner_project_key": "",
        "fix_indent": True,
        "confirm": True,
    })
    assert response.status_code == 200, response.text
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "done"
    assert engine.get_status() == Engine.Status.IDLE


def test_fix_cancel_after_first_write_reports_partial(repair_client, monkeypatch):
    """探针 4：首文件已写后取消 → written/partial=true 与增量计数。"""
    client, engine, tmp_path, _, config = repair_client
    owner = tmp_path / "Owner"
    owner_key = _bind(config, owner)
    tl = owner / "game" / "tl" / "chinese"
    a = tl / "a.rpy"
    b = tl / "b.rpy"
    a.write_text('\told "A"\n', encoding="utf-8")
    b.write_text('\told "B"\n', encoding="utf-8")

    real_fix = ErrorRepairer.auto_fix_file
    wrote_first = threading.Event()
    release = threading.Event()
    calls = {"n": 0}

    def gated(self, file_path, *args, **kwargs):
        ok, count = real_fix(self, file_path, *args, **kwargs)
        calls["n"] += 1
        if calls["n"] == 1 and count > 0:
            wrote_first.set()
            assert release.wait(5)
        return ok, count

    monkeypatch.setattr(ErrorRepairer, "auto_fix_file", gated)

    status = client.get("/api/error-repair", params={"folder": str(tl)}).json()
    response = client.post("/api/error-repair/fix", json={
        "folder": str(tl),
        "project_key": status["project_key"],
        "owner_project_key": owner_key,
        "fix_indent": True,
        "confirm": True,
    })
    assert response.status_code == 200, response.text
    job_id = response.json()["job"]["id"]
    assert wrote_first.wait(3)
    assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 200
    release.set()
    job = _finish(client, job_id, timeout=8.0)
    assert job["status"] == "cancelled"
    result = job["result"] or {}
    assert result.get("written") is True
    assert result.get("partial") is True
    assert int(result.get("fixed_files") or 0) >= 1
    assert result.get("rescan_skipped") is True
    assert "\t" not in a.read_text(encoding="utf-8")
    assert engine.get_status() == Engine.Status.IDLE


def test_fix_owner_change_before_first_write_blocks(repair_client, monkeypatch):
    """探针 5：扫描完成后、首写前切 owner → 阻止写入。"""
    client, engine, tmp_path, _, config = repair_client
    owner = tmp_path / "Owner"
    other = tmp_path / "Other"
    owner_key = _bind(config, owner)
    tl = owner / "game" / "tl" / "chinese"
    a = tl / "a.rpy"
    b = tl / "b.rpy"
    before_a = '\told "A"\n'
    before_b = '\told "B"\n'
    a.write_text(before_a, encoding="utf-8")
    b.write_text(before_b, encoding="utf-8")

    scanned = threading.Event()
    proceed = threading.Event()
    real_scan = ErrorRepairer.scan_folder

    def gated_scan(self, *args, **kwargs):
        result = real_scan(self, *args, **kwargs)
        scanned.set()
        assert proceed.wait(5)
        return result

    monkeypatch.setattr(ErrorRepairer, "scan_folder", gated_scan)

    status = client.get("/api/error-repair", params={"folder": str(tl)}).json()
    response = client.post("/api/error-repair/fix", json={
        "folder": str(tl),
        "project_key": status["project_key"],
        "owner_project_key": owner_key,
        "fix_indent": True,
        "confirm": True,
    })
    assert response.status_code == 200, response.text
    job_id = response.json()["job"]["id"]
    assert scanned.wait(3)
    _bind(config, other)
    proceed.set()
    job = _finish(client, job_id, timeout=8.0)
    assert job["status"] in {"failed", "cancelled"}
    result = job["result"] or {}
    assert result.get("written") in (False, None)
    assert a.read_text(encoding="utf-8") == before_a
    assert b.read_text(encoding="utf-8") == before_b
    assert engine.get_status() == Engine.Status.IDLE


def test_fix_rejects_junction_escape(repair_client):
    """探针 6：选中目录下 junction 指向外部时拒绝，外部文件不得改写。"""
    client, engine, tmp_path, _, _ = repair_client
    root = tmp_path / "LinkedTL"
    outside = tmp_path / "Outside"
    root.mkdir()
    outside.mkdir()
    outside_file = outside / "outside.rpy"
    outside_file.write_text('\told "OUT"\n', encoding="utf-8")
    _make_junction(root / "linked", outside)
    before = outside_file.read_text(encoding="utf-8")
    response = client.post("/api/error-repair/fix", json={
        "folder": str(root),
        "fix_indent": True,
        "confirm": True,
    })
    assert response.status_code in {400, 409}
    assert outside_file.read_text(encoding="utf-8") == before
    assert engine.get_status() == Engine.Status.IDLE


def test_large_scan_status_stays_small_and_full_report_exports(repair_client):
    from openpyxl import load_workbook

    client, _, tmp_path, _, _ = repair_client
    tl = tmp_path / "tl"
    tl.mkdir()
    script = tl / "large.rpy"
    script.write_text('translate chinese strings:\n' + ''.join(
        f'    old "Hello {i} [name]"\n    new "你好 {i}"\n' for i in range(2000)
    ), encoding="utf-8")
    before = script.read_bytes()
    response = client.post("/api/error-repair/scan", json={"folder": str(tl)})
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "done", job
    assert job["result"]["total_issues"] == 2000
    assert job["result"]["affected_files"] == 1
    assert job["result"]["scanned_files"] == 1
    for _ in range(2):
        status = client.get("/api/error-repair", params={"folder": str(tl)})
        assert len(status.content) < 3000
        assert "report" not in status.json()["job"]["result"]
        assert status.json()["has_scan_report"] is True
    output = tmp_path / "full.xlsx"
    exported = client.post("/api/error-repair/export", json={
        "folder": str(tl), "output_path": str(output),
    })
    assert exported.status_code == 200, exported.text
    with output.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True)
        assert workbook.active.max_row == 2001
        workbook.close()
    assert script.read_bytes() == before


def test_scan_cancel_during_single_file_releases_engine_and_keeps_previous_report(repair_client, monkeypatch):
    from importlib import import_module

    client, engine, tmp_path, _, _ = repair_client
    tl = _tl(tmp_path)
    first = client.post("/api/error-repair/scan", json={"folder": str(tl)})
    assert _finish(client, first.json()["job"]["id"])["status"] == "done"
    module = import_module("module.Tool.ErrorRepairer")
    parse = module.parse_tl_document
    entered, release = threading.Event(), threading.Event()
    touched = False

    def gated_parse(lines, *, checkpoint):
        def tick():
            nonlocal touched
            if not touched:
                touched = True
                entered.set()
                assert release.wait(5)
            checkpoint()
        return parse(lines, checkpoint=tick)

    monkeypatch.setattr(module, "parse_tl_document", gated_parse)
    before = {path: path.read_bytes() for path in tl.glob("*.rpy")}
    response = client.post("/api/error-repair/scan", json={"folder": str(tl)})
    job_id = response.json()["job"]["id"]
    try:
        assert entered.wait(3)
        assert client.get("/api/error-repair", params={"folder": str(tl)}).status_code == 200
        assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 200
    finally:
        release.set()
    assert _finish(client, job_id)["status"] == "cancelled"
    assert engine.get_status() == Engine.Status.IDLE
    assert client.get("/api/error-repair", params={"folder": str(tl)}).json()["has_scan_report"]
    assert all(path.read_bytes() == value for path, value in before.items())


def test_fix_with_all_options_off_preserves_empty_blocks(repair_client):
    client, _, tmp_path, _, _ = repair_client
    tl = tmp_path / "tl"
    tl.mkdir()
    script = tl / "empty.rpy"
    original = b'translate chinese strings:\n    # intentionally empty\n'
    script.write_bytes(original)
    body = {"folder": str(tl), "fix_indent": False, "confirm": True}
    response = client.post("/api/error-repair/fix", json=body)
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "done" and job["result"]["fixed_items"] == 0
    assert script.read_bytes() == original
    assert not list(tl.glob("*.bak"))
    response = client.post("/api/error-repair/fix", json={**body, "fix_empty_blocks": True})
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "done" and job["result"]["fixed_items"] == 1
    assert script.read_bytes() == b"    # intentionally empty\n"
    [backup] = list(tl.glob("*.bak"))
    assert backup.read_bytes() == original


def test_checker_failure_is_not_reported_as_clean_scan(repair_client, monkeypatch):
    client, engine, tmp_path, _, _ = repair_client
    tl = _tl(tmp_path)

    def fail_check(*args):
        raise ValueError("scan probe failed")

    monkeypatch.setattr(ErrorRepairer, "_detect_quote_issue", fail_check)
    response = client.post("/api/error-repair/scan", json={"folder": str(tl)})
    job = _finish(client, response.json()["job"]["id"])
    assert job["status"] == "failed"
    assert "scan probe failed" in job["error"]
    assert "total_issues" not in job["result"]
    assert not client.get("/api/error-repair", params={"folder": str(tl)}).json()["has_scan_report"]
    assert engine.get_status() == Engine.Status.IDLE


def _scan_job(client, folder: Path, **body):
    status = client.get("/api/error-repair", params={"folder": str(folder)}).json()
    response = client.post("/api/error-repair/scan", json={
        "folder": str(folder),
        "project_key": status["project_key"],
        "owner_project_key": status["owner_project_key"],
        **body,
    })
    assert response.status_code == 200, response.text
    return _finish(client, response.json()["job"]["id"])


def _fix_job(client, folder: Path, **body):
    status = client.get("/api/error-repair", params={"folder": str(folder)}).json()
    response = client.post("/api/error-repair/fix", json={
        "folder": str(folder),
        "project_key": status["project_key"],
        "owner_project_key": status["owner_project_key"],
        "confirm": True,
        **body,
    })
    assert response.status_code == 200, response.text
    return _finish(client, response.json()["job"]["id"])


def test_scan_returns_type_breakdown_and_manual_sample(repair_client):
    """扫描结果分组：按类型统计 + 少量需人工样本，不返回全量列表。"""
    client, engine, tmp_path, _, _ = repair_client
    tl = tmp_path / "game" / "tl" / "chinese"
    tl.mkdir(parents=True)
    (tl / "tabs.rpy").write_text('\told "A"\n\tnew "A"\n', encoding="utf-8")
    (tl / "placeholder.rpy").write_text(
        'translate chinese strings:\n'
        '    old "Hello [name]"\n'
        '    new "你好"\n',
        encoding="utf-8",
    )
    job = _scan_job(client, tl)
    assert job["status"] == "done", job
    result = job["result"]
    breakdown = result["issues_by_type"]
    assert breakdown and sum(breakdown.values()) == result["total_issues"]
    assert result["fixable_issues"] + result["manual_issues"] == result["total_issues"]
    assert result["fixable_issues"] >= 1
    assert result["manual_issues"] >= 1
    sample = result["manual_sample"]
    assert sample and len(sample) <= 10
    assert all(item["fixable"] is False for item in sample)
    assert {"file", "line", "type", "message"} <= set(sample[0])
    assert result["environment_issues"] == []
    assert "report" not in result
    assert engine.get_status() == Engine.Status.IDLE


def test_scan_manual_sample_is_capped(repair_client):
    """需人工样本最多 10 条，避免轮询快照被大目录撑爆。"""
    client, _, tmp_path, _, _ = repair_client
    tl = tmp_path / "game" / "tl" / "chinese"
    tl.mkdir(parents=True)
    (tl / "many.rpy").write_text(
        'translate chinese strings:\n'
        + ''.join(f'    old "hello {i} [name]"\n    new "你好 {i}"\n' for i in range(15)),
        encoding="utf-8",
    )
    job = _scan_job(client, tl)
    assert job["status"] == "done", job
    result = job["result"]
    assert result["total_issues"] >= 15
    assert len(result["manual_sample"]) == 10
    assert result["manual_issues"] >= 10


def test_fix_reports_rescan_fields_and_applies_new_options(repair_client):
    """一键修复：修完复扫，返回分组结果；新增选项生效。"""
    client, engine, tmp_path, _, _ = repair_client
    tl = _tl(tmp_path)
    job = _fix_job(
        client,
        tl,
        fix_indent=True,
        fix_trailing_spaces=True,
        fix_empty_blocks=True,
        fix_duplicates=True,
    )
    assert job["status"] == "done", job
    result = job["result"]
    assert result["action"] == "fix"
    assert result["written"] is True
    assert result["fixed_files"] >= 1
    assert result["fixed_items"] >= 1
    assert result["rescan_skipped"] is False
    assert isinstance(result["fixed_by_type"], dict)
    assert isinstance(result["remaining_by_type"], dict)
    assert isinstance(result["remaining_issues"], int)
    assert isinstance(result["manual_sample"], list)
    assert result["environment_fixes"] == []
    assert not (tl / "tabs.rpy").read_text(encoding="utf-8").startswith("\t")
    assert list(tl.glob("*.bak"))
    assert engine.get_status() == Engine.Status.IDLE


def test_scan_detects_incremental_leftover_and_fix_moves_it(repair_client):
    """tl/<lang>_new 残留目录：扫描报环境问题，修复可原样移出游戏目录。"""
    client, engine, tmp_path, _, config = repair_client
    owner = tmp_path / "Owner"
    owner_key = _bind(config, owner)
    assert owner_key
    tl = owner / "game" / "tl" / "chinese"
    common = tl / "common.rpy"
    common.write_text(
        'translate chinese strings:\n'
        '    old "Self-voicing disabled."\n'
        '    new "已关闭自带语音。"\n',
        encoding="utf-8",
    )
    leftover = owner / "game" / "tl" / "chinese_new"
    leftover.mkdir()
    (leftover / "common.rpy").write_text(
        'translate chinese strings:\n'
        '    old "Self-voicing disabled."\n'
        '    new "Self-voicing disabled."\n',
        encoding="utf-8",
    )

    scan_job = _scan_job(client, tl)
    assert scan_job["status"] == "done", scan_job
    environment = scan_job["result"]["environment_issues"]
    assert len(environment) == 1, environment
    issue = environment[0]
    assert issue["kind"] == "incremental_leftover"
    assert issue["type"] == "incremental_leftover"
    assert issue["language"] == "chinese"
    assert issue["issue_count"] >= 1
    assert {"merge", "move"} <= set(issue["actions"])

    fix_job = _fix_job(client, tl, fix_environment=True, environment_action="move")
    assert fix_job["status"] == "done", fix_job
    result = fix_job["result"]
    assert result["rescan_skipped"] is False
    assert len(result["environment_fixes"]) == 1
    assert result["environment_fixes"][0]["ok"] is True
    assert result["environment_fixes"][0]["action"] == "move"
    assert result["written"] is True
    assert not leftover.exists()
    assert common.is_file()
    assert json.loads(json.dumps(result))  # 结果必须可序列化给前端
    assert engine.get_status() == Engine.Status.IDLE


def _nc(path) -> str:
    # API 返回的路径经过 normcase；Windows 上大小写与 tmp_path 不同。
    return os.path.normcase(os.path.abspath(str(path)))


_ENV_ONLY = {
    "fix_indent": False,
    "fix_trailing_spaces": False,
    "fix_empty_blocks": False,
    "fix_duplicates": False,
    "fix_dialogue_quotes": False,
    "fix_indent_level": False,
    "fix_quotes": False,
    "fix_environment": True,
}


def _leftover(project: Path, language: str) -> Path:
    leftover = project / "game" / "tl" / f"{language}_new"
    leftover.mkdir(parents=True)
    (leftover / "common.rpy").write_text(
        f'translate {language} strings:\n    old "Hi"\n    new "Hi {language}"\n',
        encoding="utf-8",
    )
    return leftover


def test_environment_card_action_only_handles_selected_leftover(repair_client):
    client, engine, tmp_path, _, config = repair_client
    owner = tmp_path / "Owner"
    _bind(config, owner)
    tl = owner / "game" / "tl" / "chinese"
    (tl / "common.rpy").write_text(
        'translate chinese strings:\n    old "Hi"\n    new "你好"\n', encoding="utf-8",
    )
    chinese_new = _leftover(owner, "chinese")
    japanese_new = _leftover(owner, "japanese")
    scan = _scan_job(client, tl)
    assert scan["status"] == "done", scan
    assert {_nc(issue["path"]) for issue in scan["result"]["environment_issues"]} == {
        _nc(chinese_new), _nc(japanese_new),
    }

    job = _fix_job(
        client, tl, **_ENV_ONLY, environment_action="move", environment_path=str(japanese_new),
    )

    assert job["status"] == "done", job
    result = job["result"]
    assert [_nc(item["path"]) for item in result["environment_fixes"]] == [_nc(japanese_new)]
    assert result["environment_fixes"][0]["ok"] is True
    assert result["written"] is True
    assert result["partial"] is False
    assert result["fixed_files"] == 0
    assert not japanese_new.exists()
    assert (chinese_new / "common.rpy").is_file()
    [moved] = list((owner / "_renpybox_backup").iterdir())
    assert moved.name.startswith("japanese_new_")

    # 已处理掉的目录不在新扫描结果里：拒绝且不碰其他残留目录。
    stale = _fix_job(
        client, tl, **_ENV_ONLY, environment_action="move", environment_path=str(japanese_new),
    )
    assert stale["status"] == "failed", stale
    assert "扫描结果" in stale["error"]
    assert stale["result"].get("written") in (False, None)
    assert (chinese_new / "common.rpy").is_file()
    assert len(list((owner / "_renpybox_backup").iterdir())) == 1

    # 主目录之外的任意路径同样拒绝。
    outside = tmp_path / "Elsewhere" / "chinese_new"
    outside.mkdir(parents=True)
    rejected = _fix_job(
        client, tl, **_ENV_ONLY, environment_action="move", environment_path=str(outside),
    )
    assert rejected["status"] == "failed", rejected
    assert outside.is_dir() and (chinese_new / "common.rpy").is_file()

    status = client.get("/api/error-repair", params={"folder": str(tl)}).json()
    without_flag = client.post("/api/error-repair/fix", json={
        "folder": str(tl),
        "project_key": status["project_key"],
        "owner_project_key": status["owner_project_key"],
        "confirm": True,
        "environment_path": str(chinese_new),
    })
    assert without_flag.status_code == 400
    assert engine.get_status() == Engine.Status.IDLE


def test_environment_only_merge_reports_written_and_backup(repair_client, monkeypatch):
    client, engine, tmp_path, _, config = repair_client
    owner = tmp_path / "Owner"
    _bind(config, owner)
    tl = owner / "game" / "tl" / "chinese"
    leftover = _leftover(owner, "chinese")
    calls = []

    def fake_fix(self, issue, action="merge"):
        calls.append((issue["path"], action))
        return {
            "path": issue["path"], "action": action, "ok": True,
            "message": "merged", "backup": str(owner / "_renpybox_backup" / "merge.zip"),
        }

    monkeypatch.setattr(ErrorRepairer, "fix_environment_issue", fake_fix)
    job = _fix_job(client, tl, **_ENV_ONLY, environment_action="merge", environment_path=str(leftover))

    assert job["status"] == "done", job
    result = job["result"]
    assert [(_nc(path), action) for path, action in calls] == [(_nc(leftover), "merge")]
    assert result["written"] is True
    assert result["partial"] is False
    assert result["success"] is True
    assert result["environment_fixes"][0]["backup"].endswith("merge.zip")
    assert engine.get_status() == Engine.Status.IDLE


def test_environment_merge_failure_after_backup_still_reports_written(repair_client, monkeypatch):
    client, engine, tmp_path, _, config = repair_client
    owner = tmp_path / "Owner"
    _bind(config, owner)
    tl = owner / "game" / "tl" / "chinese"
    leftover = _leftover(owner, "chinese")
    backup = owner / "_renpybox_backup" / "merge.zip"

    def fail_after_backup(self, issue, action="merge"):
        backup.parent.mkdir(parents=True)
        backup.write_bytes(b"verified archive")
        return {
            "path": issue["path"], "action": action, "ok": False,
            "message": "merge failed", "backup": str(backup),
        }

    monkeypatch.setattr(ErrorRepairer, "fix_environment_issue", fail_after_backup)
    job = _fix_job(client, tl, **_ENV_ONLY, environment_action="merge", environment_path=str(leftover))

    assert job["status"] == "done", job
    result = job["result"]
    assert result["written"] is True
    assert result["partial"] is True
    assert result["fixed_files"] == 0
    assert result["fixed_items"] == 0
    assert result["environment_fixes"][0]["ok"] is False
    assert result["environment_fixes"][0]["backup"] == str(backup)
    assert backup.read_bytes() == b"verified archive"
    assert engine.get_status() == Engine.Status.IDLE


def _dedupe_project(tmp_path: Path, config: Config) -> tuple[Path, dict[str, Path]]:
    owner = tmp_path / "Owner"
    _bind(config, owner)
    tl = owner / "game" / "tl" / "chinese"
    files = {}
    for name in ("a", "b", "c"):
        files[name] = tl / f"{name}.rpy"
        files[name].write_text(
            f'translate chinese strings:\n    old "Same"\n    new "{name}"\n', encoding="utf-8",
        )
    return tl, files


_DEDUPE_ONLY = {
    "fix_indent": False,
    "fix_trailing_spaces": False,
    "fix_empty_blocks": False,
    "fix_duplicates": True,
    "fix_dialogue_quotes": False,
    "fix_indent_level": False,
    "fix_quotes": False,
}


def test_dedupe_owner_switch_after_first_write_blocks_later_writes(repair_client, monkeypatch):
    client, engine, tmp_path, _, config = repair_client
    tl, files = _dedupe_project(tmp_path, config)
    before_c = files["c"].read_bytes()
    real_rewrite = ErrorRepairer._rewrite_with_backup
    writes = []

    def switching_rewrite(self, file_path, *args, **kwargs):
        real_rewrite(self, file_path, *args, **kwargs)
        writes.append(str(file_path))
        if len(writes) == 1:
            _bind(config, tmp_path / "Other")

    monkeypatch.setattr(ErrorRepairer, "_rewrite_with_backup", switching_rewrite)
    job = _fix_job(client, tl, **_DEDUPE_ONLY)

    assert job["status"] == "failed", job
    assert "项目已切换" in job["error"]
    result = job["result"]
    assert [_nc(path) for path in writes] == [_nc(files["b"])]
    assert result["written"] is True
    assert result["partial"] is True
    assert result["fixed_files"] == 1
    assert result["fixed_items"] == 1
    assert result["rescan_skipped"] is True
    assert files["c"].read_bytes() == before_c
    assert not list(tl.glob("c.rpy.*.bak"))
    assert engine.get_status() == Engine.Status.IDLE


def test_dedupe_cancel_after_first_file_keeps_written_counts(repair_client, monkeypatch):
    client, engine, tmp_path, _, config = repair_client
    tl, files = _dedupe_project(tmp_path, config)
    before_c = files["c"].read_bytes()
    real_rewrite = ErrorRepairer._rewrite_with_backup
    wrote_first = threading.Event()
    release = threading.Event()

    def gated_rewrite(self, file_path, *args, **kwargs):
        real_rewrite(self, file_path, *args, **kwargs)
        if not wrote_first.is_set():
            wrote_first.set()
            assert release.wait(5)

    monkeypatch.setattr(ErrorRepairer, "_rewrite_with_backup", gated_rewrite)
    status = client.get("/api/error-repair", params={"folder": str(tl)}).json()
    response = client.post("/api/error-repair/fix", json={
        "folder": str(tl),
        "project_key": status["project_key"],
        "owner_project_key": status["owner_project_key"],
        "confirm": True,
        **_DEDUPE_ONLY,
    })
    assert response.status_code == 200, response.text
    job_id = response.json()["job"]["id"]
    assert wrote_first.wait(3)
    assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 200
    release.set()
    job = _finish(client, job_id)

    assert job["status"] == "cancelled", job
    result = job["result"]
    assert result["written"] is True
    assert result["partial"] is True
    assert result["fixed_files"] == 1
    assert result["fixed_items"] == 1
    assert result["rescan_skipped"] is True
    assert files["b"].read_text(encoding="utf-8") == '    # old "Same"\n    # new "b"\n'
    assert files["c"].read_bytes() == before_c
    assert engine.get_status() == Engine.Status.IDLE


def _manual_issue_folder(tmp_path: Path) -> Path:
    tl = tmp_path / "game" / "tl" / "chinese"
    tl.mkdir(parents=True)
    (tl / "syntax.rpy").write_text("label start\n    pass\n", encoding="utf-8")
    return tl


_ALL_OFF = {**_DEDUPE_ONLY, "fix_duplicates": False}


def test_budget_exhausted_rescan_does_not_fabricate_fixed_counts(repair_client, monkeypatch):
    client, engine, tmp_path, _, _ = repair_client
    tl = _manual_issue_folder(tmp_path)
    real_scan = ErrorRepairer.scan_folder
    calls = {"n": 0}

    def exhausting_scan(self, *args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 2:
            kwargs["should_cancel"] = lambda: True
        return real_scan(self, *args, **kwargs)

    monkeypatch.setattr(ErrorRepairer, "scan_folder", exhausting_scan)
    job = _fix_job(client, tl, **_ALL_OFF)

    assert job["status"] == "done", job
    result = job["result"]
    assert calls["n"] == 2
    assert result["rescan_skipped"] is True
    assert result["remaining_issues"] is None
    assert result["remaining_by_type"] == {}
    assert result["manual_sample"] == []
    assert result["fixed_items"] == 0
    assert result["fixed_by_type"] == {}
    assert result["written"] is False
    assert "复扫未完成" in result["message"]
    assert engine.get_status() == Engine.Status.IDLE


def test_cancelled_rescan_reports_unknown_remaining(repair_client, monkeypatch):
    client, engine, tmp_path, manager, _ = repair_client
    tl = _manual_issue_folder(tmp_path)
    real_scan = ErrorRepairer.scan_folder
    calls = {"n": 0}

    def cancelling_scan(self, *args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 2:
            [job] = [job for job in manager.list() if job.kind == "error_repair_fix"]
            job.cancel_event.set()
        return real_scan(self, *args, **kwargs)

    monkeypatch.setattr(ErrorRepairer, "scan_folder", cancelling_scan)
    job = _fix_job(client, tl, **_ALL_OFF)

    assert job["status"] == "cancelled", job
    result = job["result"]
    assert result.get("rescan_skipped") is True
    assert result.get("remaining_issues") is None
    assert result.get("fixed_items") == 0
    assert result.get("written") is False
    assert engine.get_status() == Engine.Status.IDLE
