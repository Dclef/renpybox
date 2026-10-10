"""AssetSuiteOps：批量修正、姓名提取与路径边界。"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from module.Tool.AssetSuiteOps import (
    AssetSuiteError,
    apply_batch_corrections,
    export_batch_correction_workbook,
    export_name_glossary,
    extract_character_names,
    resolve_under_root,
)


def _write_result_check(folder: Path, name: str, payload: dict) -> None:
    (folder / name).write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8-sig")


def test_batch_export_merges_groups_and_english_headers(tmp_path: Path):
    input_dir = tmp_path / "checks"
    output_dir = tmp_path / "out"
    input_dir.mkdir()
    output_dir.mkdir()
    _write_result_check(
        input_dir,
        "结果检查_术语.json",
        {"script.rpy|术语": {"Hello": "你好"}},
    )
    _write_result_check(
        input_dir,
        "result_check_similarity.json",
        {"script.rpy": {"Hello": "您好"}},
    )
    _write_result_check(
        input_dir,
        "result_check_untranslated.json",
        {"script.rpy": {"Skip": "跳过"}},
    )

    result = export_batch_correction_workbook(input_dir, output_dir, english=True, confirm_overwrite=False)
    assert result["success"] is True
    assert result["count"] == 1
    book = openpyxl.load_workbook(result["path"])
    sheet = book.active
    assert sheet.cell(1, 1).value == "File"
    assert sheet.cell(1, 5).value == "Correction (Edit This Column)"
    assert sheet.cell(2, 1).value == "script.rpy"
    assert "术语" in sheet.cell(2, 2).value
    assert "similarity" in sheet.cell(2, 2).value
    assert sheet.cell(2, 5).value in {"你好", "您好"}  # 文件枚举顺序不属于导出契约
    assert sheet.cell(2, 3).data_type == "s"


def test_batch_apply_rpy_match_bom_crlf_and_escape(tmp_path: Path):
    root = tmp_path / "tl"
    root.mkdir()
    target = root / "script.rpy"
    body = 'translate chinese strings:\r\n\r\n    old "line\\nA"\r\n    new "旧译文"\r\n'
    target.write_bytes(b"\xef\xbb\xbf" + body.encode("utf-8"))

    workbook = tmp_path / "批量修正.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["File", "Error", "Source", "Translation", "Correction"])
    sheet.append(["script.rpy", "x", "line\nA", "旧译文", "新译文"])
    book.save(workbook)

    result = apply_batch_corrections(workbook, root, confirm=True)
    assert result["success"] is True
    assert result["applied_changes"] == 1
    raw = target.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    assert b"\r\r\n" not in raw
    text = raw.decode("utf-8-sig")
    assert 'new "新译文"' in text
    assert target.with_suffix(".rpy.bak").exists()


def test_batch_apply_rejects_escape_and_keeps_unmatched(tmp_path: Path):
    root = tmp_path / "tl"
    root.mkdir()
    (root / "ok.txt").write_text("alpha", encoding="utf-8")
    workbook = tmp_path / "批量修正.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["f", "e", "s", "t", "c"])
    sheet.append(["../secret.txt", "x", "a", "b", "c"])
    sheet.append(["ok.txt", "x", "src", "alpha", "beta"])
    sheet.append(["ok.txt", "x", "other", "alpha", "gamma"])  # 同译文不同原文，文本模式按首次 dst
    book.save(workbook)

    result = apply_batch_corrections(workbook, root, confirm=True)
    assert result["applied_changes"] == 1
    assert any("越界" in (item.get("reason") or "") or "绝对" in (item.get("reason") or "") or "拒绝" in (item.get("reason") or "") for item in result["unmatched"])
    assert (root / "ok.txt").read_text(encoding="utf-8") == "beta"
    assert (root / "ok.txt.bak").exists()


def test_batch_cancel_before_write(tmp_path: Path):
    root = tmp_path / "tl"
    root.mkdir()
    target = root / "a.txt"
    target.write_text("old", encoding="utf-8")
    workbook = tmp_path / "批量修正.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["f", "e", "s", "t", "c"])
    sheet.append(["a.txt", "x", "s", "old", "new"])
    book.save(workbook)

    with pytest.raises(Exception):
        apply_batch_corrections(workbook, root, confirm=True, cancel_check=lambda: True)
    assert target.read_text(encoding="utf-8") == "old"
    assert not target.with_suffix(".txt.bak").exists()


def test_resolve_under_root_blocks_absolute(tmp_path: Path):
    root = tmp_path / "tl"
    root.mkdir()
    with pytest.raises(AssetSuiteError):
        resolve_under_root(root, str(tmp_path / "outside.txt"))


def test_name_extract_context_json_and_export(tmp_path: Path):
    src = tmp_path / "game"
    src.mkdir()
    (src / "chars.rpy").write_text(
        'define eileen = Character("Eileen")\n'
        'default eileen.mood = "happy"\n'
        'label start:\n'
        '    e "Hi"\n'
        'define eileen = Character("Eileen")\n'
        'default eileen.mood = "happy"\n'
        'default eileen.age = 18\n'
        'label later:\n'
        '    pass\n',
        encoding="utf-8",
    )
    (src / "nested.json").write_text(
        json.dumps({"people": [{"name": "Alice \"A\""}, {"name": "Eileen"}]}, ensure_ascii=False),
        encoding="utf-8",
    )
    (src / "bad.json").write_text("{not json", encoding="utf-8")

    extracted = extract_character_names(src)
    assert extracted["count"] >= 2
    eileen = next(item for item in extracted["entries"] if item["src"] == "Eileen")
    assert eileen["context"].count("\n") >= 2
    assert any(item["src"] == 'Alice "A"' for item in extracted["entries"])
    assert any("JSON 解析失败" in warning for warning in extracted["warnings"])

    txt = export_name_glossary(extracted["entries"], format="txt")
    assert "Eileen -> Eileen #角色姓名" in txt["content"]
    js = export_name_glossary(extracted["entries"], format="json")
    payload = json.loads(js["content"])
    assert payload[0]["src"]
    assert payload[0]["dst"]
    assert "comment" in payload[0]


def test_name_zero_result_does_not_invent_entries(tmp_path: Path):
    empty = tmp_path / "empty"
    empty.mkdir()
    result = extract_character_names(empty)
    assert result["empty"] is True
    assert result["entries"] == []
