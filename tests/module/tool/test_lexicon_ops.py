"""LexiconOps：Excel 往返、去重合并、分类、禁翻扫描、公式注入防护。"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from module.Tool.LexiconOps import (
    LexiconOpsError,
    apply_translate_results,
    auto_categorize_entries,
    build_glossary_header_map,
    build_preserve_header_map,
    categorize_term,
    collect_glossary_translate_tasks,
    decode_llm_jsonline,
    dedupe_glossary_entries,
    dedupe_preserve_entries,
    export_glossary_excel_bytes,
    export_preserve_excel_bytes,
    import_glossary_excel_bytes,
    import_preserve_excel_bytes,
    merge_candidate_entries,
    merge_glossary_entries,
    normalize_glossary_src,
    normalize_preserve_rows,
    scan_preserve_variables,
)


def test_glossary_excel_roundtrip_and_header_aliases():
    rows = [
        {"src": "Alice", "dst": "爱丽丝", "type": "角色", "comment": "主角"},
        {"src": "Sword", "dst": "剑", "type": "物品", "comment": ""},
    ]
    data = export_glossary_excel_bytes(rows)
    imported = import_glossary_excel_bytes(data)
    assert len(imported) == 2
    assert imported[0]["src"] == "Alice"
    assert imported[0]["dst"] == "爱丽丝"

    # 英文表头别名
    from openpyxl import Workbook
    import io

    wb = Workbook()
    sheet = wb.active
    sheet.append(["source", "target", "category", "note"])
    sheet.append(["Bob", "鲍勃", "角色", "x"])
    buf = io.BytesIO()
    wb.save(buf)
    alias_rows = import_glossary_excel_bytes(buf.getvalue())
    assert alias_rows[0]["src"] == "Bob"
    assert build_glossary_header_map(["原文", "译文"]) == {"src": 0, "dst": 1}


def test_formula_injection_export_prefixes_tab():
    import io

    data = export_glossary_excel_bytes(
        [{"src": "=1+1", "dst": "+cmd", "type": "@x", "comment": "-y"}]
    )
    wb = load_workbook(io.BytesIO(data))
    sheet = wb.active
    assert str(sheet.cell(2, 1).value).startswith("\t=")
    assert sheet.cell(2, 1).number_format == "@"
    assert sheet.cell(2, 1).data_type == "s"


def test_excel_size_limit_reject():
    with pytest.raises(LexiconOpsError):
        import_glossary_excel_bytes(b"x" * (5 * 1024 * 1024 + 1))


def test_dedupe_merge_and_candidates():
    assert normalize_glossary_src('  "Alice"  ') == "alice"
    merged = merge_glossary_entries(
        {"src": "Alice", "dst": "", "type": "", "comment": "a"},
        {"src": "Alice", "dst": "爱丽丝", "type": "角色", "comment": "longer note here"},
    )
    assert merged["dst"] == "爱丽丝"
    assert merged["type"] == "角色"

    deduped = dedupe_glossary_entries(
        [
            {"src": "Alice", "dst": "", "type": "", "comment": ""},
            {"src": "alice", "dst": "爱丽丝", "type": "角色", "comment": "n"},
        ]
    )
    assert len(deduped) == 1
    assert deduped[0]["dst"] == "爱丽丝"

    result_rows, added, updated, count_map = merge_candidate_entries(
        [{"src": "Old", "dst": "旧", "type": "", "comment": "", "candidate": False}],
        [{"source": "New", "target": "", "type": "术语", "count": 3}],
    )
    assert added == 1
    assert updated == 0
    assert any(r["src"] == "New" and r["candidate"] is True for r in result_rows)
    assert count_map[normalize_glossary_src("New")] == 3


def test_classify_keywords():
    assert categorize_term("Dark Forest") == "地名"
    assert categorize_term("magic sword") == "物品"
    changed, rows = auto_categorize_entries(
        [
            {"src": "school campus", "type": ""},
            {"src": "Alice", "type": "角色"},
        ]
    )
    assert changed == 1
    assert rows[0]["type"] == "地名"
    assert rows[1]["type"] == "角色"


def test_preserve_excel_and_variable_scan(tmp_path, monkeypatch):
    game = tmp_path / "game"
    game.mkdir()
    (game / "script.rpy").write_text(
        'label start:\n    "Hello [player_name] and [mom]."\n',
        encoding="utf-8",
    )
    (game / "tl" / "chinese").mkdir(parents=True)
    (game / "tl" / "chinese" / "old.rpy").write_text('"[player_name]"\n', encoding="utf-8")

    class FakeConfig:
        input_folder = str(game)
        output_folder = ""
        renpy_game_folder = ""

    result = scan_preserve_variables(FakeConfig())
    srcs = {item["src"] for item in result["entries"]}
    assert "[player_name]" in srcs
    assert "[mom]" in srcs
    # tl 内不应污染（扫描会跳过 tl 路径下的文件）
    assert result["cleared"] is False

    data = export_preserve_excel_bytes([{"src": "[x]", "comment": "c"}])
    imported = import_preserve_excel_bytes(data)
    assert imported[0]["src"] == "[x]"
    assert build_preserve_header_map(["text", "note"]) == {"src": 0, "comment": 1}
    assert normalize_preserve_rows(["[a]", {"src": "[b]", "info": "i"}])[1]["comment"] == "i"

    deduped = dedupe_preserve_entries(
        [{"src": "[a]", "comment": ""}, {"src": "[A]", "comment": "keep"}]
    )
    assert len(deduped) == 1
    assert deduped[0]["comment"] == "keep"


def test_translate_helpers():
    tasks = collect_glossary_translate_tasks(
        [
            {"src": "a", "dst": ""},
            {"src": "b", "dst": "b"},
            {"src": "c", "dst": "已译"},
        ]
    )
    assert tasks == [(0, "a"), (1, "b")]
    rows, applied = apply_translate_results(
        [{"src": "a", "dst": ""}, {"src": "c", "dst": "已译"}],
        [(0, "甲"), (1, "不应覆盖")],
    )
    assert applied == 1
    assert rows[0]["dst"] == "甲"
    assert rows[1]["dst"] == "已译"
    assert decode_llm_jsonline('{"0": "甲"}\n{"1": "乙"}', 2) == ["甲", "乙"]
