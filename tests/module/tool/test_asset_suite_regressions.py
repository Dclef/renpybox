"""资源套件独立回归：真实工作簿、字节保留、写前取消与链接边界。"""
from __future__ import annotations

import json
import os
import subprocess
from pathlib import Path

import openpyxl
import pytest

from module.Extract.EmojiReplacer import apply_replacements_dir, backup_folder
from module.Extract.HakimiSuiteRunner import HakimiSuiteRunner
from module.Tool import AssetSuiteOps as ops


def _workbook(path: Path, rows: list[list[str]]) -> Path:
    book = openpyxl.Workbook()
    book.active.append(["File", "Error", "Source", "Translation", "Correction"])
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()
    return path


def _junction(link: Path, target: Path) -> None:
    if os.name != "nt":
        link.symlink_to(target, target_is_directory=True)
        return
    env = dict(os.environ, ASSET_REVIEW_LINK=str(link), ASSET_REVIEW_TARGET=str(target))
    outcome = subprocess.run(
        ["pwsh.exe", "-NoProfile", "-Command",
         "New-Item -ItemType Junction -Path $env:ASSET_REVIEW_LINK -Target $env:ASSET_REVIEW_TARGET -ErrorAction Stop | Out-Null"],
        env=env, capture_output=True, text=True,
    )
    if outcome.returncode:
        pytest.skip(f"当前环境无法建立临时 junction：{outcome.stderr}")


def _game(tmp_path: Path) -> Path:
    project = tmp_path / "Project"
    game = project / "game"
    game.mkdir(parents=True)
    (game / "script.rpy").write_text('define hero = Character("Eileen")\n', encoding="utf-8")
    return project


def test_review_excel_text_preserves_formula_like_source(tmp_path: Path):
    book = openpyxl.Workbook()
    samples = ["=A1", "+Anne", "-Name", "@handle", "\t+already"]
    for row, value in enumerate(samples, 1):
        ops.write_text_cell(book.active.cell(row, 1), value)
    path = tmp_path / "text.xlsx"
    book.save(path)
    book.close()
    loaded = openpyxl.load_workbook(path, data_only=False)
    try:
        assert [cell.value for cell in loaded.active["A"]] == samples
        assert all(cell.data_type == "s" for cell in loaded.active["A"])
    finally:
        loaded.close()


def test_review_workbook_export_reports_success_after_atomic_save(tmp_path: Path):
    checks = tmp_path / "checks"
    checks.mkdir()
    (checks / "result_check_typo.json").write_text(
        json.dumps({"script.rpy": {"Hello": "Old"}}), encoding="utf-8-sig"
    )
    result = ops.export_batch_correction_workbook(checks, tmp_path / "out")
    assert result["success"] is True
    assert Path(result["path"]).is_file()


def test_review_rpy_matches_source_and_keeps_original_backup(tmp_path: Path):
    target = tmp_path / "script.rpy"
    before = (
        'translate chinese strings:\r\n\r\n'
        '    old "first"\r\n    new "same"\r\n\r\n'
        '    old u"second"\r\n    new u"same"\r\n'
    ).encode("utf-8")
    target.write_bytes(b"\xef\xbb\xbf" + before)
    backup = target.with_suffix(".rpy.bak")
    backup.write_bytes(b"original backup must survive")
    workbook = _workbook(tmp_path / "fix.xlsx", [["script.rpy", "x", "second", "same", 'B\\C"D\nE']])
    result = ops.apply_batch_corrections(workbook, tmp_path, confirm=True)
    assert result["success"] is True
    assert result["applied_changes"] == 1
    expected = before.replace(b'    new u"same"', b'    new "B\\\\C\\"D\\nE"')
    assert target.read_bytes() == b"\xef\xbb\xbf" + expected
    assert backup.read_bytes() == b"original backup must survive"
    assert result["backups"] == [str(backup)]


@pytest.mark.parametrize("relative", ["../outside.txt", "C:/outside.txt", "C:outside.txt", "//server/share/a.txt", "/outside.txt"])
def test_review_workbook_path_rejects_escape(tmp_path: Path, relative: str):
    with pytest.raises(ops.AssetSuiteError):
        ops.resolve_under_root(tmp_path, relative)


def test_review_workbook_rejects_junction_escape(tmp_path: Path):
    root = tmp_path / "tl"
    outside = tmp_path / "outside"
    root.mkdir()
    outside.mkdir()
    target = outside / "a.txt"
    target.write_bytes(b"old")
    _junction(root / "linked", outside)
    workbook = _workbook(tmp_path / "fix.xlsx", [["linked/a.txt", "x", "src", "old", "new"]])
    result = ops.apply_batch_corrections(workbook, root, confirm=True)
    assert result["success"] is False
    assert result["applied_changes"] == 0
    assert target.read_bytes() == b"old"
    assert not (outside / "a.txt.bak").exists()


def test_review_cancel_after_read_does_not_start_write(tmp_path: Path, monkeypatch):
    target = tmp_path / "a.txt"
    target.write_bytes(b"old")
    workbook = _workbook(tmp_path / "fix.xlsx", [["a.txt", "x", "src", "old", "new"]])
    state = {"cancelled": False}
    original = ops._read_text_bytes

    def read_and_cancel(path):
        data = original(path)
        state["cancelled"] = True
        return data

    monkeypatch.setattr(ops, "_read_text_bytes", read_and_cancel)
    with pytest.raises(ops.AssetSuiteCancelled):
        ops.apply_batch_corrections(workbook, tmp_path, confirm=True, cancel_check=lambda: state["cancelled"])
    assert target.read_bytes() == b"old"
    assert not target.with_suffix(".txt.bak").exists()


def test_review_partial_cancel_keeps_written_counts(tmp_path: Path, monkeypatch):
    first = tmp_path / "a.txt"
    second = tmp_path / "b.txt"
    first.write_bytes(b"old")
    second.write_bytes(b"old")
    workbook = _workbook(tmp_path / "fix.xlsx", [
        ["a.txt", "x", "src", "old", "new"],
        ["b.txt", "x", "src", "old", "new"],
    ])
    state = {"cancelled": False}
    original = ops._atomic_write_bytes

    def write_and_cancel(*args, **kwargs):
        original(*args, **kwargs)
        state["cancelled"] = True

    monkeypatch.setattr(ops, "_atomic_write_bytes", write_and_cancel)
    try:
        result = ops.apply_batch_corrections(workbook, tmp_path, confirm=True, cancel_check=lambda: state["cancelled"])
    except ops.AssetSuiteCancelled as exc:
        result = getattr(exc, "result", None) or getattr(exc, "partial_result", None)
    assert isinstance(result, dict), "取消结果必须携带已生效部分，不能只抛无统计异常"
    assert result["applied_changes"] == 1
    assert result["applied_files"] == 1
    assert result["written"] == [str(first)]
    assert result["backups"] == [str(first.with_suffix(".txt.bak"))]
    assert first.read_bytes() == b"new"
    assert second.read_bytes() == b"old"


def test_review_text_preserves_mixed_newlines_outside_replacement(tmp_path: Path):
    target = tmp_path / "mixed.txt"
    original = b"first\r\nold\nlast\r\n"
    target.write_bytes(original)
    workbook = _workbook(tmp_path / "fix.xlsx", [["mixed.txt", "x", "source", "old", "new"]])
    result = ops.apply_batch_corrections(workbook, tmp_path, confirm=True)
    assert target.read_bytes() == original.replace(b"old", b"new")
    assert result["applied_changes"] == 1


def test_review_emoji_rejects_junction_escape(tmp_path: Path):
    root = tmp_path / "tl"
    outside = tmp_path / "outside"
    root.mkdir()
    outside.mkdir()
    target = outside / "a.rpy"
    target.write_bytes(b'old "{b}Name{/b}"\r\n')
    _junction(root / "linked", outside)
    try:
        apply_replacements_dir(root, {"{b}": "EMOJI"})
    except Exception:
        pass
    assert target.read_bytes() == b'old "{b}Name{/b}"\r\n'


def test_review_structure_rejects_output_junction_escape(tmp_path: Path):
    project = _game(tmp_path)
    outside = tmp_path / "outside"
    outside.mkdir()
    _junction(project / "translate_output", outside)
    try:
        HakimiSuiteRunner().run(project, "chinese", confirm_overwrite=True)
    except Exception:
        pass
    assert list(outside.iterdir()) == [], "结构输出不得经预存 junction 写到项目范围外"


def test_review_safe_backup_detects_selected_target_game(tmp_path: Path):
    mapping_project = _game(tmp_path)
    target_project = tmp_path / "OtherProject"
    target = target_project / "game" / "tl" / "chinese"
    target.mkdir(parents=True)
    (target / "a.rpy").write_bytes(b'old "Name"\n')
    backup = backup_folder(target, project_root=mapping_project)
    assert not backup.resolve().is_relative_to((target_project / "game").resolve())


def test_review_names_json_real_format_preserves_escaped_names(tmp_path: Path):
    (tmp_path / "names.json").write_text(
        json.dumps({"people": [{"name": 'Alice "A"'}, {"name": "艾琳"}]}, ensure_ascii=False), encoding="utf-8"
    )
    result = ops.extract_character_names(tmp_path)
    exported = ops.export_name_glossary(result["entries"], format="json")
    payload = json.loads(exported["content"])
    assert {row["src"] for row in payload} == {'Alice "A"', "艾琳"}
    assert all(row["src"] == row["dst"] for row in payload)
    assert all(row["comment"] == "角色姓名" for row in payload)


def test_review_names_json_imports_with_existing_glossary_reader(tmp_path: Path):
    from module.TableManager import TableManager

    exported = ops.export_name_glossary([{"src": "Eileen", "context": "define e"}], format="json")
    output = tmp_path / "names.json"
    output.write_text(exported["content"], encoding="utf-8")
    imported = TableManager("GLOSSARY", []).load_from_json_file(str(output))
    assert [entry["src"] for entry in imported] == ["Eileen"]
    assert imported[0]["dst"] == "Eileen"
    assert imported[0]["info"] == "角色姓名"


@pytest.fixture
def review_api(tmp_path: Path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from api.jobs import JobManager
    from api.routes.asset_suite import router
    from api.routes.jobs import router as jobs_router
    from module.Config import Config
    from module.Engine.Engine import Engine

    project = _game(tmp_path)
    config = Config()
    config.renpy_project_path = str(project)
    config.renpy_game_folder = str(project / "game")
    config.renpy_tl_folder = str(project / "game" / "tl" / "chinese")
    engine = Engine()
    monkeypatch.setattr(Engine, "get", classmethod(lambda cls: engine))
    # 测试不得触碰真实用户配置。
    monkeypatch.setattr(Config, "save", lambda self, *args, **kwargs: self)
    app = FastAPI()
    app.state.config = config
    app.state.jobs = JobManager()
    app.include_router(router)
    app.include_router(jobs_router)
    with TestClient(app) as client:
        yield client, config, engine, project
        app.state.jobs.shutdown()


def _finished_job(client, job_id: str):
    import time
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in {"done", "failed", "cancelled"} and not (job.get("result") or {}).get("worker_active"):
            return job
        time.sleep(0.01)
    pytest.fail(f"测试任务没有收尾：{job}")


@pytest.mark.parametrize("block", ["barrier", "single"])
def test_review_api_rejects_barrier_and_single_tasks(review_api, block: str):
    client, _, engine, project = review_api
    if block == "barrier":
        engine.set_stop_barrier(True)
    else:
        with engine.lock:
            engine.single_task_count = 1
    response = client.post("/api/asset-suite/names/extract", json={"input_dir": str(project)})
    assert response.status_code == 409


def test_review_api_cancelled_structure_does_not_update_config(review_api, monkeypatch):
    import threading
    import time
    from module.Extract.HakimiSuiteRunner import HakimiResult

    client, config, engine, project = review_api
    started = threading.Event()

    def slow_result(self, target_path, tl_name, **kwargs):
        started.set()
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline and not kwargs["cancel_check"]():
            time.sleep(0.005)
        return HakimiResult(base_dir=project / "translate_output", names_count=1)

    monkeypatch.setattr(HakimiSuiteRunner, "run", slow_result)
    before = config.renpy_game_folder
    request = client.post("/api/asset-suite/structure", json={"path": str(project), "confirm_overwrite": True})
    job_id = request.json()["job"]["id"]
    assert started.wait(2)
    assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 200
    job = _finished_job(client, job_id)
    assert job["status"] == "cancelled"
    assert engine.get_status() == engine.Status.IDLE
    assert config.renpy_game_folder == before, "取消任务不能在返回时再更新当前项目配置"


def test_review_api_project_switch_keeps_new_config_and_filters_old_job(review_api, monkeypatch, tmp_path: Path):
    import threading
    from module.Extract.HakimiSuiteRunner import HakimiResult

    client, config, _, project = review_api
    started = threading.Event()
    release = threading.Event()

    def slow_result(self, target_path, tl_name, **kwargs):
        started.set()
        assert release.wait(3)
        return HakimiResult(base_dir=project / "translate_output", names_count=1)

    monkeypatch.setattr(HakimiSuiteRunner, "run", slow_result)
    request = client.post("/api/asset-suite/structure", json={"path": str(project), "confirm_overwrite": True})
    assert started.wait(2)
    other = tmp_path / "Other"
    (other / "game").mkdir(parents=True)
    config.renpy_project_path = str(other)
    config.renpy_game_folder = str(other / "game")
    config.renpy_tl_folder = str(other / "game" / "tl" / "chinese")
    release.set()
    job = _finished_job(client, request.json()["job"]["id"])
    assert job["status"] == "done"
    assert config.renpy_game_folder == str(other / "game")
    assert client.get("/api/asset-suite/status").json()["job"] is None


def test_review_hakimi_modes_have_real_pool_difference(tmp_path: Path):
    project = _game(tmp_path)
    (project / "game" / "data.json").write_text('{"text": "Only external words"}', encoding="utf-8")
    (project / "game" / "deep.rpy").write_text('$ unknown_call("Only deep words")\n', encoding="utf-8")
    pools = {}
    for mode in ("1", "2", "3"):
        result = HakimiSuiteRunner().run(project, "chinese", mode=mode, confirm_overwrite=True)
        book = openpyxl.load_workbook(result.excel_dir / "replace_text.xlsx")
        try:
            pools[mode] = {cell.value for cell in book.active["A"][1:]}
        finally:
            book.close()
    assert "Only external words" not in pools["1"]
    assert "Only external words" in pools["2"]
    assert "Only deep words" not in pools["2"]
    assert "Only deep words" in pools["3"]


def test_review_name_context_uses_longest_actual_definition(tmp_path: Path):
    (tmp_path / "chars.rpy").write_text(
        'define e = Character("Eileen")\n' +
        '\n' + '\n' + '\n' +
        'define e = Character("Eileen")\n' +
        'default mood = "happy"\n' +
        'default history = "Much longer second definition context"\n' +
        'label start:\n', encoding="utf-8"
    )
    result = ops.extract_character_names(tmp_path)
    assert result["count"] == 1
    assert "Much longer second definition context" in result["entries"][0]["context"]


def test_review_structure_cancel_after_first_workbook_reports_partial(tmp_path: Path, monkeypatch):
    project = _game(tmp_path)
    state = {"cancelled": False}
    original = HakimiSuiteRunner._save_to_excel

    def save_then_cancel(self, *args, **kwargs):
        original(self, *args, **kwargs)
        state["cancelled"] = True

    monkeypatch.setattr(HakimiSuiteRunner, "_save_to_excel", save_then_cancel)
    with pytest.raises(ops.AssetSuiteCancelled) as captured:
        HakimiSuiteRunner().run(project, "chinese", cancel_check=lambda: state["cancelled"])
    result = captured.value.result
    assert result["partial"] is True
    assert result["written"] == [str(project / "translate_output" / "1_Excels" / "names.xlsx")]
    assert not (project / "translate_output" / "1_Excels" / "others.xlsx").exists()


def test_review_structure_failed_second_write_keeps_backup_and_counts(tmp_path: Path, monkeypatch):
    project = _game(tmp_path)
    output = project / "translate_output"
    output.mkdir()
    (output / "manual.rpy").write_bytes(b"manual work")
    original = HakimiSuiteRunner._save_to_excel
    calls = {"value": 0}

    def fail_second(self, *args, **kwargs):
        calls["value"] += 1
        if calls["value"] == 2:
            raise OSError("模拟磁盘写入失败")
        original(self, *args, **kwargs)

    monkeypatch.setattr(HakimiSuiteRunner, "_save_to_excel", fail_second)
    with pytest.raises(ops.AssetSuiteError) as captured:
        HakimiSuiteRunner().run(project, "chinese", confirm_overwrite=True)
    result = captured.value.result
    assert result["partial"] is True
    assert len(result["written"]) == 1
    assert (Path(result["backup_path"]) / "manual.rpy").read_bytes() == b"manual work"


def test_review_emoji_preserves_bom_and_mixed_newlines(tmp_path: Path):
    target = tmp_path / "a.rpy"
    raw = b'\xef\xbb\xbf' + b'old "{b}Name{/b}"\r\nnew "Name"\n'
    target.write_bytes(raw)
    assert apply_replacements_dir(tmp_path, {"{b}": "EMOJI"}) == (1, 0, 1)
    assert target.read_bytes() == raw.replace(b"{b}", b"EMOJI")


def test_review_structure_other_selected_project_does_not_mix_config(review_api, tmp_path: Path):
    client, config, _, project = review_api
    other = tmp_path / "Other"
    (other / "game").mkdir(parents=True)
    (other / "game" / "script.rpy").write_text('define e = Character("Alice")\n', encoding="utf-8")
    before = (config.renpy_project_path, config.renpy_game_folder, config.renpy_tl_folder)
    response = client.post("/api/asset-suite/structure", json={"path": str(other), "confirm_overwrite": True})
    job = _finished_job(client, response.json()["job"]["id"])
    assert job["status"] == "done"
    assert (config.renpy_project_path, config.renpy_game_folder, config.renpy_tl_folder) == before
    assert Path(job["result"]["output_dir"]) == other / "translate_output"


def test_review_api_cancelled_apply_keeps_real_partial_result(review_api, monkeypatch, tmp_path: Path):
    import threading
    import time

    client, _, engine, project = review_api
    target = project / "game" / "tl" / "chinese"
    target.mkdir(parents=True)
    first = target / "a.txt"
    second = target / "b.txt"
    first.write_bytes(b"old")
    second.write_bytes(b"old")
    workbook = _workbook(tmp_path / "fix.xlsx", [
        ["a.txt", "x", "src", "old", "new"],
        ["b.txt", "x", "src", "old", "new"],
    ])
    written = threading.Event()
    original = ops._atomic_write_bytes

    def write_then_wait(*args, **kwargs):
        original(*args, **kwargs)
        written.set()
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline and not kwargs["cancel_check"]():
            time.sleep(0.005)

    monkeypatch.setattr(ops, "_atomic_write_bytes", write_then_wait)
    response = client.post("/api/asset-suite/corrections/apply", json={
        "workbook": str(workbook), "translation_root": str(target), "confirm": True,
    })
    job_id = response.json()["job"]["id"]
    assert written.wait(2)
    assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 200
    job = _finished_job(client, job_id)
    assert job["status"] == "cancelled"
    assert job["result"]["applied_changes"] == 1
    assert job["result"]["written"] == [str(first)]
    assert job["result"]["backups"] == [str(first.with_suffix(".txt.bak"))]
    assert second.read_bytes() == b"old"
    assert engine.get_status() == engine.Status.IDLE
