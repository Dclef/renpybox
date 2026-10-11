"""HtmlTextConverter：HTML/TXT/Excel 往返与空内容不写。"""
from __future__ import annotations

from pathlib import Path

import pytest

from module.Tool.HtmlTextConverter import (
    build_html_content,
    excel_to_txt,
    html_to_txt,
    read_html_strings,
    read_html_translation_pairs,
    txt_to_html,
)


def test_txt_html_roundtrip_plain_and_payload(tmp_path: Path) -> None:
    txt = tmp_path / "a.txt"
    txt.write_text("one\n\nthree\n", encoding="utf-8")
    html = tmp_path / "a.html"
    count, written = txt_to_html(txt, html, keep_data=True)
    assert count == 3
    assert written.is_file()
    strings = read_html_strings(html)
    assert strings == ["one", "", "three"]
    pairs = read_html_translation_pairs(html)
    assert pairs[0] == ("one", "one")

    out = tmp_path / "b.txt"
    n, _ = html_to_txt(html, out, pairs=False)
    assert n == 3
    assert out.read_text(encoding="utf-8").split("\n") == ["one", "", "three"]


def test_html_pairs_without_payload_fails(tmp_path: Path) -> None:
    html = tmp_path / "plain.html"
    html.write_text(build_html_content(["a", "b"], keep_data=False), encoding="utf-8")
    with pytest.raises(ValueError, match="payload"):
        html_to_txt(html, tmp_path / "x.txt", pairs=True)
    assert not (tmp_path / "x.txt").exists()


def test_legacy_literal_eval_payload(tmp_path: Path) -> None:
    html = tmp_path / "legacy.html"
    # 旧版 str(list) 形态
    payload = "[{'line': 1, 'original': 'b', 'target': 'b', 'current': 'B'}, {'line': 0, 'original': 'a', 'current': 'A'}]"
    html.write_text(
        f"<html><body><div id='data'>{payload}</div></body></html>",
        encoding="utf-8",
    )
    strings = read_html_strings(html)
    assert strings == ["A", "B"]
    pairs = read_html_translation_pairs(html)
    assert pairs == [("a", "A"), ("b", "B")]


def test_excel_aliases_skip_metadata_and_empty(tmp_path: Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    meta = book.active
    meta.title = "Metadata"
    meta["A1"] = "x"
    sheet = book.create_sheet("Sheet1")
    sheet["A1"] = "Source (Do Not Edit)"
    sheet["B1"] = "Translation (Do Not Edit)"
    sheet["A2"] = "hello"
    sheet["B2"] = "你好"
    sheet["A3"] = "skip"
    sheet["B3"] = None
    path = tmp_path / "book.xlsx"
    book.save(path)

    out = tmp_path / "out.txt"
    count, written = excel_to_txt(path, out, column="译文")
    assert count == 1
    assert written.read_text(encoding="utf-8") == "你好"

    empty_book = Workbook()
    empty_book.active.title = "Metadata"
    empty_path = tmp_path / "empty.xlsx"
    empty_book.save(empty_path)
    with pytest.raises(ValueError, match="可导出"):
        excel_to_txt(empty_path, tmp_path / "empty.txt", column="译文")
    assert not (tmp_path / "empty.txt").exists()


def test_empty_txt_does_not_write_html(tmp_path: Path) -> None:
    txt = tmp_path / "empty.txt"
    txt.write_text("", encoding="utf-8")
    with pytest.raises(ValueError, match="空"):
        txt_to_html(txt, tmp_path / "out.html")
    assert not (tmp_path / "out.html").exists()
